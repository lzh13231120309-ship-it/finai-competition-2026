import json
from pathlib import Path

import pytest

from server import config
from server import finance_service as svc


@pytest.fixture
def task(tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'DATA_DIR', tmp_path)
    monkeypatch.setattr(config, 'UPLOAD_DIR', tmp_path / 'uploads')
    monkeypatch.setattr(config, 'LOG_DIR', tmp_path / 'logs')
    folder = config.UPLOAD_DIR / 'testtask'
    folder.mkdir(parents=True)
    (folder / '公告.txt').write_text('甲公司于2025年10月9日质押100万股。', encoding='utf-8')
    return 'testtask'


def test_uploaded_file_boundary(task):
    with pytest.raises(ValueError):
        svc.resolve_attachment(task, '../secret.txt')
    with pytest.raises(ValueError):
        svc.resolve_attachment('../outside', '公告.txt')
    assert svc.resolve_attachment(task, '公告.txt').name == '公告.txt'


def test_every_field_requires_real_support(task):
    doc = svc.get_document(task, '公告.txt')
    evidence = doc['pages'][0]['blocks'][0]['evidence_id']
    valid = svc.validate_structured_output(task, {'filename': '公告.txt', 'required_fields': ['公司', '质押股数'],
        'records': [{'公司': {'value': '甲公司', 'evidence_ids': [evidence]},
                     '质押股数': {'value': '100万股', 'evidence_ids': [evidence]}}]})
    assert valid['valid'] is True
    bad = svc.validate_structured_output(task, {'filename': '公告.txt', 'required_fields': ['公司', '质押股数'],
        'records': [{'公司': {'value': '甲公司', 'evidence_ids': [evidence]},
                     '质押股数': {'value': '999万股', 'evidence_ids': [evidence]}}]})
    assert bad['valid'] is False
    assert bad['review_queue']
    missing = svc.validate_structured_output(task, {'filename': '公告.txt', 'required_fields': ['公司'],
        'records': [{'公司': {'value': '甲公司', 'evidence_ids': ['invented']}}]})
    assert missing['valid'] is False


def test_runs_and_exports_preserve_all_records(task):
    records = [{'名称': {'value': '=HYPERLINK("evil")', 'evidence_ids': []}} for _ in range(350)]
    run = svc.save_run(task, {'kind': 'structured', 'records': records, 'warnings': ['需人工复核']})
    for fmt in ('json', 'xlsx', 'docx', 'md'):
        output = svc.export_financial_report(task, {'run_id': run['run_id'], 'format': fmt})
        path = Path(output['path'])
        assert path.exists()
    result = json.loads((svc.run_dir(task, run['run_id']) / '报告.json').read_text())
    assert len(result['records']) == 350
    from openpyxl import load_workbook
    wb = load_workbook(svc.run_dir(task, run['run_id']) / '报告.xlsx')
    assert wb['结构化数据'].max_row == 351
    assert wb['结构化数据']['A2'].data_type != 'f'


def test_run_isolation(task):
    run = svc.save_run(task, {'kind': 'test'})
    with pytest.raises((FileNotFoundError, ValueError)):
        svc.load_run('differenttask', run['run_id'])
    with pytest.raises(ValueError):
        svc.run_dir(task, '../outside')


def test_full_run_can_be_read_without_shell_or_preview_loss(task):
    run = svc.save_run(task, {'kind':'test', 'facts':[{'code':'revenue','value':str(i)} for i in range(80)]})
    page = svc.read_financial_run(task, {'run_id':run['run_id'], 'section':'facts', 'offset':75, 'limit':10})
    assert [f['value'] for f in page['items']] == ['75','76','77','78','79']
    assert page['total'] == 80 and page['has_more'] is False


def test_human_confirmation_is_new_run_and_does_not_rewrite_source(task):
    facts = [{'fact_id':'f1','code':'revenue','raw_label':'营业收入','company':'甲公司','period':'2025-12-31',
        'period_type':'annual','scope':'consolidated','currency':'CNY','unit':'元','scale':1,
        'value':'100','raw_value':'100','status':'needs_review','confidence':.62,'source':{'document_id':'doc'},'evidence_ids':['e1']}]
    reviews = [{'code':'unresolved_metadata','message':'营业收入 缺少或待确认：ocr_confidence。','fact_ids':['f1']}]
    original = svc.save_run(task, {'kind':'analysis','facts':facts,'review_queue':reviews})
    new = svc.confirm_financial_facts(task, {'run_id':original['run_id'],'fact_ids':['f1']})
    assert new['run_id'] != original['run_id']
    assert svc.load_run(task, original['run_id'])['facts'][0]['status'] == 'needs_review'
    confirmed = svc.load_run(task,new['run_id'])
    assert confirmed['facts'][0]['value'] == '100'
    assert confirmed['facts'][0]['status'] == 'verified'
    assert confirmed['human_confirmations'][0]['fact_id'] == 'f1'


def test_missing_units_cannot_be_confirmed_by_single_click(task):
    run=svc.save_run(task,{'kind':'analysis','facts':[{'fact_id':'bad','unit':'unknown','status':'needs_review'}]})
    with pytest.raises(ValueError):
        svc.confirm_financial_facts(task, {'run_id':run['run_id'],'fact_ids':['bad']})


def _valuation_run():
    """一份最小但完整的估值运行，用于渲染测试。

    字段与 valuation.from_metrics 的真实返回保持一致（三情景在**顶层**），
    这样测试才能抓住"读 run['scenarios'] 导致估值情景表永远不生成"这类错误。
    """
    return {
        'run_id': 'v1', 'kind': 'valuation', 'created_at': 1791576235.0,
        'parent_run_id': 'a1', 'forensic_run_id': 'f1',
        'base': {
            'label': '基准情景',
            'parameters': {'revenue_base': '52552910000', 'revenue_growth': '0.153386',
                           'gross_margin': '0.6052', 'opex_ratio': '0.2331',
                           'tax_rate': '0.2413', 'capex_to_revenue': '0.1233',
                           'net_debt': '-17772518000', 'wacc': '0.10',
                           'terminal_growth': '0.03', 'years': '5',
                           'receivable_days': '4.096', 'inventory_days': '95.535',
                           'payable_days': '27.744', 'nwc_to_revenue': '0.0845'},
            'valuation': {
                'enterprise_value': '253287484790.16187675360',
                'equity_value': '278615356790.16187675360',
                'value_per_share': '24.77435', 'terminal_value_share': '0.72',
                'multiples': {
                    'implied_ev_ebitda': '10.39040', 'implied_pe': '17.55801',
                    'comparables': {'pe': {'multiple': '25', 'basis': '15868274000',
                                           'equity_value': '396706850000',
                                           'value_per_share': '35.275'}},
                    'comparable_range': {'low': '396706850000', 'high': '396706850000',
                                         'text': '可比倍数法隐含权益价值 396706850000 元'},
                    'dcf_vs_comparable_gap': '-0.2975', 'warnings': []}},
            'projection': [{'year': 1, 'revenue': '60613794572.29', 'ebit': '22555060988.88',
                            'fcff': '12866505841.26'}],
            'sensitivity': {
                'wacc_axis': ['0.08', '0.10'], 'growth_axis': ['0.02', '0.03'],
                'matrix': [{'terminal_growth': '0.02',
                            'cells': [{'wacc': '0.08', 'equity_value': '300000000000', 'ok': True},
                                      {'wacc': '0.10', 'equity_value': '250000000000', 'ok': True}]},
                           {'terminal_growth': '0.03',
                            'cells': [{'wacc': '0.08', 'equity_value': '320000000000', 'ok': True},
                                      {'wacc': '0.10', 'equity_value': '260000000000', 'ok': True}]}],
                'note': '每个单元格是按对应 WACC 与永续增长率重算的权益价值（元）。'}},
        'optimistic': {'label': '乐观情景', 'parameters': {}, 'valuation': {'multiples': {}},
                       'projection': []},
        'stress': {'label': '压力情景', 'parameters': {}, 'valuation': {'multiples': {}},
                   'projection': [],
                   'risk_adjustments': [{'source': 'R_EXP_01 期间费用率异常下降',
                                         'severity': 'low', 'parameter': 'opex_ratio',
                                         'from': '0.2331', 'to': '0.2346',
                                         'path': '费用率上调 → EBIT 与终值下降'}],
                   'unmatched_risks': []},
        'range': {'equity_value_low': '277154690900.2038204547109466',
                  'equity_value_high': '324365701081.7261202422744908',
                  'equity_value_mid': '278615356790.1618767536030606',
                  'value_per_share_low': '24.64447417911913990529764899',
                  'value_per_share_high': '28.84245660405886334341399112',
                  'text': '三情景权益价值区间 277154690900.2038204547109466 – 324365701081.7261202422744908 元',
                  'disclaimer': '这是按情景分布的合理价值区间，不是目标价，也不构成投资建议。'},
        'inputs_used': {'revenue_base': {'source': 'fact', 'detail': '最新一期营业收入',
                                         'period': '2025-12-31'}},
        'data_quality': {'from_company_data': ['revenue_base', 'net_debt'],
                         'from_default_assumptions': ['wacc', 'terminal_growth'],
                         'overridden_by_user': [], 'blocking_missing': [],
                         'text': '驱动变量中 2 项取自公司已核验数据，2 项降级或沿用默认假设。'},
        'degraded_inputs': [{'parameter': 'minority_interest', 'label': '少数股东权益',
                             'impact': '沿用默认假设，结果可信度下降。'}],
        'calibration_notes': [], 'calibration_warnings': [], 'missing_inputs': [],
        'calibration_checks': [{'code': 'ebit_calibration', 'status': 'pass',
                                'implied_ebit': '19555517000', 'reported_ebit': '20984492000',
                                'difference_ratio': '-0.068', 'note': '差异来自其他收益等项目。'}],
        'stress_diagnostic': {'adjustment_count': 1,
                              'text': '压力情景相对基准情景变动 -0.52%',
                              'insufficient': True,
                              'reason': '压力情景与基准情景的权益价值差异不足 2%，'
                                        '本次压力测试没有构成实质压力。'},
        'ok': True,
    }


def test_valuation_markdown_has_every_required_section():
    md = svc._markdown(_valuation_run())
    for heading in ('一、估值结论', '二、驱动变量与来源', '三、数据质量声明',
                    '四、逐年预测与 FCFF', '五、敏感性矩阵', '六、压力情景传导',
                    '七、倍数交叉验证', '八、模型自检', '九、口径与边界说明'):
        assert heading in md, heading


def test_valuation_markdown_never_leaks_full_precision_decimals():
    """报告里的金额必须格式化；run.json 里的全精度值不得直接贴进交付物。

    回归背景：旧导出把顶层键铺成一张表，一行里塞进
    277154690900.2038204547109466 这种 28 位小数，人根本读不了。
    """
    md = svc._markdown(_valuation_run())
    assert '277154690900.2038204547109466' not in md
    assert '278,615,356,790.16' in md
    assert '15.34%' in md and '12.33%' in md


def test_valuation_markdown_separates_company_data_from_assumptions():
    md = svc._markdown(_valuation_run())
    assert '默认假设' in md
    assert '工具未能从本次数据校准' in md or '引擎未能从本次数据校准' in md
    assert '原文事实' in md


def test_valuation_markdown_reports_insufficient_stress():
    md = svc._markdown(_valuation_run())
    assert '没有构成实质压力' in md
    assert 'R_EXP_01' in md


def test_valuation_sheets_read_scenarios_from_top_level():
    """估值情景表必须从顶层 base/optimistic/stress 取数。

    回归背景：旧实现读 run['scenarios']，而估值运行把三情景存在顶层，
    于是"估值情景"表**从来没生成过**，导出的 xlsx 里只有价值区间，
    而且这一点不会报错。
    """
    sheets = svc._sheets(_valuation_run())
    assert '估值情景' in sheets
    assert len(sheets['估值情景']) == 3
    assert sheets['估值情景'][0]['企业价值'] == '253287484790.16187675360'
    for name in ('假设来源', '数据质量', '压力传导', '压力有效性', '模型自检',
                 '倍数交叉验证', '敏感性矩阵', '价值区间'):
        assert name in sheets, name
