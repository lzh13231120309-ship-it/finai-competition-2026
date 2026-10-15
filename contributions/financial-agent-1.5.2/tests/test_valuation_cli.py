"""Exercise CLI behavior with offline inputs, including failure diagnostics."""
import csv
import json
import math
import subprocess
import sys
from decimal import Decimal
from pathlib import Path

PROJECT = Path(__file__).resolve().parents[1]


def command(*arguments):
    return subprocess.run([sys.executable, '-m', 'server.valuation_cli', *map(str, arguments)],
                          cwd=PROJECT, text=True, capture_output=True, timeout=60)


def inputs(tmp_path, **overrides):
    analysis = tmp_path / '已核验分析.json'
    analysis.write_text(json.dumps({'run_id':'analysis-fixture', 'company':'演示公司', 'facts':[]}, ensure_ascii=False), encoding='utf-8')
    options = tmp_path / '模型选项.json'
    parameters = dict(revenue_base='1000', revenue_growth='0', gross_margin='.3', opex_ratio='.1',
                      tax_rate='0', da_to_revenue='0', capex_to_revenue='0', receivable_days='0',
                      inventory_days='0', payable_days='0', wacc='.1', terminal_growth='0', shares='100')
    parameters.update(overrides)
    options.write_text(json.dumps({'overrides':parameters,'forecast':{'years':5,'terminal_roic':'.2'}}), encoding='utf-8')
    return analysis, options


def test_help_discovers_both_commands():
    result = command('--help')
    assert result.returncode == 0, result.stderr
    assert 'build' in result.stdout and 'market' in result.stdout


def test_build_json_preserves_identity_and_recalculable_model(tmp_path):
    analysis, options = inputs(tmp_path)
    folder = tmp_path / '估值结果'
    result = command('build', '--run', analysis, '--options', options, '--output', folder, '--formats', 'json')
    assert result.returncode == 0, result.stderr
    summary = json.loads(result.stdout)
    output = json.loads((folder/'run.json').read_text(encoding='utf-8'))
    assert output['run_id'] == summary['run_id']
    assert output['kind'] == 'valuation_model'
    assert output['company'] == '演示公司'
    assert output['parent_run_id'] == 'analysis-fixture'
    assert Decimal(output['scenarios']['base']['valuation']['enterprise_value']) == Decimal(2000)
    assert Path(summary['exports']['json']).is_file()


def test_blocked_model_keeps_diagnostics_and_returns_nonzero(tmp_path):
    analysis, options = inputs(tmp_path, wacc='.02', terminal_growth='.03')
    folder = tmp_path/'blocked'
    result = command('build','--run',analysis,'--options',options,'--output',folder,'--formats','json')
    assert result.returncode == 2, result.stderr
    output = json.loads((folder/'run.json').read_text(encoding='utf-8'))
    assert not output['ok']
    assert output['model_plan']['status'] == 'blocked'
    assert json.loads(result.stdout)['ok'] is False


def test_findings_from_saved_forensic_run_change_stress(tmp_path):
    analysis, options = inputs(tmp_path)
    findings = tmp_path/'forensic.json'
    findings.write_text(json.dumps({'run_id':'forensic-fixture','findings':[{'rule_id':'R_MARGIN_01','severity':'critical'}]}),encoding='utf-8')
    folder = tmp_path/'with-findings'
    result = command('build','--run',analysis,'--options',options,'--findings',findings,'--output',folder,'--formats','json')
    assert result.returncode == 0, result.stderr
    output = json.loads((folder/'run.json').read_text())
    assert output['forensic_run_id'] == 'forensic-fixture'
    assert Decimal(output['scenarios']['stress']['valuation']['equity_value']) < Decimal(output['scenarios']['base']['valuation']['equity_value'])


def test_invalid_json_is_actionable_without_traceback(tmp_path):
    bad = tmp_path/'broken.json'; bad.write_text('{broken',encoding='utf-8')
    result = command('build','--run',bad,'--output',tmp_path/'output','--formats','json')
    assert result.returncode == 2
    assert 'broken.json' in result.stderr
    assert 'Traceback' not in result.stderr


def test_unknown_export_format_does_not_start_build(tmp_path):
    analysis, options = inputs(tmp_path)
    folder = tmp_path/'output'
    result = command('build','--run',analysis,'--options',options,'--output',folder,'--formats','csv')
    assert result.returncode == 2
    assert not folder.exists()


def test_local_market_csv_beta_without_network(tmp_path):
    from datetime import date,timedelta
    stock = tmp_path/'股票行情.csv'; bench = tmp_path/'基准行情.csv'
    with stock.open('w',newline='',encoding='utf-8') as a, bench.open('w',newline='',encoding='utf-8') as b:
        sw,bw=csv.writer(a),csv.writer(b)
        sw.writerow(['date','close','adj_close']); bw.writerow(['date','close','adj_close'])
        for i in range(85):
            day=(date(2025,1,1)+timedelta(days=i)).isoformat()
            level=math.exp(i*.004+math.sin(i*.3)*.02)
            sw.writerow([day,100*level**2,100*level**2]); bw.writerow([day,100*level,100*level])
    output = tmp_path/'market.json'
    result = command('market','--symbol','DEMO','--benchmark','INDEX','--price-file',stock,
                     '--benchmark-file',bench,'--currency','CNY','--output',output)
    assert result.returncode == 0, result.stderr
    data=json.loads(output.read_text())
    assert data['ok'] and data['beta_estimate']['ok']
    assert abs(data['beta_estimate']['beta']-2) < 1e-10
    assert data['as_of'] == '2025-03-26'
    assert 'risk_free_rate' not in data['market_fields']
    assert json.loads(result.stdout)['ok']


def test_all_report_formats_are_real_files(tmp_path):
    import openpyxl
    from docx import Document
    analysis, options = inputs(tmp_path)
    folder=tmp_path/'reports'
    result=command('build','--run',analysis,'--options',options,'--output',folder)
    assert result.returncode == 0, result.stderr + result.stdout
    summary=json.loads(result.stdout)
    assert set(summary['exports']) == {'xlsx','docx','json','md'}
    wb=openpyxl.load_workbook(summary['exports']['xlsx'],data_only=False)
    assert any(c.data_type=='f' for ws in wb.worksheets for row in ws for c in row)
    wb.close()
    assert Document(summary['exports']['docx']).paragraphs
    assert '自动估值' in Path(summary['exports']['md']).read_text(encoding='utf-8')
    assert json.loads(Path(summary['exports']['json']).read_text())['run_id'] == summary['run_id']


def test_existing_run_is_not_overwritten(tmp_path):
    analysis, options=inputs(tmp_path)
    folder=tmp_path/'saved'; folder.mkdir()
    previous=folder/'run.json'; previous.write_text('{"run_id":"old-run"}',encoding='utf-8')
    result=command('build','--run',analysis,'--options',options,'--output',folder,'--formats','json')
    assert result.returncode == 2
    assert json.loads(previous.read_text())['run_id'] == 'old-run'


def test_foreign_forensic_parent_is_rejected(tmp_path):
    analysis, options=inputs(tmp_path)
    foreign=tmp_path/'foreign.json'
    foreign.write_text(json.dumps({'run_id':'foreign','parent_run_id':'other-company','findings':[]}),encoding='utf-8')
    result=command('build','--run',analysis,'--options',options,'--findings',foreign,'--output',tmp_path/'output','--formats','json')
    assert result.returncode == 2
    assert not (tmp_path/'output'/'run.json').exists()


def test_sourced_field_market_inputs_are_supported_offline(tmp_path):
    analysis, options=inputs(tmp_path)
    settings=json.loads(options.read_text())
    del settings['overrides']['wacc']
    settings['market']={name:{'value':value,'source':'explicit test snapshot','as_of':'2026-10-01','currency':'CNY'}
        for name,value in {'risk_free_rate':'.03','equity_risk_premium':'.05','beta':'1.2','cost_of_debt':'.04','debt_weight':'.25'}.items()}
    options.write_text(json.dumps(settings),encoding='utf-8')
    result=command('build','--run',analysis,'--options',options,'--output',tmp_path/'sourced','--formats','json')
    assert result.returncode == 0, result.stderr + result.stdout
    output=json.loads((tmp_path/'sourced'/'run.json').read_text())
    assert output['discount_rate']['method'] == 'CAPM'
    assert Decimal(output['discount_rate']['wacc']) == Decimal('.0775')
    assert output['discount_rate']['market_inputs']['field_sources']['beta']['source'] == 'explicit test snapshot'
