"""Complete local document extraction with stable, source-addressable evidence.

PDF boxes use top-left PDF points; image boxes use top-left pixels. Spreadsheet
coordinates use A1 addresses. OCR table recovery is conservative and always
requires review; cached extraction never discards rows, columns or pages.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import statistics
import subprocess
import tempfile
from pathlib import Path
from typing import Any

PARSER_VERSION = 'financial-documents-v5'
MAX_FILE_BYTES = 150 * 1024 * 1024
MAX_PAGES = 2000
MAX_CELLS = 2_000_000
MAX_IMAGE_PIXELS = 80_000_000
MAX_ARCHIVE_BYTES = 300 * 1024 * 1024
IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.bmp', '.tif', '.tiff', '.webp', '.gif', '.heic'}
TEXT_EXTS = {'.txt', '.md', '.markdown', '.json', '.yaml', '.yml', '.log', '.ini', '.toml', '.xml', '.html', '.htm', '.py', '.js', '.ts', '.java', '.c', '.cpp', '.go', '.rs', '.sql', '.sh', '.r', '.m'}
_NUMERIC = re.compile(r'^[（(]?[+\-−－]?[¥￥$€£]?\d[\d,，]*(?:[.．]\d+)?[%％]?[）)]?$')


def _check_archive(path: Path) -> None:
    import zipfile
    with zipfile.ZipFile(path) as archive:
        size = sum(item.file_size for item in archive.infolist())
        if size > MAX_ARCHIVE_BYTES:
            raise ValueError(f'文档解压大小 {size} 超过资源上限 {MAX_ARCHIVE_BYTES}；未进行截断解析')


def _decode(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ('utf-8-sig', 'gb18030', 'utf-16'):
        try:
            return data.decode(encoding)
        except UnicodeError:
            pass
    raise ValueError('无法确定文本编码，请另存为 UTF-8 后重新导入')


def _value(value: Any) -> str:
    if value is None:
        return ''
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    if isinstance(value, float):
        return str(value) if math.isfinite(value) else repr(value)
    return str(value)


def _display(value: Any, number_format: str) -> str:
    # Percent display carries the scale explicitly; raw values remain in metadata.
    cleaned = re.sub(r'"[^"]*"|\\.', '', number_format)
    if isinstance(value,(float,int)) and not isinstance(value,bool) and '%' in cleaned:
        match = re.search(r'\.([0#]+)',cleaned)
        decimals = len(match.group(1)) if match else 0
        return f'{value*100:.{decimals}f}%'
    return _value(value)


def _column(index: int) -> str:
    result = ''
    while index:
        index, remainder = divmod(index-1, 26)
        result = chr(65+remainder) + result
    return result


def _bbox(items: list[dict]) -> list[float] | None:
    boxes = [item['bbox'] for item in items if item.get('bbox')]
    if not boxes:
        return None
    return [min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes)]


def _intersects(a: list | None, b: list | None) -> bool:
    return bool(a and b and a[0] <= b[2] and a[2] >= b[0] and a[1] <= b[3] and a[3] >= b[1])


def _blocks(doc: dict, page: int, items: list[dict]) -> list[dict]:
    return [dict(item, evidence_id=f"{doc['sha256'][:16]}:p{page}:b{i+1}") for i, item in enumerate(items)]


def _page(doc: dict, number: int, text: str, items: list[dict], method: str = 'text', warnings: list | None = None, **extra) -> dict:
    position = 0
    for item in items:
        if 'start' not in item:
            found = text.find(item['text'], position)
            if found < 0: found = text.find(item['text'])
            if found >= 0:
                item.update(start=found,end=found+len(item['text']))
                position = found+len(item['text'])
    page = dict(page=number, text=text, method=method, blocks=_blocks(doc, number, items), warnings=warnings or [], **extra)
    doc['pages'].append(page)
    return page


def _table(doc: dict, page: dict, rows: list[list[str]], method: str, bbox=None, confidence=1.0, **extra) -> dict:
    ordinal = len([t for t in doc['tables'] if t['page'] == page['page']]) + 1
    context_blocks = [b for b in page['blocks'] if not b.get('table_id') and (not bbox or (b.get('bbox') and bbox[1]-120 <= b['bbox'][1] <= bbox[3]+25))]
    context = '\n'.join(b['text'] for b in context_blocks) if bbox else page['text']
    table = dict(table_id=f"{doc['sha256'][:16]}:p{page['page']}:t{ordinal}", page=page['page'], rows=rows,
                 bbox=bbox, method=method, confidence=confidence,
                 evidence_ids=[b['evidence_id'] for b in page['blocks'] if not bbox or _intersects(b['bbox'], bbox)],
                 context=context, context_evidence_ids=[b['evidence_id'] for b in context_blocks], **extra)
    cell_ids = []; source_position = 0
    for r,row in enumerate(rows):
        ids = []
        for c,value in enumerate(row):
            eid = f"{table['table_id']}:r{r+1}:c{c+1}"
            boxes = extra.get('cell_bboxes',[])
            box = boxes[r][c] if r < len(boxes) and c < len(boxes[r]) else bbox
            block = dict(evidence_id=eid,text=value,bbox=box,confidence=confidence,table_id=table['table_id'],row=r+1,column=c+1)
            coords = extra.get('cell_coordinates',[])
            if r < len(coords) and c < len(coords[r]): block['coordinate'] = coords[r][c]
            if extra.get('sheet'): block['sheet'] = extra['sheet']
            for matrix,key in [('number_formats','number_format'),('raw_values','raw_value')]:
                if matrix in extra: block[key] = extra[matrix][r][c]
            found = page['text'].find(value,source_position) if value else -1
            if found >= 0:
                block.update(start=found,end=found+len(value))
                source_position = found+len(value)
            elif not value:
                block.update(start=source_position,end=source_position)
            page['blocks'].append(block); ids.append(eid)
        cell_ids.append(ids)
    table['cell_evidence_ids'] = cell_ids
    table['evidence_ids'] = list(dict.fromkeys(table['evidence_ids'] + [eid for row in cell_ids for eid in row]))
    doc['tables'].append(table)
    return table


def _rows_text(rows: list[list[str]]) -> str:
    return '\n'.join('\t'.join(row) for row in rows)


def _line_items(text: str) -> list[dict]:
    # Keep line offsets so read_document can attach only relevant evidence.
    result = []; offset = 0
    for line in text.splitlines(keepends=True):
        result.append(dict(text=line.rstrip('\r\n'), bbox=None, confidence=1.0, start=offset, end=offset+len(line)))
        offset += len(line)
    return result


def _parse_delimited(doc: dict, path: Path) -> None:
    text = _decode(path)
    rows = list(csv.reader(io.StringIO(text, newline=''), delimiter='\t' if path.suffix.lower() == '.tsv' else ','))
    width = max((len(r) for r in rows), default=0)
    if len(rows)*width > MAX_CELLS:
        raise ValueError('表格单元格数量超过资源上限；未进行截断解析')
    normalized = [row + ['']*(width-len(row)) for row in rows]
    page = _page(doc, 1, _rows_text(normalized), _line_items(_rows_text(normalized)))
    _table(doc, page, normalized, 'csv', cell_coordinates=[[_column(c+1)+str(r+1) for c in range(width)] for r in range(len(rows))])
    if any(len(row) != width for row in rows):
        doc['warnings'].append('分隔文本各行列数不同，已用空单元格补齐；请核对列对齐')


def _parse_excel(doc: dict, path: Path) -> None:
    import openpyxl
    _check_archive(path)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    count = 0
    try:
        for pnum, ws in enumerate(formulas.worksheets, 1):
            width, height = ws.max_column or 0, ws.max_row or 0
            count += width*height
            if count > MAX_CELLS:
                raise ValueError('Excel 单元格数量超过资源上限；未进行截断解析')
            rows = []; coords = []; formats = []; raw_values = []; formula_map = {}; warning = []; items = []
            cached_sheet = values[ws.title]
            for row, cache_row in zip(ws.iter_rows(), cached_sheet.iter_rows()):
                output = []; addresses = []; row_formats = []; row_raw = []
                for cell, cached in zip(row, cache_row):
                    address = f'{_column(cell.column)}{cell.row}' if hasattr(cell, 'column') else f'{_column(len(output)+1)}{len(rows)+1}'
                    raw = cell.value
                    if cell.data_type == 'f':
                        formula_map[address] = raw
                        if cached.value is None:
                            output.append(raw)
                            warning.append(f'{ws.title}!{address} 公式没有已计算缓存，保留公式，不推算结果')
                        else:
                            output.append(_display(cached.value,cell.number_format or 'General'))
                    else:
                        output.append(_display(raw,cell.number_format or 'General'))
                    if cell.data_type == 'e' or cached.data_type == 'e':
                        warning.append(f'{ws.title}!{address} 包含 Excel 错误值 {_value(cached.value or raw)}，不得作为有效财务数字')
                    row_formats.append(cell.number_format or 'General'); row_raw.append(_value(cached.value if cell.data_type == 'f' else raw))
                    addresses.append(address)
                rows.append(output); coords.append(addresses); formats.append(row_formats); raw_values.append(row_raw)
                items.append(dict(text='\t'.join(output), bbox=None, confidence=1.0, sheet=ws.title, cell_range=f'A{len(rows)}:{_column(width)}{len(rows)}'))
            page = _page(doc, pnum, f'工作表：{ws.title}\n'+_rows_text(rows), items, warnings=warning, sheet=ws.title, coordinate_system='A1')
            _table(doc, page, rows, 'xlsx', sheet=ws.title, cell_coordinates=coords, number_formats=formats, raw_values=raw_values, formulas=formula_map,
                   merged_cells=[str(r) for r in getattr(ws, 'merged_cells', []).ranges] if hasattr(ws, 'merged_cells') else [])
            doc['warnings'].extend(warning)
        doc['warnings'].append('Excel 公式缓存可能早于文件修改时间；未重新计算公式，请核对原工作簿计算状态')
    finally:
        formulas.close(); values.close()


def _parse_word(doc: dict, path: Path) -> None:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    _check_archive(path)
    document = Document(path)
    items = []; parts = []; tables = []
    for element in document.element.body.iterchildren():
        if element.tag.endswith('}p'):
            text = Paragraph(element, document).text
            parts.append(text); items.append(dict(text=text, bbox=None, confidence=1.0))
        elif element.tag.endswith('}tbl'):
            table = Table(element, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if sum(len(row) for row in rows) > MAX_CELLS:
                raise ValueError('Word 表格单元格数量超过资源上限；未进行截断解析')
            parts.append(_rows_text(rows)); items.append(dict(text=_rows_text(rows), bbox=None, confidence=1.0))
            tables.append(rows)
    warning = 'Word 原始页码依赖排版引擎，当前以一个逻辑页保存文档顺序；不作为原始页码引用'
    page = _page(doc, 1, '\n'.join(parts), items, warnings=[warning])
    for rows in tables:
        _table(doc, page, rows, 'docx')
    doc['warnings'].append(warning)


class OcrBackendUnavailable(RuntimeError):
    """某个 OCR 后端在本机不可用（缺程序 / 缺依赖 / 系统不兼容）。

    继承 RuntimeError 是为了让调用方既有的 `except RuntimeError`（见 `_parse_pdf` 的
    降级分支）继续生效；同时它可被识别为"可换下一个后端再试"，与"后端跑了但失败"区分开。
    """


# 后端按优先级排列。**只在这一处做选择**——`_ocr()` 是对外唯一入口，换后端不改调用方，
# 也不会出现"同一次运行两套 OCR 结果"（同输入同输出仍然成立）。
OCR_BACKENDS = ('finocr', 'rapidocr')

_RAPIDOCR_ENGINE = None


def _ocr_finocr(image_path: Path) -> list[dict]:
    """macOS Apple Vision（本机原生，离线）。返回 blocks：[{text,bbox,confidence,tokens}]。"""
    binary = Path(os.environ.get('FINOCR_PATH', str(Path(__file__).parent/'native'/'finocr')))
    if not binary.is_file():
        raise OcrBackendUnavailable(f'未找到本机 OCR 程序 {binary}')
    if not os.access(binary, os.X_OK):
        raise OcrBackendUnavailable(f'本机 OCR 程序不可执行 {binary}')
    try:
        completed = subprocess.run([str(binary), str(image_path)], capture_output=True, text=True, timeout=180)
    except OSError as error:  # 例如 Linux 上遇到 macOS 的 Mach-O 二进制：Exec format error
        raise OcrBackendUnavailable(f'本机 OCR 程序无法在当前系统运行（{error.strerror or error}）') from error
    if completed.returncode:
        raise RuntimeError('离线 OCR 失败：'+completed.stderr.strip())
    return json.loads(completed.stdout)['blocks']


def _rapidocr_engine():
    """惰性加载并缓存 RapidOCR 引擎（模型随 wheel 自带，首次调用会加载约 1–2 秒）。"""
    global _RAPIDOCR_ENGINE
    if _RAPIDOCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR
        _RAPIDOCR_ENGINE = RapidOCR()
    return _RAPIDOCR_ENGINE


def _ocr_rapidocr(image_path: Path) -> list[dict]:
    """跨平台后备后端（ONNX Runtime + PaddleOCR 模型，纯 CPU，离线）。

    **可选依赖**：不装 `rapidocr-onnxruntime` 时抛 OcrBackendUnavailable，不影响 macOS 主路径。
    之所以不做成硬依赖：它会给 .app 增加上百 MB，而主目标环境（macOS）已有原生 OCR。
    """
    try:
        engine = _rapidocr_engine()
    except Exception as error:
        raise OcrBackendUnavailable('未安装可选后端 rapidocr-onnxruntime') from error
    result, _elapse = engine(str(image_path))
    blocks = []
    for item in result or []:
        box, text = item[0], item[1]
        score = float(item[2]) if len(item) > 2 else 1.0
        xs = [float(point[0]) for point in box]
        ys = [float(point[1]) for point in box]
        bbox = [min(xs), min(ys), max(xs), max(ys)]
        token = {'text': str(text), 'bbox': bbox, 'confidence': score}
        blocks.append({**token, 'tokens': [dict(token)]})
    return blocks


def ocr_backends() -> list[dict]:
    """各 OCR 后端的可用性与原因（供 /health、工具说明与打包自检展示）。"""
    out = []
    for name in OCR_BACKENDS:
        if name == 'finocr':
            binary = Path(os.environ.get('FINOCR_PATH', str(Path(__file__).parent/'native'/'finocr')))
            if not binary.is_file():
                out.append({'backend': name, 'available': False, 'reason': f'未找到 {binary}'})
            elif not os.access(binary, os.X_OK):
                out.append({'backend': name, 'available': False, 'reason': f'不可执行 {binary}'})
            else:
                out.append({'backend': name, 'available': True, 'reason': f'使用 {binary}'})
        else:
            try:
                import rapidocr_onnxruntime  # noqa: F401
                out.append({'backend': name, 'available': True, 'reason': '已安装 rapidocr-onnxruntime'})
            except Exception:
                out.append({'backend': name, 'available': False,
                            'reason': '未安装（可选）：pip install rapidocr-onnxruntime'})
    return out


def _ocr(image_path: Path, scale_x=1.0, scale_y=1.0) -> list[dict]:
    """按 OCR_BACKENDS 顺序取第一个可用后端；全部不可用才抛错。

    坐标一律归一到**原图像素**（top-left），所以两个后端对调用方是同构的。
    """
    unavailable = []
    for name in OCR_BACKENDS:
        try:
            items = _ocr_finocr(image_path) if name == 'finocr' else _ocr_rapidocr(image_path)
        except OcrBackendUnavailable as error:
            unavailable.append(f'{name}：{error}')
            continue
        for item in items:
            for entry in [item] + item.get('tokens', []):
                x0,y0,x1,y1 = entry['bbox']
                entry['bbox'] = [round(x0*scale_x, 3), round(y0*scale_y, 3), round(x1*scale_x, 3), round(y1*scale_y, 3)]
        return sorted(items, key=lambda b: (round(b['bbox'][1]/8), b['bbox'][0]))
    raise RuntimeError('本机没有可用的离线 OCR 后端——' + '；'.join(unavailable) +
                       '。macOS 请使用带 server/native/finocr 的版本；其他系统可 pip install rapidocr-onnxruntime。')


def _spatial_rows(tokens: list[dict]) -> list[list[dict]]:
    rows: list[list[dict]] = []
    for token in sorted(tokens, key=lambda t: ((t['bbox'][1]+t['bbox'][3])/2, t['bbox'][0])):
        box = token['bbox']; center = (box[1]+box[3])/2
        target = None
        for row in reversed(rows[-4:]):
            row_center = statistics.median((c['bbox'][1]+c['bbox'][3])/2 for c in row)
            typical_height = statistics.median(c['bbox'][3]-c['bbox'][1] for c in row)
            if abs(center-row_center) <= max(2, min(box[3]-box[1], typical_height)*.55):
                target = row; break
        if target is None:
            rows.append([token])
        else:
            target.append(token)
    return [sorted(row, key=lambda t: t['bbox'][0]) for row in rows]


def _spatial_tables(doc: dict, page: dict, tokens: list[dict], method: str) -> None:
    # Cluster text fragments into cells by gaps; only numeric grids become tables.
    rows = _spatial_rows(tokens)
    grouped = []
    for row in rows:
        cells = []
        for token in row:
            if not cells:
                cells.append(dict(token)); continue
            prior = cells[-1]
            gap = token['bbox'][0]-prior['bbox'][2]
            height = min(token['bbox'][3]-token['bbox'][1], prior['bbox'][3]-prior['bbox'][1])
            if gap <= max(3, height*.65):
                prior['text'] += ' ' + token['text']
                prior['bbox'] = _bbox([prior,token])
                prior['confidence'] = min(prior.get('confidence',1),token.get('confidence',1))
            else:
                cells.append(dict(token))
        grouped.append(cells)
    runs = []; run = []
    for row in grouped:
        if len(row) >= 2:
            run.append(row)
        else:
            if run: runs.append(run)
            run = []
    if run: runs.append(run)
    for run in runs:
        numeric_rows = [row for row in run if sum(bool(_NUMERIC.fullmatch(c['text'].strip())) for c in row) >= 1]
        if len(run) < 3 or len(numeric_rows) < 2:
            continue
        # Require repeated column positions; split columns rather than compress empty cells.
        typical_height = statistics.median(c['bbox'][3]-c['bbox'][1] for row in run for c in row)
        tolerance = max(8,typical_height*1.5)
        layouts = []
        for right in (False,True):
            def cell_x(cell):
                return cell['bbox'][2] if right and _NUMERIC.fullmatch(cell['text'].strip()) else cell['bbox'][0]
            anchors = []
            for row in sorted(run,key=len,reverse=True):
                for cell in row:
                    x = cell_x(cell)
                    if not any(abs(x-a) <= tolerance for a in anchors): anchors.append(x)
            anchors.sort()
            if len(anchors) < 2 or len(anchors) > max(len(r) for r in run)+2: continue
            aligned = []; boxes = []; valid = True
            for row in run:
                output = ['']*len(anchors); cell_boxes = [None]*len(anchors)
                for cell in row:
                    col = min(range(len(anchors)),key=lambda c: abs(cell_x(cell)-anchors[c]))
                    if output[col]: valid = False; break
                    output[col] = cell['text']; cell_boxes[col] = cell['bbox']
                aligned.append(output); boxes.append(cell_boxes)
            if valid: layouts.append((len(anchors),aligned,boxes))
        if not layouts:
            # 兜底：以"数值列"为骨架，其余文本全部并进一个科目名列。
            #
            # 为什么必须有这条：报表科目名按层级缩进，同一页会出现很多个不同的左边界。
            # 实测一页现金流量表有 51 行、30 行含数值，但左边界多达 11 个
            # （子项缩进 + 附註列 + 4 个数值列），而上面的列数上限是"最长行 + 2" = 7，
            # 于是两种布局都被否决，整页被丢弃 —— 现金流量表一张表都没建起来，
            # 购建固定资产支出、投资/融资活动现金流全部静默消失，而报表本身完全正常。
            #
            # 财务报表的稳定结构是"一个科目名列 + 若干右对齐数值列"，比"每个缩进层级
            # 自成一列"更接近真实版式。所以数值列用右边界聚类，其余文本全部归入第 0 列。
            numeric_cells = [c for row in run for c in row
                             if _NUMERIC.fullmatch(c['text'].strip())]
            if len(numeric_cells) >= 2:
                edge_tolerance = max(6, typical_height)
                edges = []
                for cell in sorted(numeric_cells, key=lambda c: c['bbox'][2]):
                    x = cell['bbox'][2]
                    if not any(abs(x - e) <= edge_tolerance for e in edges):
                        edges.append(x)
                edges.sort()
                # 上限 8 列：再多说明这一页不是一张表，而是并排的多个小表
                if 2 <= len(edges) <= 8:
                    aligned, boxes, valid = [], [], True
                    for row in run:
                        label_parts = []
                        out = [''] * (len(edges) + 1)
                        cell_boxes = [None] * (len(edges) + 1)
                        for cell in row:
                            text = cell['text'].strip()
                            if _NUMERIC.fullmatch(text):
                                col = min(range(len(edges)),
                                          key=lambda k: abs(cell['bbox'][2] - edges[k])) + 1
                                if out[col]:
                                    valid = False
                                    break
                                out[col] = cell['text']
                                cell_boxes[col] = cell['bbox']
                            else:
                                label_parts.append(cell)
                        if not valid:
                            break
                        out[0] = ' '.join(part['text'].strip() for part in label_parts).strip()
                        if label_parts:
                            cell_boxes[0] = _bbox(label_parts)
                        aligned.append(out)
                        boxes.append(cell_boxes)
                    if valid and any(row[0] for row in aligned):
                        layouts.append((len(edges) + 1, aligned, boxes))
        if not layouts: continue
        _,aligned,boxes = min(layouts,key=lambda layout:layout[0])
        # 置信度要区分"文字坐标推断"与"OCR 坐标推断"这两个完全不同可靠性的来源：
        # spatial-text 用的是 PDF 原生文字的精确坐标（列边界只是从 x 位置推的），
        # 而 OCR 的每个字本身就可能认错。原来两者共用 0.72/0.62 的上限，
        # 结果原生文字表也被打上低置信度、事实被 exclude 出分析 ——
        # 实测一份年报 57 张 spatial-text 表（占 130 张的 44%）几乎没贡献可用数据。
        # 两者仍然都带 review_required=True，界面会照常提示复核。
        ceiling = .92 if method == 'spatial-text' else .62
        confidence = min(ceiling, statistics.mean(c.get('confidence',1) for row in run for c in row))
        box = _bbox([c for row in run for c in row])
        _table(doc, page, aligned, method, bbox=box, confidence=round(confidence,3),
               cell_bboxes=boxes,
               review_required=True)
        warning = '表格按文字坐标推断，列边界、合并单元格及数字符号需对照原页复核'
        if warning not in page['warnings']: page['warnings'].append(warning)


def _ocr_page(doc: dict, pnum: int, image: Path, width: float, height: float, scale_x=1.0, scale_y=1.0) -> dict:
    items = _ocr(image,scale_x,scale_y)
    ordered_rows = _spatial_rows(items)
    text = '\n'.join('  '.join(item['text'] for item in row) for row in ordered_rows)
    warnings = ['OCR 识别文本及财务数字须对照原页复核；识别置信度不是财务数据准确率']
    if not text.strip(): warnings.append('OCR 未识别出文字；页面可能为空或图像质量不足')
    page = _page(doc,pnum,text,items,'ocr',warnings,width=width,height=height,coordinate_system='top-left')
    tokens = [dict(token,evidence_id=b['evidence_id']) for b in page['blocks'] for token in b.get('tokens', [b])]
    _spatial_tables(doc,page,tokens,'spatial-ocr')
    return page


def _parse_image(doc: dict, path: Path) -> None:
    from PIL import Image, ImageOps
    with Image.open(path) as image:
        frames = getattr(image, 'n_frames', 1)
        if frames > MAX_PAGES: raise ValueError('图像帧数超过资源上限；未进行截断解析')
        with tempfile.TemporaryDirectory(prefix='finocr-') as directory:
            for i in range(frames):
                image.seek(i)
                if image.width*image.height > MAX_IMAGE_PIXELS:
                    raise ValueError('图像像素数量超过资源上限；未进行截断解析')
                frame = ImageOps.exif_transpose(image.copy()).convert('RGB')
                target = Path(directory)/f'{i+1}.png'; frame.save(target)
                _ocr_page(doc,i+1,target,frame.width,frame.height)


def _parse_pdf(doc: dict, path: Path) -> None:
    import pymupdf as fitz
    import pdfplumber
    with fitz.open(path) as pdf:
        if pdf.needs_pass: raise ValueError('PDF 已加密，需要先提供可读取的解密副本')
        if len(pdf) > MAX_PAGES: raise ValueError('PDF 页数超过资源上限；未进行截断解析')
        with tempfile.TemporaryDirectory(prefix='finpdf-') as directory, pdfplumber.open(path) as plumber:
            for index, source in enumerate(pdf):
                native = source.get_text('text',sort=True)
                images = source.get_image_info()
                image_area = sum(fitz.Rect(im['bbox']).get_area() for im in images)
                coverage = min(1,image_area/max(1,source.rect.get_area()))
                # Small digital stamps on full-page scans do not suppress OCR.
                use_ocr = (bool(images) and not native.strip()) or (coverage >= .45 and len(re.sub(r'\s','',native)) < 1000) or (not native.strip() and bool(source.get_drawings()))
                ocr_page = None; ocr_note = ''
                if use_ocr:
                    scale = min(3.0, math.sqrt(MAX_IMAGE_PIXELS/max(1,source.rect.get_area())))
                    pixmap = source.get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False)
                    target = Path(directory)/f'{index+1}.png'; pixmap.save(target)
                    try:
                        ocr_page = _ocr_page(doc,index+1,target,source.rect.width,source.rect.height,source.rect.width/pixmap.width,source.rect.height/pixmap.height)
                    except RuntimeError as error:
                        # OCR 不可用（非 macOS / 二进制缺失）时不能让整份年报失败：
                        # 这一页退回原生文本层并留下明确告警（A 股年报常见整页插图/印章页）。
                        ocr_note = f'本页需要 OCR 但本机 OCR 不可用，已退回文本层（可能漏数）：{error}'
                if ocr_page is not None:
                    page = ocr_page
                    if native.strip(): page['native_text'] = native
                else:
                    items = []
                    for block in source.get_text('dict',sort=True)['blocks']:
                        if block.get('type') != 0: continue
                        for line in block['lines']:
                            text = ''.join(span['text'] for span in line['spans'])
                            items.append(dict(text=text,bbox=[round(v,3) for v in line['bbox']],confidence=1.0))
                    page = _page(doc,index+1,native,items,width=source.rect.width,height=source.rect.height,coordinate_system='top-left')
                    table_count = len(doc['tables'])
                    try:
                        for found in plumber.pages[index].find_tables():
                            rows = [[_value(c) for c in row] for row in found.extract()]
                            # 注意：**单列表格必须保留**，不能在这里丢掉。
                            # 无框线页面上 pdfplumber 经常只框住数值列，给出"每行一个单元格"
                            # 的退化表；看起来没用，但它保留了正确的**行分组**，
                            # 而 ocr_table_rebuild.upgrade_tables 正是靠这份行分组，
                            # 配合页面文本块的 x 坐标把「科目名 + 附註 + 两年数值」重建回来
                            # （method 变成 ocr_row_rebuild）。实测曾在这里过滤掉单列表格，
                            # 结果主表精确值被五年摘要页的四舍五入值顶替（收入 52,552,910
                            # 变成 52,553,000 千元），已核验事实反而下降 —— 已回退。
                            if rows and len(rows)>1:
                                _table(doc,page,rows,'pdfplumber',bbox=list(found.bbox),confidence=.95,cell_bboxes=[[list(cell) if cell else None for cell in row.cells] for row in found.rows])
                    except Exception as error:
                        page['warnings'].append(f'有框表格解析失败，尝试坐标推断：{type(error).__name__}')
                    if len(doc['tables']) == table_count:
                        words = [dict(text=w[4],bbox=list(w[:4]),confidence=1.0) for w in source.get_text('words',sort=True)]
                        _spatial_tables(doc,page,words,'spatial-text')
                    if ocr_note: page['warnings'].append(ocr_note)


def parse_document(path: Path | str, cache_dir: Path | str | None = None) -> dict:
    """Extract all supported content or raise an explicit error; never truncate."""
    path = Path(path)
    if not path.is_file(): raise ValueError('文件不存在或不是普通文件')
    size = path.stat().st_size
    if size > MAX_FILE_BYTES: raise ValueError(f'文件大小 {size} 超过资源上限 {MAX_FILE_BYTES}；未进行截断解析')
    ext = path.suffix.lower()
    supported = TEXT_EXTS | IMAGE_EXTS | {'.pdf','.csv','.tsv','.xlsx','.xlsm','.docx'}
    if ext not in supported: raise ValueError(f'结构化解析不支持 {ext or "无扩展名"}，请转换为 PDF、xlsx、docx 或 UTF-8 文本')
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024*1024), b''): digest.update(chunk)
    sha = digest.hexdigest()
    cache = Path(cache_dir)/f'{sha}-{PARSER_VERSION}-{ext[1:]}.json' if cache_dir else None
    if cache and cache.is_file():
        try:
            result = json.loads(cache.read_text(encoding='utf-8'))
            if result['sha256'] == sha and result['parser_version'] == PARSER_VERSION:
                result['filename'] = path.name
                return result
        except (ValueError,KeyError,OSError): pass
    doc = dict(document_id=sha,filename=path.name,sha256=sha,parser_version=PARSER_VERSION,
               page_count=0,pages=[],tables=[],text='',warnings=[])
    if ext == '.pdf': _parse_pdf(doc,path)
    elif ext in IMAGE_EXTS: _parse_image(doc,path)
    elif ext in {'.xlsx','.xlsm'}: _parse_excel(doc,path)
    elif ext in {'.csv','.tsv'}: _parse_delimited(doc,path)
    elif ext == '.docx': _parse_word(doc,path)
    else:
        text = _decode(path)
        if ext == '.json': json.loads(text) # Preserve original JSON formatting while validating.
        _page(doc,1,text,_line_items(text))
    doc['page_count'] = len(doc['pages'])
    doc['text'] = '\n\n'.join(page['text'] for page in doc['pages'])
    doc['warnings'] = list(dict.fromkeys(doc['warnings'] + [w for p in doc['pages'] for w in p['warnings']]))
    if cache:
        try:
            cache.parent.mkdir(parents=True,exist_ok=True)
            with tempfile.NamedTemporaryFile('w',encoding='utf-8',dir=cache.parent,delete=False) as stream:
                temporary = Path(stream.name)
                json.dump(doc,stream,ensure_ascii=False,allow_nan=False)
            os.replace(temporary,cache)
        except OSError as error:
            doc['warnings'].append(f'缓存写入失败，完整解析结果仍然可用：{type(error).__name__}')
    return doc


def read_document(document: dict, page: int | None = None, table_id: str | None = None,
                  offset: int = 0, limit: int = 12000, query: str = '') -> dict:
    """Read a bounded character window and attach exact page/table evidence.

    Query matches return contextual blocks. table_id returns intersecting complete
    rows and their row numbers; use next_offset to retrieve subsequent windows.
    """
    if not isinstance(offset,int) or offset < 0: raise ValueError('offset 必须是非负整数')
    if not isinstance(limit,int) or limit < 1 or limit > 100000: raise ValueError('limit 必须介于 1 和 100000')
    pages = document['pages']
    if page is not None:
        pages = [p for p in pages if p['page'] == page]
        if not pages: raise ValueError('页码不存在')
    tables = [t for t in document['tables'] if t['page'] in {p['page'] for p in pages}]
    selected_table = None
    if table_id:
        selected_table = next((t for t in tables if t['table_id'] == table_id),None)
        if selected_table is None: raise ValueError('表格 ID 不存在或不属于指定页')
        pages = [p for p in pages if p['page'] == selected_table['page']]
        text = _rows_text(selected_table['rows'])
    elif query:
        matches = [b for p in pages for b in p['blocks'] if query.casefold() in b['text'].casefold()]
        text = '\n\n'.join(b['text'] for b in matches)
    else:
        text = '\n\n'.join(p['text'] for p in pages)
    chunk = text[offset:offset+limit]; end = min(len(text),offset+limit)
    evidence = []
    if selected_table:
        selected_ids = set(); position = 0
        for r,row in enumerate(selected_table['rows']):
            length = len('\t'.join(row))
            if position+length > offset and position < end:
                selected_ids.update(selected_table.get('cell_evidence_ids',[[]]*len(selected_table['rows']))[r])
            position += length+1
        evidence = [dict(b,page=p['page']) for p in pages for b in p['blocks'] if b['evidence_id'] in selected_ids]
    elif query:
        position = 0
        for p in pages:
            for b in p['blocks']:
                if query.casefold() in b['text'].casefold():
                    length = len(b['text'])
                    if position+length > offset and position < end: evidence.append(dict(b,page=p['page']))
                    position += length+2
    else:
        position = 0
        for p in pages:
            if position+len(p['text']) > offset and position < end:
                for b in p['blocks']:
                    if 'start' not in b or (position+b['end'] > offset and position+b['start'] < end):
                        evidence.append(dict(b,page=p['page']))
            position += len(p['text'])+2
    output_tables = []
    for table in ([selected_table] if selected_table else tables):
        if table is None: continue
        if not selected_table and not any(e['evidence_id'] in table['evidence_ids'] for e in evidence): continue
        summary = {k:v for k,v in table.items() if k not in {'rows','cell_coordinates','cell_bboxes','cell_evidence_ids','number_formats','raw_values','context','formulas','evidence_ids'}}
        if selected_table:
            selected_rows = []; numbers = []; position = 0
            for i,row in enumerate(table['rows']):
                length = len('\t'.join(row))
                if position+length > offset and position < end:
                    selected_rows.append(row); numbers.append(i+1)
                position += length+1
            summary.update(rows=selected_rows,row_numbers=numbers,total_rows=len(table['rows']))
            for key in ('cell_coordinates','cell_bboxes','cell_evidence_ids','number_formats','raw_values'):
                if key in table: summary[key] = [table[key][i-1] for i in numbers]
            if 'formulas' in table:
                addresses = {c for row in summary.get('cell_coordinates',[]) for c in row}
                summary['formulas'] = {c:v for c,v in table['formulas'].items() if c in addresses}
        else: summary['total_rows'] = len(table['rows'])
        output_tables.append(summary)
    # Preserve all evidence IDs but keep selected evidence text bounded for tools.
    budget = limit
    bounded = []
    for entry in evidence:
        item = dict(entry)
        value = item['text']
        item['text'] = value[:budget]
        item['text_complete'] = len(value) <= budget
        budget = max(0,budget-len(item['text']))
        item.pop('tokens',None)
        bounded.append(item)
    return dict(document_id=document['document_id'],filename=document['filename'],text=chunk,evidence=bounded,
                tables=output_tables,offset=offset,next_offset=end,total_characters=len(text),has_more=end<len(text),
                matched_blocks=len(evidence),warnings=document.get('warnings',[]))
