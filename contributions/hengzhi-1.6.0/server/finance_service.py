"""Task-scoped financial tools, immutable run artifacts and evidence-grounded exports."""
from __future__ import annotations

import hashlib
import json
import os
import re
import time
import uuid
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import quote

from . import config

VERSION = 'financial-tools-v1'
SUPPORTED = {'.pdf', '.xlsx', '.xlsm', '.csv', '.tsv', '.docx', '.txt', '.md', '.json', '.png', '.jpg', '.jpeg', '.webp'}


def _id(value: str) -> str:
    if not re.fullmatch(r'[A-Za-z0-9_-]{1,100}', str(value)):
        raise ValueError('无效的任务或运行标识')
    return str(value)


def resolve_attachment(task_id: str, filename: str) -> Path:
    folder = (config.UPLOAD_DIR / _id(task_id)).resolve()
    if not filename or Path(filename).name != filename or '/' in filename or '\\' in filename:
        raise ValueError('只允许读取当前任务已上传的文件名')
    path = (folder / filename).resolve()
    if path.parent != folder:
        raise ValueError('文件不在当前任务附件目录内')
    if not path.is_file():
        raise FileNotFoundError(f'未找到附件：{filename}')
    return path


def run_dir(task_id: str, run_id: str) -> Path:
    return config.DATA_DIR / 'financial' / 'runs' / _id(task_id) / _id(run_id)


def _write_json(path: Path, value: dict | list):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix('.tmp')
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    temp.replace(path)


def get_document(task_id: str, filename: str) -> dict:
    from .financial_documents import parse_document
    path = resolve_attachment(task_id, filename)
    doc = parse_document(path, cache_dir=config.DATA_DIR / 'financial' / 'documents')
    # Keep filename scoped to this task even when a content-identical document was cached.
    doc = dict(doc, filename=filename)
    registry = config.DATA_DIR / 'financial' / 'tasks' / _id(task_id) / f"{_id(doc['document_id'])}.json"
    _write_json(registry, {'filename': filename, 'sha256': doc['sha256']})
    return doc


def evidence_index(doc: dict) -> dict[str, dict]:
    out = {}
    for page in doc.get('pages', []):
        for block in page.get('blocks', []):
            out[block['evidence_id']] = {**block, 'page': page['page']}
    for table in doc.get('tables', []):
        out[table['table_id']] = {'evidence_id': table['table_id'], 'page': table.get('page', 1),
            'text': '\n'.join('\t'.join(str(c) for c in row) for row in table['rows']),
            'bbox': table.get('bbox'), 'sheet': table.get('sheet'), 'table_id': table['table_id']}
    return out


def find_evidence(task_id: str, document_id: str, evidence_id: str) -> tuple[dict, dict, Path]:
    registry = config.DATA_DIR / 'financial' / 'tasks' / _id(task_id) / f'{_id(document_id)}.json'
    if not registry.exists():
        raise FileNotFoundError('该证据不属于当前任务')
    entry = json.loads(registry.read_text(encoding='utf-8'))
    doc = get_document(task_id, entry['filename'])
    if doc['sha256'] != entry['sha256']:
        raise ValueError('源文件已变化，请重新分析')
    evidence = evidence_index(doc).get(evidence_id)
    if not evidence:
        raise FileNotFoundError('未找到证据位置')
    return doc, evidence, resolve_attachment(task_id, entry['filename'])


def inspect_document(task_id: str, args: dict) -> dict:
    from .financial_documents import read_document
    from .notes_index import build_notes_index, disclosure_hint
    doc = get_document(task_id, str(args.get('filename', '')))
    result = read_document(doc, page=args.get('page'), table_id=args.get('table_id'),
        offset=max(0, int(args.get('offset') or 0)), limit=min(16000, max(100, int(args.get('limit') or 10000))),
        query=str(args.get('query') or ''))
    index = build_notes_index(doc.get('pages'), doc.get('text'))
    result.update(document_id=doc['document_id'], filename=doc['filename'], warnings=doc.get('warnings', []),
        table_catalog=[{k: t.get(k) for k in ('table_id', 'page', 'sheet', 'confidence', 'method')} |
            {'rows': len(t['rows']), 'columns': max(map(len, t['rows']), default=0)} for t in doc.get('tables', [])],
        # 附注索引与关键披露定位：让"数据缺口"能变成"见附注 X（第 N 页）"。
        notes_index=index.get('notes', []),
        notes_section_start=index.get('notes_section_start'),
        disclosure_map=disclosure_hint(index),
        notes_index_note=index.get('note', ''),
        evidence_link_template=f"/api/financial/evidence/{task_id}/{doc['document_id']}?evidence_id=证据ID")
    return result


def _dedupe_reviews(reviews: list) -> list:
    """待复核清单去重。

    同一个冲突会在两个阶段各产生一次：`extract_financial_facts` 在抽取阶段跑一次
    `_conflicts`，`analyze_financials` 对**全部**事实再跑一次，两边结果在
    analyze_financial_report 里被合并（`review + analysis['review_queue']`）。
    于是同一条冲突在报告里出现两遍，用户看到的"待复核 N 条"直接翻倍 ——
    实测农夫山泉 6 条真冲突被报成 12 条。这里按 (代码, 事实ID集合, 说明) 去重。
    """
    seen, unique = set(), []
    for review in reviews or []:
        if not isinstance(review, dict):
            unique.append(review)
            continue
        key = (review.get('code'), tuple(sorted(review.get('fact_ids') or [])),
               review.get('message'))
        if key in seen:
            continue
        seen.add(key)
        unique.append(review)
    return unique


def save_run(task_id: str, payload: dict) -> dict:
    run = {**payload, 'run_id': uuid.uuid4().hex, 'task_id': _id(task_id),
           'created_at': time.time(), 'tool_version': VERSION}
    if isinstance(run.get('review_queue'), list):
        run['review_queue'] = _dedupe_reviews(run['review_queue'])
    for fact in run.get('facts', []):
        fact['can_confirm'] = _can_confirm(fact, run.get('review_queue', []))
    # 可复现凭证：记录代码指纹、Prompt 哈希、输入哈希与输出哈希。
    # 竞赛要求每次运行生成 repro.json，并能在答辩时复现同一 run_id。
    try:
        from . import repro
        run['repro'] = repro.build(run)
    except Exception as error:                            # pragma: no cover
        run['repro'] = {'error': f'复现凭证生成失败：{error}'}
    _write_json(run_dir(task_id, run['run_id']) / 'run.json', run)
    return run


def _can_confirm(fact: dict, reviews: list) -> bool:
    from datetime import date
    if fact.get('status') != 'needs_review' or fact.get('unit') in {None, '', 'unknown', '%'}:
        return False
    if fact.get('scope') not in {'consolidated', 'parent'} or fact.get('currency') in {None, '', 'unknown'}:
        return False
    if fact.get('period_type') not in {'instant', 'annual', 'ytd', 'quarter'} or not fact.get('evidence_ids'):
        return False
    try:
        date.fromisoformat(fact.get('period', ''))
        from decimal import Decimal
        if not Decimal(fact.get('value', '')).is_finite():
            return False
    except (TypeError, ValueError, ArithmeticError):
        return False
    relevant = [r for r in reviews if isinstance(r, dict) and fact.get('fact_id') in r.get('fact_ids', [])]
    for entry in relevant:
        if entry.get('code') == 'excluded_fact':
            continue
        if entry.get('code') != 'unresolved_metadata' or not re.search(r'待确认：ocr_confidence[。.]?$', entry.get('message', '')):
            return False
    return bool(relevant)


def confirm_financial_facts(task_id: str, args: dict) -> dict:
    """Human UI only: confirm source reading, preserve original run and values."""
    from .financial_analysis import analyze_financials
    original = load_run(task_id, str(args.get('run_id') or ''))
    selected = set(args.get('fact_ids') or [])
    if not selected:
        raise ValueError('请选择已对照原文确认的财务事实')
    facts = json.loads(json.dumps(original.get('facts', [])))
    matched = [f for f in facts if f.get('fact_id') in selected]
    if len(matched) != len(selected) or any(not _can_confirm(f, original.get('review_queue', [])) for f in matched):
        raise ValueError('仅能确认元数据齐全、因OCR或坐标解析置信度不足而待复核的取数；单位、期间或来源冲突需先解决')
    confirmations = list(original.get('human_confirmations', []))
    for fact in matched:
        fact['status'] = 'verified'
        fact['human_confirmed'] = True
        confirmations.append({'fact_id': fact['fact_id'], 'confirmed_at': time.time(), 'value': fact['value'],
                              'source_run_id': original['run_id'], 'evidence_ids': fact['evidence_ids']})
    analysis = analyze_financials(facts)
    retained_reviews = [r for r in original.get('review_queue', []) if not selected.intersection(r.get('fact_ids', []))]
    analysis['review_queue'] = retained_reviews + analysis.get('review_queue', [])
    analysis['warnings'] = original.get('warnings', []) + analysis.get('warnings', [])
    run = save_run(task_id, {**original, **analysis, 'kind': 'analysis', 'parent_run_id': original['run_id'],
                            'human_confirmations': confirmations})
    result = _summary(run)
    result['exports'] = [export_financial_report(task_id, {'run_id':run['run_id'], 'format':fmt}) for fmt in ('xlsx','md','json')]
    return result


def load_run(task_id: str, run_id: str) -> dict:
    return json.loads((run_dir(task_id, run_id) / 'run.json').read_text(encoding='utf-8'))


def list_runs(task_id: str) -> list[dict]:
    folder = config.DATA_DIR / 'financial' / 'runs' / _id(task_id)
    output = []
    if folder.exists():
        for file in folder.glob('*/run.json'):
            try:
                run = json.loads(file.read_text(encoding='utf-8'))
                output.append({k: run.get(k) for k in ('run_id', 'kind', 'created_at', 'filenames', 'company', 'currency', 'parent_run_id')})
            except (OSError, ValueError):
                continue
    return sorted(output, key=lambda r: r['created_at'] or 0, reverse=True)


def read_financial_run(task_id: str, args: dict) -> dict:
    run = load_run(task_id, str(args.get('run_id') or ''))
    section = str(args.get('section') or 'facts')
    if section not in {'facts', 'metrics', 'checks', 'review_queue', 'records', 'field_results',
                       'signals', 'findings', 'tieout_diagnosis', 'scenarios', 'experiments',
                       'issues', 'competition_metrics', 'claims', 'quality_dimensions',
                       'restatement_changes', 'reason_candidates', 'disclosure_map',
                       'quarter_facts', 'restatement_suspects'}:
        raise ValueError('不支持的结果部分')
    items = ((run.get('claim_review') or {}).get('claims', []) if section == 'claims' else
             (run.get('quality_review') or {}).get('dimensions', []) if section == 'quality_dimensions' else
             run.get(section, []))
    if section == 'scenarios' and isinstance(items, dict):
        items = [{**({'name': k} if isinstance(v, dict) else {}), **(v if isinstance(v, dict) else {'value': v})}
                 for k, v in items.items()]
    for key in ('code', 'period'):
        if args.get(key):
            items = [item for item in items if isinstance(item, dict) and item.get(key) == args[key]]
    offset = max(0, int(args.get('offset') or 0))
    limit = min(50, max(1, int(args.get('limit') or 15)))
    selected = items[offset:offset+limit]
    return {'ok': True, 'run_id': run['run_id'], 'section': section, 'items': selected, 'total': len(items), 'offset': offset,
            'next_offset': offset+len(selected), 'has_more': offset+len(selected) < len(items)}


def _decorate_fact(task_id: str, fact: dict) -> dict:
    source = fact.get('source') or {}
    docid = source.get('document_id', '')
    ids = fact.get('evidence_ids') or []
    out = dict(fact)
    if docid and ids:
        amount_id = None
        if source.get('row') is not None and source.get('column') is not None and source.get('table_id'):
            target = f"{source['table_id']}:r{source['row']+1}:c{source['column']+1}"
            if target in ids:
                amount_id = target
        out['evidence_url'] = f'/api/financial/evidence/{task_id}/{docid}?evidence_id={quote(amount_id or ids[0], safe="")}'
    return out


def extract_financial_data(task_id: str, args: dict) -> dict:
    from .financial_analysis import extract_financial_facts
    filename = str(args.get('filename') or '')
    doc = get_document(task_id, filename)
    extracted = extract_financial_facts(doc, company=str(args.get('company') or ''),
        period=str(args.get('period') or ''), scope=str(args.get('scope') or 'auto'))
    extracted['facts'] = [_decorate_fact(task_id, f) for f in extracted.get('facts', [])]
    run = save_run(task_id, {**extracted, 'kind': 'extraction', 'filenames': [filename],
        'documents': [{k: doc.get(k) for k in ('document_id', 'filename', 'sha256', 'parser_version')}],
        'parameters': {k: args.get(k) for k in ('company', 'period', 'scope')}})
    return _summary(run)


def analyze_financial_report(task_id: str, args: dict) -> dict:
    from .financial_analysis import extract_financial_facts, analyze_financials
    filenames = args.get('filenames') or []
    if not filenames:
        folder = config.UPLOAD_DIR / _id(task_id)
        filenames = [f.name for f in sorted(folder.glob('*')) if f.suffix.lower() in SUPPORTED and f.is_file()]
    if not filenames:
        raise ValueError('请先上传财报 PDF、Excel、Word 或图片')
    facts, warnings, review, docs = [], [], [], []
    disclosure = []
    full_docs = []
    for filename in filenames:
        doc = get_document(task_id, str(filename))
        full_docs.append(doc)
        extraction = extract_financial_facts(doc, company=str(args.get('company') or ''),
            period=str(args.get('period') or ''), scope=str(args.get('scope') or 'auto'))
        facts.extend(extraction.get('facts', []))
        warnings.extend(doc.get('warnings', []) + extraction.get('warnings', []))
        review.extend(extraction.get('review_queue', []))
        docs.append({k: doc.get(k) for k in ('document_id', 'filename', 'sha256', 'parser_version')})
        # 附注索引：把"工具未抽取"与"文档确未披露"区分开。
        # 报告里的"数据缺口"若改成"见附注 19（第 197 页）"，性质完全不同 ——
        # 前者是取数范围问题，后者才是披露缺失。
        try:
            from .notes_index import build_notes_index, disclosure_hint
            index = build_notes_index(doc.get('pages'), doc.get('text'))
            for line in disclosure_hint(index):
                disclosure.append({**line, 'file': doc.get('filename', '')})
            if index.get('notes'):
                warnings.append(
                    f"{doc.get('filename', '')}：识别到 {len(index['notes'])} 条财务报表附注"
                    f"（附注章节起始于工具页码 {index.get('notes_section_start')}）。"
                    '引用附注页码时请注明是工具页码，并与印刷页脚核对。')
        except Exception as error:                        # pragma: no cover
            warnings.append(f'附注索引未能建立（{error}）；涉及附注的取数请人工定位。')
    facts = [_decorate_fact(task_id, f) for f in facts]
    analysis = analyze_financials(facts)
    analysis['warnings'] = warnings + analysis.get('warnings', [])
    analysis['review_queue'] = review + analysis.get('review_queue', [])
    analysis['disclosure_map'] = disclosure
    # 交接文档第 5 项的确定性补齐：单季/环比 / 会计口径变化 / 业绩变动原因定位。
    # 引擎原生已覆盖同比、可推导单季、现金转化率与非经常性损益汇总占比，这里只补真正缺的。
    # 全部为补充计算，缺数据返回空并给出说明，不用默认值填平。
    try:
        from .financial_gaps import gap_analysis, quarter_facts_from_documents, restatement_suspects
        # 年报「分季度主要财务指标」表：四个单季值 + 用「四季之和 == 年度值」勾稽。
        # 有了它，单看一份年报也能算环比（决赛是评委给样例数据，不能指望凑齐多期文件）。
        quarter_facts, quarter_reviews, quarter_notes = quarter_facts_from_documents(full_docs, facts)
        if quarter_facts:
            analysis['quarter_facts'] = quarter_facts
        gap_metrics, gap_reviews, restatement, reasons, gap_notes = gap_analysis(
            facts, full_docs, extra_facts=quarter_facts)
        # 口径未确定的多值组：单独存、单独说明，不得当口径变化引用。
        analysis['restatement_suspects'] = restatement_suspects(facts)
        analysis['metrics'] = analysis.get('metrics', []) + gap_metrics
        analysis['review_queue'] = analysis['review_queue'] + quarter_reviews + gap_reviews
        analysis['restatement_changes'] = restatement
        analysis['reason_candidates'] = reasons
        analysis['gap_notes'] = list(quarter_notes) + list(gap_notes)
        if quarter_facts:
            analysis.setdefault('warnings', []).append(
                f'已从年报「分季度主要财务指标」表抽取 {len(quarter_facts)} 个单季值'
                '（period_type=quarter，derived=true）；勾稽未通过的行标为 needs_review，不参与环比。')
        if gap_metrics:
            analysis.setdefault('warnings', []).append(
                f'已补充计算 {len(gap_metrics)} 项环比指标（code 以 _qoq 结尾，derived=true）；'
                '这是本项目补充计算，引用时须与引擎原生指标区分。')
        if restatement:
            analysis.setdefault('warnings', []).append(
                f'发现 {len(restatement)} 组同一科目同一期间的多值情况（可能为重述或口径调整）；'
                '差异已列出，成因须人工读附注确认，不得据此推断为差错。')
        if analysis['restatement_suspects']:
            analysis.setdefault('warnings', []).append(
                f'另有 {len(analysis["restatement_suspects"])} 组同科目同期间的多值差异因口径未确定'
                '（合并/母公司未知）未计入重述候选，已存于 restatement_suspects，不得当口径变化引用。')
    except Exception as error:                            # 补齐失败不能吞掉主分析结果
        analysis.setdefault('warnings', []).append(f'环比/口径变化/业绩原因补齐未完成：{error}')
    run = save_run(task_id, {**analysis, 'kind': 'analysis', 'filenames': filenames, 'documents': docs,
        'parameters': {k: args.get(k) for k in ('company', 'period', 'scope')}})
    # Generate deterministic deliverables without unrestricted shell/file tools.
    exports = [export_financial_report(task_id, {'run_id': run['run_id'], 'format': fmt}) for fmt in ('xlsx', 'md', 'json')]
    result = _summary(run)
    result['exports'] = exports
    return result


def _compact_metric(metric: dict) -> dict:
    """指标的紧凑视图：去掉 evidence_ids / dependency_fact_ids。

    这两个字段占了指标记录的绝大部分体积（一条指标动辄十几个证据 ID），
    而模型判断"数值是多少、怎么算的"只需要 code / period / scope / value / formula。
    证据链通过 complete_result 与 read_financial_run 按需取，不必塞进每次上下文。
    """
    return {key: metric.get(key) for key in
            ('code', 'period', 'period_type', 'scope', 'currency', 'value', 'unit', 'formula', 'reason')
            if metric.get(key) is not None}


def _compact_check(check: dict) -> dict:
    return {key: check.get(key) for key in
            ('code', 'status', 'period', 'scope', 'difference', 'tolerance', 'formula')
            if check.get(key) is not None}


def _compact_review(review: dict) -> dict:
    return {key: review.get(key) for key in ('code', 'message', 'fact_ids', 'alias_of')
            if review.get(key) is not None}


def _facts_digest(facts: list, max_periods: int = 6) -> list:
    """把事实压成「科目 × 口径 → 各期数值」的紧凑视图。

    为什么必须压：`analyze_financial_report` 的完整产出（250 条事实 + 86 项指标 +
    待复核清单）超过 24 万字符，而单个工具结果进模型上下文前会被截到 12000 字符。
    不压的话模型只看得到开头几条事实，于是产出三类自相矛盾的结果：
      · 2024 年的折旧摊销写「—」，而程序其实抽到了 2,727,727 千元；
      · 每股基本盈利填 0，而同一份报告的结论里写 1.411；
      · 把 9 项附注里写得清清楚楚的数据写成"数据缺口"。
    压缩后 250 条事实约 9 千字符，模型能看到全貌，再按需用 read_financial_run 取明细。
    """
    grouped: dict[tuple, dict] = {}
    for fact in facts or []:
        if fact.get('status') != 'verified':
            continue
        key = (fact.get('code'), fact.get('scope'), fact.get('currency'))
        entry = grouped.setdefault(key, {'code': fact.get('code'), 'scope': fact.get('scope'),
                                         'unit': fact.get('unit'), 'tier': fact.get('source_tier'),
                                         'values': {}, 'pages': []})
        period = str(fact.get('period') or '')
        if period not in entry['values']:
            entry['values'][period] = fact.get('raw_value')
        page = (fact.get('source') or {}).get('page')
        if page and page not in entry['pages']:
            entry['pages'].append(page)
    digest = []
    for entry in grouped.values():
        periods = sorted(entry['values'], reverse=True)[:max_periods]
        digest.append({'code': entry['code'], 'scope': entry['scope'], 'unit': entry['unit'],
                       'tier': entry['tier'], 'values': {p: entry['values'][p] for p in periods},
                       'pages': sorted(entry['pages'])[:6]})
    digest.sort(key=lambda item: (str(item['code']), str(item['scope'])))
    return digest


def _summary(run: dict) -> dict:
    """给模型的紧凑结果。

    体积是这里的第一约束：模型看到的单次工具结果上限是 12000 字符，
    超出后 agent 的通用压缩逻辑会把**每个列表都砍到 5–10 条** —— 那正是
    "模型只看到 5 条事实、于是把 2024 年折旧写成 —、把附注可查数据写成数据缺口"
    的由来。这里用「短键 + 数组 + 公式目录」把载荷压到预算内，
    让通用压缩逻辑没有东西可砍；逐条明细仍可通过 complete_result 取。

    字段含义见 legend；事实按 (科目, 口径, 原文单位, 证据等级) 分组，
    避免把"百万元摘要页"与"千元主表"的同一科目混在一个数值序列里。
    """
    facts = run.get('facts', [])
    reviews = run.get('review_queue', [])
    review_counts: dict[str, int] = {}
    for review in reviews:
        code = str(review.get('code') or 'unknown')
        review_counts[code] = review_counts.get(code, 0) + 1
    warnings = list(dict.fromkeys(str(w) for w in run.get('warnings', [])))

    grouped: dict[tuple, dict] = {}
    for fact in facts:
        if fact.get('status') != 'verified':
            continue
        key = (fact.get('code'), fact.get('scope'), fact.get('unit'), fact.get('source_tier'))
        entry = grouped.setdefault(key, {'v': {}, 'p': []})
        period = str(fact.get('period') or '')
        entry['v'].setdefault(period, fact.get('raw_value'))
        page = (fact.get('source') or {}).get('page')
        if page and page not in entry['p']:
            entry['p'].append(page)
    fact_rows = []
    for (code, scope, unit, tier), entry in grouped.items():
        periods = sorted(entry['v'], reverse=True)[:5]
        fact_rows.append([code, scope, unit, tier,
                          {p: entry['v'][p] for p in periods}, sorted(entry['p'])[:5]])
    fact_rows.sort(key=lambda row: (str(row[0]), str(row[1]), str(row[2])))

    # 同一指标代码在不同期间可能因可得科目不同而有不同公式
    # （例：cash_and_deposits 在缺长期银行存款的年度退化为 cash）。
    # 静默取第一条会让公式目录与部分期间的实际算式不符 —— 这正是
    # "可验证/可复现"的漏洞。此处保留全部不同公式：唯一时给字符串，
    # 多版本时按「首次出现期间」给出列表。
    formulas, formula_seen, formula_extra = {}, {}, {}
    metric_rows = []
    for metric in run.get('metrics', []):
        code = metric.get('code')
        fml = metric.get('formula')
        if code and fml:
            seen = formula_seen.setdefault(code, [])
            if fml not in seen:
                seen.append(fml)
                formula_extra.setdefault(code, []).append(
                    {'formula': fml, 'period': metric.get('period'),
                     'scope': metric.get('scope')})
            formulas[code] = fml if len(seen) == 1 else [
                {'formula': e['formula'], 'period': e['period'], 'scope': e['scope']}
                for e in formula_extra[code]]
        metric_rows.append([code, metric.get('period'), metric.get('scope'),
                            metric.get('value'), metric.get('unit')])
    check_rows, failed = [], 0
    for check in run.get('checks', []):
        if check.get('status') != 'pass':
            failed += 1
        check_rows.append([check.get('code'), check.get('status'), check.get('period'),
                           check.get('scope'), check.get('difference'), check.get('tolerance')])

    return {'ok': True, 'run_id': run['run_id'], 'kind': run.get('kind'),
        'fact_count': len(facts), 'verified_count': sum(f.get('status') == 'verified' for f in facts),
        'legend': ('facts: [科目代码, 口径, 原文单位, 证据等级, {期间: 原文数值}, [工具页码]]；'
                   'metrics: [指标代码, 期间, 口径, 数值, 单位]，公式见 formula_catalog；'
                   '若某指标的 formula_catalog 值为列表，表示该指标在不同期间因可得科目不同'
                   '而存在多个算式（列表内按首次出现期间给出），引用时必须对齐该期间的算式；'
                   'checks: [核验代码, 结论, 期间, 口径, 差额, 容差]。'
                   '数值为**原文列报数值**（单位见该行第 3 项），未经跨单位换算；'
                   '每股类指标的 unit 为「元/股」，不是金额，禁止再按千元/百万元缩放。'),
        'facts': fact_rows,
        'metrics': metric_rows,
        'formula_catalog': formulas,
        'checks': check_rows,
        'check_summary': {'total': len(check_rows), 'failed': failed},
        'signals': run.get('signals', []),
        'review_summary': review_counts, 'review_count': len(reviews),
        'review_queue': [_compact_review(r) for r in reviews[:6]],
        'review_queue_note': '此处仅列前 6 条，完整清单见 complete_result 的 review_queue。',
        'warnings': [w[:180] for w in warnings[:10]], 'warning_count': len(warnings),
        'disclosure_map': run.get('disclosure_map', []),
        'disclosure_map_note': ('这是本次文档里已定位到的披露位置。报告中的"数据缺口"'
                                '必须区分「本工具未抽取」与「文档确未披露」：'
                                '此处列出的项目属于前者，应写"见附注 X（第 N 页）"。'),
        'complete_result': f"/api/financial/artifacts/{run['task_id']}/{run['run_id']}/run.json"}


def _normal(text) -> str:
    return re.sub(r'[\s,，]', '', str(text)).casefold()


def _supports_literal(raw, support: str) -> bool:
    needle, haystack = _normal(raw), _normal(support)
    if not needle or needle not in haystack:
        return False
    if re.fullmatch(r'\d{4}-\d{1,2}-\d{1,2}', needle):
        from datetime import date
        try:
            date.fromisoformat(needle)
        except ValueError:
            return False
        return bool(re.search(r'(?<![\d-])' + re.escape(needle) + r'(?![\d-])', haystack))
    # Do not accept 12.5 as evidence for 112.5, or +100 as evidence for -100.
    if re.fullmatch(r'[-+−]?\d+(?:\.\d+)?(?:%|％|万股|股|万元|亿元|元|万|亿)?', needle):
        prefix = r'(?<![\d.\-−(（])' if needle[:1] not in '-−' else r'(?<![\d.])'
        return bool(re.search(prefix + re.escape(needle) + r'(?![\d.])', haystack))
    return True


def validate_structured_output(task_id: str, args: dict) -> dict:
    doc = get_document(task_id, str(args.get('filename') or ''))
    index = evidence_index(doc)
    records = args.get('records')
    required = args.get('required_fields') or []
    if not isinstance(records, list) or not records:
        raise ValueError('records 必须是非空记录列表；每个字段采用 {value, evidence_ids} 格式')
    review, field_results = [], []
    for number, record in enumerate(records, 1):
        if not isinstance(record, dict):
            review.append({'row': number, 'reason': '记录不是对象'})
            continue
        fields = record.get('fields', record)
        if not isinstance(fields, dict) or not fields:
            review.append({'row': number, 'reason': '记录字段必须为非空对象'})
            continue
        for name in required:
            if name not in fields:
                review.append({'row': number, 'field': name, 'reason': '缺少必填字段'})
        for name, field in fields.items():
            issues = []
            if not isinstance(field, dict) or 'value' not in field:
                issues.append('字段必须包含 value 和 evidence_ids')
                field = {'value': field, 'evidence_ids': []}
            value, ids = field.get('value'), field.get('evidence_ids') or []
            if value is None or value == '':
                if name in required:
                    issues.append('必填字段缺失，已留空，不能补猜')
                status = 'missing'
            else:
                if not isinstance(ids, list) or not ids:
                    issues.append('非空字段缺少证据')
                    ids = []
                unknown = [i for i in ids if not isinstance(i, str) or i not in index]
                if unknown:
                    issues.append('引用了不存在的证据ID')
                support = '\n'.join(index[i]['text'] for i in ids if isinstance(i, str) and i in index)
                raw = field.get('raw_value', value)
                if not _supports_literal(raw, support):
                    issues.append('原始字段值未出现在所引证据中；需补精确原文或人工复核')
                if 'raw_value' in field and _normal(value) != _normal(raw):
                    issues.append('归一化值与原始值不同，转换尚未由程序核验')
                status = 'needs_review' if issues else 'grounded'
            result = {'row': number, 'field': name, 'value': value, 'evidence_ids': ids,
                      'status': status, 'issues': issues}
            field_results.append(result)
            review.extend({'row': number, 'field': name, 'reason': reason} for reason in issues)
    run = save_run(task_id, {'kind': 'structured', 'filenames': [doc['filename']], 'records': records,
        'field_results': field_results, 'review_queue': review, 'warnings': [
            '此工具核验字段格式、完整性和原文支持；不自动证明主体归属、因果或所有语义关系。'],
        'documents': [{k: doc.get(k) for k in ('document_id', 'filename', 'sha256', 'parser_version')}]})
    return {'ok': True, 'valid': not review, 'run_id': run['run_id'], 'fields': field_results,
            'review_queue': review, 'warning': run['warnings'][0]}


def _sheets(run: dict) -> dict[str, list[dict]]:
    sheets = {}
    if run.get('records'):
        sheets['结构化数据'] = [{name: (field.get('value') if isinstance(field, dict) else field)
            for name, field in record.get('fields', record).items()} for record in run['records'] if isinstance(record, dict)]
        sheets['字段证据'] = run.get('field_results', [])
    if run.get('facts'):
        sheets['财务事实'] = [{
            '科目': f.get('raw_label'), '标准科目代码': f.get('code'), '公司': f.get('company'),
            '期间': f.get('period'), '期间类型': f.get('period_type'), '口径': f.get('scope'),
            '标准金额': f.get('value'), '标准金额单位': f.get('value_unit') or f.get('currency') or '未知',
            '币种': f.get('currency'), '原始金额': f.get('raw_value'), '原始单位': f.get('unit'),
            '换算倍率': f.get('scale'), '状态': f.get('status'), '启发式置信度': f.get('confidence'),
            '来源': f.get('source'), '证据ID': f.get('evidence_ids'), '原文链接': f.get('evidence_url')
        } for f in run['facts']]
    for key, name in [('metrics', '财务指标'), ('checks', '勾稽核验'),
                      ('signals', '异常信号'), ('findings', '法证发现'),
                      ('tieout_diagnosis', '勾稽反查'), ('review_queue', '待复核')]:
        if run.get(key):
            sheets[name] = [x if isinstance(x, dict) else {'说明': x} for x in run[key]]
    # 消融实验是三层嵌套（实验 → 指标 → 基线与实验组），直接摊平进 Excel
    # 会得到一列套一列的 JSON 字符串，评审根本读不了。这里按
    # "一行一个指标" 铺开：实验、消融变量、基线、实验组、指标、两侧取值、更好的一侧。
    experiments = run.get('experiments') or []
    if experiments:
        ablation_rows = []
        for exp in experiments:
            if not isinstance(exp, dict):
                continue
            baseline_label = (exp.get('baseline') or {}).get('label')
            treatment_label = (exp.get('treatment') or {}).get('label')
            for row in exp.get('metrics') or []:
                ablation_rows.append({
                    '实验': exp.get('title'), '消融变量': exp.get('variable'),
                    '基线': baseline_label, '实验组': treatment_label,
                    '指标': row.get('label'), '基线取值': row.get('baseline'),
                    '实验组取值': row.get('treatment'), '更好的一侧': row.get('better'),
                    '判定': exp.get('verdict'), '结论': exp.get('conclusion'),
                })
        if ablation_rows:
            sheets['消融实验'] = ablation_rows
    # 估值情景是嵌套结构，单独铺平成一行一情景，便于直接进 Excel。
    #
    # 注意取的是**顶层** base/optimistic/stress：save_run 保存的是
    # valuation.from_metrics 的原始返回，三情景在顶层而不是在 'scenarios' 键下。
    # 曾经写成 run['scenarios']，于是估值运行的"估值情景"表**从来没生成过** ——
    # 导出的 xlsx 里只有价值区间，没有三情景对照，而这一点不会报错。
    scenarios = run.get('scenarios') or {key: run.get(key) for key in ('base', 'optimistic', 'stress')}
    rows = []
    for key in ('base', 'optimistic', 'stress'):
        sc = scenarios.get(key)
        if not isinstance(sc, dict):
            continue
        val = sc.get('valuation') or {}
        params = sc.get('parameters') or {}
        multiples = val.get('multiples') or {}
        rows.append({
            '情景': sc.get('label', key),
            '企业价值': val.get('enterprise_value'), '权益价值': val.get('equity_value'),
            '每股价值': val.get('value_per_share'),
            'WACC': val.get('wacc'), '永续增长率': val.get('terminal_growth'),
            '隐含EV/EBITDA': multiples.get('implied_ev_ebitda'),
            '隐含P/E': multiples.get('implied_pe'),
            '终值占比': val.get('terminal_value_share'),
            '收入增速': params.get('revenue_growth'), '毛利率': params.get('gross_margin'),
            '期间费用率': params.get('opex_ratio'), '有效税率': params.get('tax_rate'),
            '资本开支率': params.get('capex_to_revenue'), '净债务': params.get('net_debt'),
            '应收天数': params.get('receivable_days'),
            '存货天数': params.get('inventory_days'),
            '应付天数': params.get('payable_days'),
            '净营运资本占收入': params.get('nwc_to_revenue'),
            '风险传导': sc.get('risk_adjustments'),
            '未匹配风险': sc.get('unmatched_risks'),
        })
    if rows:
        sheets['估值情景'] = rows
    if run.get('range'):
        sheets['价值区间'] = [run['range']]
    # 以下四张表是本轮新增的可验证证据，直接引用 run.json 字段，不做二次计算。
    if run.get('inputs_used'):
        sheets['假设来源'] = [{'驱动变量': key, '来源': item.get('source'),
                              '依据': item.get('detail'), '期间': item.get('period'),
                              '证据ID': item.get('evidence_ids')}
                             for key, item in run['inputs_used'].items()]
    if run.get('data_quality'):
        quality = run['data_quality']
        sheets['数据质量'] = [{'项目': '取自公司已核验数据',
                              '数量': len(quality.get('from_company_data') or []),
                              '明细': quality.get('from_company_data')},
                             {'项目': '沿用默认假设',
                              '数量': len(quality.get('from_default_assumptions') or []),
                              '明细': quality.get('from_default_assumptions')},
                             {'项目': '使用者覆盖',
                              '数量': len(quality.get('overridden_by_user') or []),
                              '明细': quality.get('overridden_by_user')},
                             {'项目': '阻塞性缺失',
                              '数量': len(run.get('missing_inputs') or []),
                              '明细': run.get('missing_inputs')},
                             {'项目': '说明', '数量': '', '明细': quality.get('text')}]
    stress = run.get('stress') or {}
    if stress.get('risk_adjustments') or stress.get('unmatched_risks'):
        sheets['压力传导'] = [
            {'类型': '已传导', '规则': item.get('source'), '严重度': item.get('severity'),
             '参数': item.get('parameter'), '从': item.get('from'), '到': item.get('to'),
             '路径': item.get('path')} for item in (stress.get('risk_adjustments') or [])]
        sheets['压力传导'] += [
            {'类型': '未匹配', '规则': item.get('rule_id'), '严重度': '', '参数': '',
             '从': '', '到': '', '路径': item.get('reason')}
            for item in (stress.get('unmatched_risks') or [])]
    if run.get('stress_diagnostic'):
        sheets['压力有效性'] = [run['stress_diagnostic']]
    if run.get('calibration_checks'):
        sheets['模型自检'] = run['calibration_checks']
    base_valuation = (run.get('base') or {}).get('valuation') or {}
    if base_valuation.get('multiples'):
        sheets['倍数交叉验证'] = [base_valuation['multiples']]
    sensitivity = (run.get('base') or {}).get('sensitivity') or {}
    if sensitivity.get('matrix'):
        wacc_axis = sensitivity.get('wacc_axis') or []
        rows = []
        for row in sensitivity['matrix']:
            entry = {'永续增长率 / WACC': row.get('terminal_growth')}
            for position, cell in enumerate(row.get('cells') or []):
                if position < len(wacc_axis):
                    entry[wacc_axis[position]] = (cell.get('equity_value')
                                                  if cell.get('ok') else '不成立')
            rows.append(entry)
        sheets['敏感性矩阵'] = rows
    sheets['运行说明'] = [{'项目': '版本', '内容': VERSION}, {'项目': '运行ID', '内容': run['run_id']},
        {'项目': '源文件', '内容': run.get('documents', [])},
        {'项目': '说明', '内容': '标准金额已换算为币种基础单位（CNY=元、USD=美元），原始金额和原始单位单列保留。ratio指标为比例小数，0.2=20%。待复核事实不作为已核验结论；启发式置信度不是正确概率。估值输出是按情景分布的合理价值区间，不是目标价。'},
        *[{'项目': '警告', '内容': w} for w in run.get('warnings', [])]]
    return sheets


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _markdown(run: dict) -> str:
    if run.get('kind') == 'valuation':
        return _valuation_markdown(run)
    # 消融实验本身就是一份结构化报告（实验 → 基线/实验组 → 指标 → 判定）。
    # 用通用渲染会把它压成"一行几十列"的表格，逐规则位置表这类内容也会丢掉。
    if run.get('kind') == 'ablation' and run.get('experiments'):
        from . import ablations
        return ablations.render_markdown(run)
    parts = ['# 财务与结构化分析报告', f"运行编号：{run['run_id']}",
             '数字、公式及出处见下表。异常信号仅表示进一步核查线索。']
    for name, rows in _sheets(run).items():
        parts.append(f'\n## {name}')
        columns = list(dict.fromkeys(k for row in rows for k in row))
        if not columns:
            continue
        def mdcell(v):
            return _cell(v).replace('|', '\\|').replace('\n', '<br>')
        parts.append('| ' + ' | '.join(columns) + ' |')
        parts.append('| ' + ' | '.join('---' for _ in columns) + ' |')
        parts.extend('| ' + ' | '.join(mdcell(row.get(c)) for c in columns) + ' |' for row in rows)
    return '\n\n'.join(parts)


# 估值运行导出的 Markdown 就是评审要读的那份交付物。通用渲染是"把顶层键
# 铺成一张表"，对估值运行会得到一行几十列的表格，人根本读不了 ——
# 实测导出的报告只有 1.2 KB、全部内容挤在一行里。所以单独渲染。
_SOURCE_LABEL = {'fact': '原文事实', 'metric': '工具指标', 'derived': '派生计算',
                 'assumption': '默认假设'}

# 报告里必须逐项交代的驱动变量（顺序即阅读顺序）。
_VALUATION_DRIVERS = (
    'revenue_base', 'revenue_growth', 'gross_margin', 'opex_ratio', 'tax_rate',
    'da_to_revenue', 'capex_to_revenue', 'receivable_days', 'inventory_days',
    'payable_days', 'nwc_days_revenue_basis', 'nwc_to_revenue',
    'gross_debt', 'cash_like', 'net_debt', 'minority_interest',
    'non_operating_assets', 'shares', 'ebitda_base', 'parent_net_profit_base',
    'wacc', 'terminal_growth', 'years', 'impairment_to_revenue',
    'near_term_fcff_haircut', 'sustainable_profit_discount',
)


def _num_text(value, digits=2):
    """把长小数渲染成可读文本；非数值原样返回。

    估值里的金额是 Decimal 全精度（例如 297805465449.5512434），
    直接写进报告既读不了也没意义 —— 但**不能改计算层**，只在渲染层格式化。
    """
    if value is None or value == '':
        return '—'
    text = str(value)
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        return text
    if number == number.to_integral_value():
        return f'{number:,.0f}'
    return f'{number:,.{digits}f}'


def _percent_text(value, digits=2):
    number = None
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return '—'
    try:
        return f'{(number * 100).quantize(Decimal(1).scaleb(-digits))}%'
    except InvalidOperation:
        return str(value)


def _md_table(columns, rows) -> list:
    def cell(value):
        return _cell(value).replace('|', '\\|').replace('\n', '<br>')
    lines = ['| ' + ' | '.join(columns) + ' |',
             '| ' + ' | '.join('---' for _ in columns) + ' |']
    lines.extend('| ' + ' | '.join(cell(row.get(c)) for c in columns) + ' |' for row in rows)
    return lines


# 以百分比渲染的驱动变量。刻意用白名单而不是"值在 0–1 之间就当百分比" ——
# 后者会把 years=5 之外的任何小数值都误当比例，且无法解释。
_PERCENT_DRIVERS = {
    'wacc', 'terminal_growth', 'revenue_growth', 'gross_margin', 'opex_ratio',
    'tax_rate', 'da_to_revenue', 'capex_to_revenue', 'nwc_to_revenue',
    'impairment_to_revenue', 'near_term_fcff_haircut', 'sustainable_profit_discount',
}


def _valuation_markdown(run: dict) -> str:
    """把估值运行渲染成一份完整的估值报告。

    章节固定为：结论 → 假设与来源 → 数据质量 → 逐年预测 → 敏感性 →
    压力传导 → 倍数交叉验证 → 自检 → 口径与边界。
    每一节都直接引用 run.json 里的字段，不做二次计算 ——
    报告里出现的每个数字都能在 run.json 里原样找到，这才叫"可复现"。
    """
    base = run.get('base') or {}
    created = run.get('created_at')
    try:
        created = time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(float(created)))
    except (TypeError, ValueError):
        created = str(created or '')
    parts = ['# 估值分析报告（DCF 三情景 + 倍数交叉验证）',
             f"运行编号：{run.get('run_id', '')}　生成时间：{created}",
             f"父分析运行：{run.get('parent_run_id', '')}　法证运行：{run.get('forensic_run_id', '') or '（未使用）'}",
             '',
             '> 本报告输出的是**按情景分布的合理价值区间**，不是目标价，也不构成投资建议。']

    # ---- 一、估值结论 ----
    parts.append('\n## 一、估值结论')
    rows = []
    for name in ('base', 'optimistic', 'stress'):
        scenario = run.get(name) or {}
        valuation = scenario.get('valuation') or {}
        multiples = valuation.get('multiples') or {}
        rows.append({
            '情景': scenario.get('label', name),
            '企业价值(元)': _num_text(valuation.get('enterprise_value')),
            '权益价值(元)': _num_text(valuation.get('equity_value')),
            '每股价值(元)': _num_text(valuation.get('value_per_share')),
            '隐含 EV/EBITDA': _num_text(multiples.get('implied_ev_ebitda')),
            '隐含 P/E': _num_text(multiples.get('implied_pe')),
        })
    parts.extend(_md_table(['情景', '企业价值(元)', '权益价值(元)', '每股价值(元)',
                            '隐含 EV/EBITDA', '隐含 P/E'], rows))
    value_range = run.get('range') or {}
    # 区间文字用渲染层重新拼：run.json 里的 text 是机器可读的全精度值
    # （例如 277154690900.2038204547109466），直接贴进报告读不了。
    if value_range.get('equity_value_low') and value_range.get('equity_value_high'):
        parts.append(f"\n**三情景权益价值区间 "
                     f"{_num_text(value_range['equity_value_low'])} – "
                     f"{_num_text(value_range['equity_value_high'])} 元**")
    elif value_range.get('text'):
        parts.append(f"\n**{value_range['text']}**")
    if value_range.get('value_per_share_low') and value_range.get('value_per_share_high'):
        parts.append(f"\n**三情景每股价值区间 "
                     f"{_num_text(value_range['value_per_share_low'])} – "
                     f"{_num_text(value_range['value_per_share_high'])} 元/股**")
    if value_range.get('disclaimer'):
        parts.append(f"\n{value_range['disclaimer']}")
    # DCF 自身的告警（终值占比过高、capex 符号反常等）必须跟着结论一起出现 ——
    # 它们直接影响结论怎么读，不能只躺在 run.json 里。
    for name in ('base', 'optimistic', 'stress'):
        scenario = run.get(name) or {}
        for warning in ((scenario.get('valuation') or {}).get('warnings') or []):
            parts.append(f"\n> ⚠️ {scenario.get('label', name)}：{warning}")

    # ---- 二、驱动变量与来源 ----
    parts.append('\n## 二、驱动变量与来源')
    parts.append('「来源」列说明这个数是**从公司自己的数据算出来的**还是**沿用的默认假设**。'
                 '凡标为「默认假设」的，都是模型对这家公司的未知，必须在阅读结论时一并考虑。')
    provenance = run.get('inputs_used') or {}
    overrides = ((run.get('parameters') or {}).get('overrides') or {})
    parameters = base.get('parameters') or {}
    keys = [k for k in _VALUATION_DRIVERS if k in parameters]
    keys += [k for k in parameters if k not in keys]
    rows = []
    for key in keys:
        item = provenance.get(key) or {}
        if key in overrides:
            source, detail = '使用者覆盖', '由调用方通过 overrides 显式指定'
        elif item:
            source = _SOURCE_LABEL.get(item.get('source', ''), item.get('source', ''))
            detail = item.get('detail', '')
            if item.get('period'):
                detail = (detail + f"（期间 {item['period']}）").strip()
        else:
            source, detail = '默认假设', '引擎未能从本次数据校准，沿用行业中性默认值'
        value = parameters.get(key)
        if isinstance(value, str) and value:
            if key in _PERCENT_DRIVERS:
                value = _percent_text(value)
            else:
                value = _num_text(value)
        rows.append({'驱动变量': key, '取值': value, '来源': source, '依据': detail})
    parts.extend(_md_table(['驱动变量', '取值', '来源', '依据'], rows))

    # ---- 三、数据质量 ----
    parts.append('\n## 三、数据质量声明')
    quality = run.get('data_quality') or {}
    if quality.get('text'):
        parts.append(quality['text'])
    rows = [
        {'项目': '取自公司已核验数据', '数量': len(quality.get('from_company_data') or []),
         '明细': '、'.join(quality.get('from_company_data') or [])},
        {'项目': '沿用默认假设', '数量': len(quality.get('from_default_assumptions') or []),
         '明细': '、'.join(quality.get('from_default_assumptions') or [])},
        {'项目': '使用者覆盖', '数量': len(quality.get('overridden_by_user') or []),
         '明细': '、'.join(quality.get('overridden_by_user') or [])},
        {'项目': '阻塞性缺失（模型不成立）', '数量': len(run.get('missing_inputs') or []),
         '明细': '、'.join(run.get('missing_inputs') or []) or '无'},
    ]
    parts.extend(_md_table(['项目', '数量', '明细'], rows))
    degraded = run.get('degraded_inputs') or []
    if degraded:
        parts.append('\n**降级输入**（模型能跑，但结果可信度下降）：')
        parts.extend(_md_table(['参数', '说明', '影响'],
                               [{'参数': d.get('parameter'), '说明': d.get('label'),
                                 '影响': d.get('impact')} for d in degraded]))
    for note in (run.get('calibration_notes') or []):
        parts.append(f'- {note}')
    warnings = run.get('calibration_warnings') or []
    if warnings:
        parts.append('\n**校准告警**：')
        parts.extend(f'- {w}' for w in warnings)

    # ---- 四、逐年预测 ----
    parts.append('\n## 四、逐年预测与 FCFF')
    for name in ('base', 'stress'):
        scenario = run.get(name) or {}
        projection = scenario.get('projection') or []
        if not projection:
            continue
        parts.append(f"\n### {scenario.get('label', name)}")
        rows = [{'年度': row.get('year'), '营业收入(元)': _num_text(row.get('revenue')),
                 'EBIT(元)': _num_text(row.get('ebit')), 'FCFF(元)': _num_text(row.get('fcff'))}
                for row in projection]
        parts.extend(_md_table(['年度', '营业收入(元)', 'EBIT(元)', 'FCFF(元)'], rows))

    # ---- 五、敏感性矩阵 ----
    sensitivity = base.get('sensitivity') or {}
    matrix = sensitivity.get('matrix') or []
    if matrix:
        parts.append('\n## 五、敏感性矩阵（WACC × 永续增长率）')
        parts.append('每个单元格是按对应 WACC 与永续增长率**完整重算**的权益价值（元），'
                     '不是线性插值。中心格应等于基准情景的权益价值，可用来验证矩阵没算错。')
        wacc_axis = sensitivity.get('wacc_axis') or []
        growth_axis = sensitivity.get('growth_axis') or []
        columns = ['永续增长率 / WACC'] + [_percent_text(w) for w in wacc_axis]
        rows = []
        for index, row in enumerate(matrix):
            growth = row.get('terminal_growth', growth_axis[index] if index < len(growth_axis) else '')
            entry = {columns[0]: _percent_text(growth)}
            for position, cell in enumerate(row.get('cells') or []):
                if position + 1 >= len(columns):
                    break
                entry[columns[position + 1]] = (_num_text(cell.get('equity_value'))
                                                if cell.get('ok') else '不成立')
            rows.append(entry)
        parts.extend(_md_table(columns, rows))
        parts.append(f"\n{sensitivity.get('note', '')}")
        terminal_share = (base.get('valuation') or {}).get('terminal_value_share')
        if terminal_share:
            parts.append(f"\n基准情景终值占企业价值 **{_percent_text(terminal_share)}**。"
                         '该比例越高，结论越依赖永续假设。')

    # ---- 六、压力情景传导 ----
    parts.append('\n## 六、压力情景传导（风险 → 估值）')
    stress = run.get('stress') or {}
    adjustments = stress.get('risk_adjustments') or []
    if adjustments:
        def _param_text(name, value):
            # 比率类参数必须按百分比渲染，否则 0.2331 → 0.2346 会一起显示成
            # "0.23 → 0.23"，看上去"没改动"，读者会误判这条规则失效。
            return (_percent_text(value, 4) if name in _PERCENT_DRIVERS
                    else _num_text(value))
        rows = [{'来源规则': a.get('source'), '严重度': a.get('severity'),
                 '参数': a.get('parameter'),
                 '从': _param_text(a.get('parameter'), a.get('from')),
                 '到': _param_text(a.get('parameter'), a.get('to')),
                 '传导路径': a.get('path')} for a in adjustments]
        parts.extend(_md_table(['来源规则', '严重度', '参数', '从', '到', '传导路径'], rows))
    else:
        parts.append('本次压力情景**没有任何参数改动**。')
    unmatched = stress.get('unmatched_risks') or []
    if unmatched:
        parts.append('\n**未匹配的风险（必须作为待补假设处理）**：')
        parts.extend(_md_table(['规则', '原因'],
                               [{'规则': u.get('rule_id'), '原因': u.get('reason')}
                                for u in unmatched]))
    diagnostic = run.get('stress_diagnostic') or {}
    if diagnostic:
        parts.append(f"\n**压力测试有效性**：{diagnostic.get('text', '')}"
                     f"（命中 {diagnostic.get('adjustment_count', 0)} 条可传导规则）")
        if diagnostic.get('insufficient'):
            parts.append(f"\n> ⚠️ {diagnostic.get('reason', '')}")

    # ---- 七、倍数交叉验证 ----
    parts.append('\n## 七、倍数交叉验证')
    multiples = (base.get('valuation') or {}).get('multiples') or {}
    if multiples.get('implied_ev_ebitda') or multiples.get('implied_pe'):
        parts.append('DCF 隐含倍数（自检：落在常识区间之外时应先怀疑输入有误）：')
        parts.extend(_md_table(['隐含 EV/EBITDA', '隐含 P/E'],
                               [{'隐含 EV/EBITDA': _num_text(multiples.get('implied_ev_ebitda')),
                                 '隐含 P/E': _num_text(multiples.get('implied_pe'))}]))
    comparables = multiples.get('comparables') or {}
    if comparables:
        parts.append('\n可比倍数法反推（倍数由使用者提供）：')
        rows = [{'方法': key, '倍数': _num_text(item.get('multiple')),
                 '基数': _num_text(item.get('basis')),
                 '隐含权益价值(元)': _num_text(item.get('equity_value')),
                 '每股价值(元)': _num_text(item.get('value_per_share'))}
                for key, item in comparables.items()]
        parts.extend(_md_table(['方法', '倍数', '基数', '隐含权益价值(元)', '每股价值(元)'], rows))
        low = multiples.get('comparable_range', {}).get('low')
        high = multiples.get('comparable_range', {}).get('high')
        if low and high:
            parts.append(f"\n**可比倍数法隐含权益价值 {_num_text(low)} – {_num_text(high)} 元**")
        elif multiples.get('comparable_range', {}).get('text'):
            parts.append(f"\n{multiples['comparable_range']['text']}")
        gap = multiples.get('dcf_vs_comparable_gap')
        if gap not in (None, ''):
            parts.append(f"\nDCF 结果与可比倍数法最低值的差异为 **{_percent_text(gap)}**；"
                         '两者不一致本身就是一个需要解释的结论（资本开支强度、增长假设、'
                         '或可比组选取），不要只报其中一个。')
    for warning in (multiples.get('warnings') or []):
        parts.append(f'- {warning}')
    if multiples.get('disclaimer'):
        parts.append(f"\n{multiples['disclaimer']}")

    # ---- 八、自检 ----
    checks = run.get('calibration_checks') or []
    if checks:
        parts.append('\n## 八、模型自检')
        rows = [{'检查项': c.get('code'), '结论': c.get('status'),
                 '模型隐含值': _num_text(c.get('implied_ebit')),
                 '报表口径值': _num_text(c.get('reported_ebit')),
                 '差异': _percent_text(c.get('difference_ratio')),
                 '说明': c.get('note')} for c in checks]
        parts.extend(_md_table(['检查项', '结论', '模型隐含值', '报表口径值', '差异', '说明'], rows))

    # ---- 九、口径与边界 ----
    parts.append('\n## 九、口径与边界说明')
    parts.append('- 金额单位为元；比例一律为分数（0.2 = 20%），仅在渲染层做百分比格式化。')
    parts.append('- FCFF = EBIT × (1 − 有效税率) + 折旧摊销 − 资本开支 − ΔNWC；'
                 'EV = Σ FCFF/(1+WACC)^t + TV/(1+WACC)^n；'
                 '权益价值 = EV − 净债务 − 少数股东权益 + 非经营性资产。')
    parts.append('- ΔNWC 已把存货与应付天数从营业成本口径折算到收入口径'
                 '（乘 1 − 毛利率），见驱动变量 `nwc_days_revenue_basis`。')
    parts.append('- WACC 与永续增长率是**行业中性默认值**，不是这家公司的；'
                 '换值必须用 overrides 并写清依据。')
    parts.append('- 压力情景的参数由法证发现自动推导，每条传导都记录来源规则；'
                 '不是手填的悲观数字。')
    parts.append('- 三表不平、单位错位、口径未统一时，应先修数据再估值。')
    return '\n\n'.join(parts)


def export_run(run: dict, fmt: str, output: Path) -> Path:
    if run.get('kind') == 'audit' and run.get('claim_review'):
        if fmt not in {'xlsx', 'docx', 'json', 'md'}:
            raise ValueError('导出格式支持 xlsx、docx、json、md')
        from .report_review_exports import export_review
        names = export_review(output, run)
        return output / names['markdown' if fmt == 'md' else fmt]
    if run.get('kind') == 'valuation_model':
        from .valuation_reports import export_model
        return export_model(run, fmt, output)
    if fmt not in {'xlsx', 'docx', 'json', 'md'}:
        raise ValueError('导出格式支持 xlsx、docx、json、md')
    output.mkdir(parents=True, exist_ok=True)
    path = output / f'报告.{fmt}'
    if fmt == 'json':
        _write_json(path, run)
    elif fmt == 'md':
        path.write_text(_markdown(run), encoding='utf-8')
    elif fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in _sheets(run).items():
            ws = wb.create_sheet(name)
            columns = list(dict.fromkeys(k for row in rows for k in row))
            ws.append(columns or ['说明'])
            for row in rows:
                values = [_cell(row.get(c)) for c in columns]
                # Raw source text must never become a spreadsheet formula.
                ws.append(["'" + v if v[:1] in '=+-@\t\r' else v for v in values])
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='174B64')
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(65, max(14, len(str(column[0].value or '')) + 8))
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
        wb.save(path)
    else:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        doc.sections[0].left_margin = Inches(.65)
        doc.sections[0].right_margin = Inches(.65)
        style = doc.styles['Normal']
        style.font.name = 'PingFang SC'
        style.font.size = Pt(9)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')
        doc.add_heading('财务与结构化分析报告', 0)
        doc.add_paragraph('运行编号：' + run['run_id'])
        doc.add_paragraph('待复核项目不作为已核验结论。异常信号表示核查线索，需要结合原文与业务背景判断。')
        # Compact human report; full facts remain in Excel/JSON.
        for name, rows in _sheets(run).items():
            doc.add_heading(name, level=1)
            if name == '财务事实':
                keys = ['科目', '期间', '标准金额', '币种', '原始金额', '原始单位', '状态', '来源']
            elif name == '财务指标':
                keys = ['code', 'period', 'value', 'unit', 'formula', 'status']
            else:
                keys = list(dict.fromkeys(k for row in rows for k in row))[:6]
            table = doc.add_table(rows=1, cols=max(1, len(keys)))
            table.style = 'Light Shading Accent 1'
            for cell, key in zip(table.rows[0].cells, keys):
                cell.text = key
            for row in rows:
                for cell, key in zip(table.add_row().cells, keys):
                    cell.text = _cell(row.get(key))
        doc.save(path)
    return path


def export_financial_report(task_id: str, args: dict) -> dict:
    run = load_run(task_id, str(args.get('run_id') or ''))
    fmt = str(args.get('format') or 'xlsx')
    path = export_run(run, fmt, run_dir(task_id, run['run_id']))
    return {'ok': True, 'run_id': run['run_id'], 'format': fmt, 'path': str(path),
        'url': f'/api/financial/artifacts/{task_id}/{run["run_id"]}/{quote(path.name)}',
        'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'size': path.stat().st_size}


def _latest_run(task_id: str, kinds: tuple) -> dict | None:
    """取最近一次指定类型的运行，供新工具自动衔接上一步结果。"""
    for entry in list_runs(task_id):
        if entry.get('kind') in kinds:
            try:
                return load_run(task_id, entry['run_id'])
            except (OSError, ValueError):
                continue
    return None


def forensic_scan(task_id: str, args: dict) -> dict:
    """在已核验的财务事实、指标与勾稽结果上运行确定性财务法证规则。"""
    from . import forensics
    run_id = str(args.get('run_id') or '')
    analysis = load_run(task_id, run_id) if run_id else _latest_run(task_id, ('analysis',))
    if not analysis:
        raise ValueError('没有可用的财报分析结果，请先调用 analyze_financial_report 完成提取与核验')
    peer = args.get('peer') if isinstance(args.get('peer'), dict) else {}
    top_k = int(args.get('top_k') or 10)
    result = forensics.scan(analysis.get('facts', []), checks=analysis.get('checks', []),
                            metrics=analysis.get('metrics', []), peer=peer, top_k=top_k)
    run = save_run(task_id, {**result, 'kind': 'forensic',
                             'parent_run_id': analysis['run_id'],
                             'filenames': analysis.get('filenames', []),
                             'parameters': {'top_k': top_k, 'peer_provided': bool(peer)}})
    findings = result.get('findings', [])
    return {
        'ok': True, 'run_id': run['run_id'], 'parent_run_id': analysis['run_id'],
        'finding_count': len(findings), 'findings': findings[:12],
        'findings_preview_only': len(findings) > 12,
        'tieout_diagnosis': result.get('tieout_diagnosis', []),
        'review_queue': result.get('review_queue', []),
        'summary': result.get('summary', {}),
        'complete_result': f"/api/financial/artifacts/{run['task_id']}/{run['run_id']}/run.json",
        'note': ('每条发现都带 rule_id、支持证据、**反证检索方向**、未解问题和估值传导路径。'
                 'risk_score 只用于排序，不等于舞弊概率；severity 是严重度，不是置信度。'
                 '完整清单可用 read_financial_run(section=findings) 分页读取。'),
    }


def scenario_valuation(task_id: str, args: dict) -> dict:
    """基准 / 乐观 / 压力三情景预测与 DCF 估值压力测试。"""
    from . import valuation
    run_id = str(args.get('run_id') or '')
    analysis = load_run(task_id, run_id) if run_id else _latest_run(task_id, ('analysis',))
    if not analysis:
        raise ValueError('没有可用的财报分析结果，请先调用 analyze_financial_report')

    findings = []
    forensic = None
    forensic_run_id = str(args.get('forensic_run_id') or '')
    if forensic_run_id:
        forensic = load_run(task_id, forensic_run_id)
    elif args.get('use_findings', True):
        forensic = _latest_run(task_id, ('forensic',))
    if forensic:
        findings = forensic.get('findings', [])

    overrides = args.get('overrides') if isinstance(args.get('overrides'), dict) else {}
    result = valuation.from_metrics(analysis.get('metrics', []), findings=findings,
                                    overrides=overrides, facts=analysis.get('facts', []))
    run = save_run(task_id, {**result, 'kind': 'valuation',
                             'parent_run_id': analysis['run_id'],
                             'forensic_run_id': (forensic or {}).get('run_id', ''),
                             'filenames': analysis.get('filenames', []),
                             'parameters': {'overrides': overrides,
                                            'findings_used': [f.get('rule_id') for f in findings]}})
    scenarios = {k: {'label': result[k].get('label'), 'parameters': result[k].get('parameters'),
                     'valuation': result[k].get('valuation'),
                     'risk_adjustments': result[k].get('risk_adjustments'),
                     'unmatched_risks': result[k].get('unmatched_risks'),
                     'projection': result[k].get('projection'),
                     'sensitivity': result[k].get('sensitivity')}
                 for k in ('base', 'optimistic', 'stress') if isinstance(result.get(k), dict)}
    return {
        'ok': bool(result.get('ok')), 'run_id': run['run_id'],
        'parent_run_id': analysis['run_id'],
        'forensic_run_id': (forensic or {}).get('run_id', ''),
        'findings_used': [f.get('rule_id') for f in findings],
        'range': result.get('range', {}),
        'scenarios': {k: {kk: vv for kk, vv in v.items() if kk != 'sensitivity'} for k, v in scenarios.items()},
        'sensitivity_base': (scenarios.get('base') or {}).get('sensitivity'),
        'risk_adjustments_stress': (scenarios.get('stress') or {}).get('risk_adjustments'),
        'inputs_used': result.get('inputs_used', {}),
        'missing_inputs': result.get('missing_inputs', []),
        # 以下四块是本轮新增的"可验证"证据：报告里的假设表、数据质量声明、
        # 压力测试有效性声明、以及 DCF 的倍数交叉验证，全部直接引用这里，
        # 不要自己复述、也不要自己重算。
        'assumption_sources': result.get('assumption_sources', {}),
        'data_quality': result.get('data_quality', {}),
        'calibration_notes': result.get('calibration_notes', []),
        'calibration_warnings': result.get('calibration_warnings', []),
        'calibration_checks': result.get('calibration_checks', []),
        'degraded_inputs': result.get('degraded_inputs', []),
        'stress_diagnostic': result.get('stress_diagnostic', {}),
        'multiples_base': ((scenarios.get('base') or {}).get('valuation') or {}).get('multiples', {}),
        'note': ('压力情景的参数由法证发现自动推导，不是手填的悲观数字；每条传导都记录来源规则。'
                 '输出是按情景分布的合理价值区间，不是目标价。'
                 'assumption_sources / inputs_used 给出每个驱动变量的来源与依据，'
                 '写假设表时必须逐项引用；data_quality 说明哪些参数取自公司数据、'
                 '哪些沿用默认假设；stress_diagnostic.insufficient 为真时，'
                 '报告中必须写明本次压力测试没有构成实质压力，不得宣称风险已充分量化。'
                 '完整敏感性矩阵与逐年预测可用 read_financial_run(section=scenarios) 读取。'),
        'complete_result': f"/api/financial/artifacts/{run['task_id']}/{run['run_id']}/run.json",
    }


def _runs_roots() -> list[Path]:
    """运行目录可能有多个根，都要搜。

    开发模式下数据目录是项目内的 `data/`，打包应用则是
    `~/Library/Application Support/FinAgent/`（由 FINAGENT_DATA_DIR 指定）。
    报告是哪个环境生成的，就得去哪个环境找它引用的那次运行 ——
    否则拿新数据核对旧报告，会把"数据版本不同"误报成"计数抄错"。
    """
    candidates = [config.DATA_DIR / 'financial' / 'runs']
    override = os.environ.get('FINAGENT_DATA_DIR')
    if override:
        candidates.append(Path(override).expanduser() / 'financial' / 'runs')
    candidates.append(Path.home() / 'Library' / 'Application Support' / 'FinAgent'
                      / 'financial' / 'runs')
    seen, out = set(), []
    for root in candidates:
        key = str(root)
        if key in seen or not root.exists():
            continue
        seen.add(key)
        out.append(root)
    return out


def _locate_runs_by_text(text: str) -> dict:
    """按报告正文里出现的 run_id，在运行目录里找对应的运行。

    报告里引用的运行可能属于**另一次任务**（报告是当时生成的，自检是后来跑的），
    所以不能只在当前 task_id 下找，要按 run_id 全局搜。

    为什么非做这一步：如果拿新跑的分析去核对旧报告，"数据版本不同"会被误报成
    "计数抄错"。实测踩过：报告写 248 条事实，重跑后是 257 条，
    四处计数全报"不符"，而报告本身没错——它引用的就是那一次运行。
    """
    from . import report_audit
    picks: dict[str, dict] = {}
    run_ids = report_audit.extract_run_ids(text or '')
    if not run_ids:
        return picks
    for root in _runs_roots():
        for run_id in run_ids:
            if run_id in {item.get('run_id') for item in picks.values()}:
                continue
            matches = sorted(root.glob(f'*/{run_id}/run.json'))
            if not matches:
                continue
            try:
                run = json.loads(matches[-1].read_text(encoding='utf-8'))
            except (OSError, ValueError):
                continue
            kind = str(run.get('kind') or '')
            if kind == 'valuation_model': kind = 'valuation'
            if kind and kind not in picks:
                picks[kind] = run
    return picks


def audit_report(task_id: str, args: dict) -> dict:
    """本地研报主张核验及覆盖率透明的内部质量评估。"""
    from . import report_audit, report_claims, report_quality, report_review_exports, report_consistency, report_formula_review

    report_path = str(args.get('report_path') or '')
    filename = str(args.get('filename') or '')
    if filename:
        report_path = str(resolve_attachment(task_id, filename))
    text = str(args.get('text') or '')
    if not text and report_path:
        text = report_audit.read_report_text(report_path)
    if not text.strip():
        raise ValueError('请提供报告正文、报告路径或当前任务的报告附件名')
    if len(text) > 2_000_000:
        raise ValueError('报告正文超过200万字符，请拆分后核查')
    located = _locate_runs_by_text(text)
    run = (load_run(task_id, str(args['run_id'])) if args.get('run_id')
           else located.get('analysis'))
    valuation = (load_run(task_id, str(args['valuation_run_id']))
                 if args.get('valuation_run_id') else located.get('valuation'))
    forensic = (load_run(task_id, str(args['forensic_run_id']))
                if args.get('forensic_run_id') else located.get('forensic'))
    if run and run.get('kind') != 'analysis':
        raise ValueError('事实依据必须是财报分析运行，不能用待核查报告自证')
    for linked, kinds in ((valuation, {'valuation', 'valuation_model'}), (forensic, {'forensic'})):
        if linked and linked.get('kind') not in kinds:
            raise ValueError('引用运行类型与用途不符')
        if linked and run and linked.get('parent_run_id') != run.get('run_id'):
            raise ValueError('估值或法证结果与财报分析不是同一批输入，请重新选择依据')
    result = report_audit.audit_report(text, run, valuation, forensic)
    pool = report_audit.build_pool((run or {}).get('facts'), (run or {}).get('metrics'), valuation)
    claim_review = report_claims.review_claims(text, pool, run)
    formula_review = report_formula_review.review_formulas(text)
    conclusion_review = report_formula_review.review_conclusions(text)
    quality_review = report_quality.evaluate_quality(text, claim_review, run, valuation, forensic, formula_review=formula_review)
    result.update(claim_review=claim_review, quality_review=quality_review,
                  review_version='research-review-v1', source_file=filename or report_path,
                  text_length=len(text), report_sha256=hashlib.sha256(text.encode()).hexdigest(),
                  runs_located={kind: item.get('run_id') for kind, item in located.items()})
    # 旧数字池匹配仅作诊断；主张级语义绑定和新覆盖率评分才是本次审阅结论。
    result['grade']['legacy_diagnostic_only'] = True
    result['grade']['rubric_scope'] = '项目旧版内部诊断，不是竞赛官方评分'
    stats = claim_review.get('stats', {})
    result['summary'] = (f"主张核查：已验证{stats.get('verified', 0)}项，"
                         f"证据矛盾{stats.get('contradicted', 0)}项，"
                         f"待核验{stats.get('unverifiable', 0)}项。"
                         '质量评分区分机器核查与人工审阅，未完成部分不计为满分。')
    result['issues'] = [
        {'severity': 'hard' if c.get('status') == 'contradicted' else 'medium',
         'type': '主张与证据矛盾' if c.get('status') == 'contradicted' else '主张待核验',
         'detail': c.get('raw'), 'line': c.get('line'), 'claim_id': c.get('id'),
         'context': c.get('reason'), 'advice': c.get('advice'),
         'evidence_ids': c.get('evidence_ids', [])}
        for c in claim_review.get('claims', []) if c.get('status') != 'verified'
    ] + [i for i in result.get('issues', []) if i.get('type') in {'计数不符', '引用不存在', '一票否决'}]
    consistency = report_consistency.check_consistency(claim_review.get('claims', []), valuation)
    result['formula_review'] = formula_review
    result['conclusion_review'] = conclusion_review
    result['consistency_review'] = consistency
    result['issues'].extend(consistency['issues'] + formula_review['issues'] + conclusion_review['issues'])
    result['summary'] += f" 公式核查{len(formula_review['checks'])}项，内部一致性与表达提示{len(consistency['issues'])+len(conclusion_review['issues'])}项。"
    result['hard_issue_count'] = sum(i.get('severity') == 'hard' for i in result['issues'])
    quality_review['delivery_blocked'] = bool(quality_review.get('delivery_blocked') or result['hard_issue_count'])
    result['reference_snapshot'] = {
        'analysis_run_id': (run or {}).get('run_id'),
        'valuation_run_id': (valuation or {}).get('run_id'),
        'forensic_run_id': (forensic or {}).get('run_id'),
        'reference_hashes': {k: hashlib.sha256(json.dumps(v, sort_keys=True, default=str).encode()).hexdigest()
                             for k, v in [('analysis', run), ('valuation', valuation), ('forensic', forensic)] if v},
        'documents': (run or {}).get('documents', []),
        'company': (run or {}).get('company'), 'currency': (run or {}).get('currency'),
        'evidence_scope': '仅核对指定本地运行；相符不等于对源文件及全部叙述独立鉴证',
    }
    saved = save_run(task_id, {**result, 'kind': 'audit', 'filenames': [filename] if filename else [],
        'parent_run_id': (run or {}).get('run_id', ''),
        'forensic_run_id': (forensic or {}).get('run_id', ''),
        'valuation_run_id': (valuation or {}).get('run_id', ''), 'warnings': [result.get('note', '')]})
    directory = run_dir(task_id, saved['run_id'])
    filenames = report_review_exports.export_review(directory, saved)
    exports = []
    for fmt, name in filenames.items():
        path = directory / name
        exports.append({'format': 'md' if fmt == 'markdown' else fmt, 'path': str(path),
                        'url': f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/{quote(name)}',
                        'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'size': path.stat().st_size})
    saved['exports'] = exports
    _write_json(directory / 'run.json', saved)
    return {**result, 'kind': 'audit', 'run_id': saved['run_id'], 'exports': exports,
            'complete_result': f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/run.json'}


def _load_run_anywhere(task_id: str, run_id: str) -> dict:
    """先按当前任务找，找不到就在所有运行根目录里找。

    同行比较里的其它公司往往是分开几次任务上传的，它们的运行不在当前 task 下，
    只认当前 task 会让"五家同行"根本凑不齐。
    """
    try:
        return load_run(task_id, run_id)
    except (OSError, ValueError):
        pass
    for root in _runs_roots():
        matches = sorted(root.glob(f'*/{run_id}/run.json'))
        if not matches:
            continue
        try:
            return json.loads(matches[-1].read_text(encoding='utf-8'))
        except (OSError, ValueError):
            continue
    raise ValueError(f'找不到运行：{run_id}')


def peer_benchmark(task_id: str, args: dict) -> dict:
    """把多家公司的分析运行摊成同行分位表。

    竞赛 MVP 硬指标里有「五家同行可比公司」，且消融实验要证明"同行分位减少
    对行业周期的误判"。这一步输出的 `peer_payload` 可以直接作为
    `forensic_scan` 的 `peer` 参数，让 peer_extremeness 用真实分位加权。
    """
    from . import peers

    run_ids = args.get('run_ids') or []
    if not isinstance(run_ids, list) or len(run_ids) < 2:
        raise ValueError('run_ids 至少传两家公司的分析运行ID（analyze_financial_report 的返回）')
    runs = []
    for run_id in run_ids:
        run = _load_run_anywhere(task_id, str(run_id))
        if run.get('kind') != 'analysis':
            raise ValueError(f'{run_id} 不是分析运行（kind={run.get("kind")}），'
                             '同行分位只能基于 analyze_financial_report 的结果')
        runs.append(run)

    result = peers.benchmark(runs, args.get('metrics'), args.get('period'),
                             args.get('target_run_id'))
    if not result.get('ok'):
        return result
    saved = save_run(task_id, {'kind': 'peer', **result,
                               'warnings': [w for w in (result.get('warning'),) if w]})
    result['run_id'] = saved['run_id']
    result['note'] = (result.get('note', '') + ' 把 peer_payload 直接作为 '
                      'forensic_scan 的 peer 参数，peer_extremeness 就会按真实分位加权。').strip()
    return result


def run_ablation(task_id: str, args: dict) -> dict:
    """跑五个消融对照实验（RAG / Critic / 规则 / 同行 / 估值联动）。

    对应竞赛《评测指标与目标值》第三节"消融实验（四个对照）"。所有实验都跑在
    **同一份已保存的分析运行**上，只改被消融的那一层，因此结果可复现、可对照。
    """
    from . import ablations

    run_id = str(args.get('run_id') or '')
    run = (load_run(task_id, run_id) if run_id
           else _latest_run(task_id, ('analysis',)))
    if not run:
        raise ValueError('没有可用的财报分析结果，请先调用 analyze_financial_report')

    forensic_run_id = str(args.get('forensic_run_id') or '')
    forensic = (load_run(task_id, forensic_run_id) if forensic_run_id
                else _latest_run(task_id, ('forensic',)))

    peer_runs = []
    for peer_id in (args.get('peer_run_ids') or []):
        peer_runs.append(_load_run_anywhere(task_id, str(peer_id)))

    report_text = None
    if args.get('report_path'):
        from . import report_audit
        report_text = report_audit.read_report_text(str(args['report_path']))
    elif args.get('report_text'):
        report_text = str(args['report_text'])

    result = ablations.run_all(
        run, forensic_run=forensic, peer_runs=peer_runs,
        synthetic_peers=int(args.get('synthetic_peers') or 0),
        report_text=report_text)

    saved = save_run(task_id, {
        'kind': 'ablation', **result,
        'parent_run_id': run['run_id'],
        'forensic_run_id': (forensic or {}).get('run_id', ''),
        'parameters': {'synthetic_peers': int(args.get('synthetic_peers') or 0),
                       'peer_count': len(peer_runs),
                       'report_provided': bool(report_text)},
        'warnings': [],
    })
    result['run_id'] = saved['run_id']
    result['markdown'] = ablations.render_markdown(result)
    result['complete_result'] = (f"/api/financial/artifacts/{task_id}/{saved['run_id']}"
                                 f"/run.json")
    result['note'] = (
        '每个实验的 baseline 与 treatment 之间只改一个变量，输入与代码路径完全相同，'
        '因此差异可以归因到被消融的那一层。数据不足的实验会返回 insufficient_data '
        '而不是用默认值填平；同行消融若用了 synthetic_peers，必须在报告里写明'
        '那是机制演示而非真实同行结论。')
    return result


def extract_announcement_fields(task_id: str, args: dict) -> dict:
    """对公告/表格做**确定性**机械解析，供模型在此基础上填字段。


    为什么要有这一步：公告抽取的失分点大半不是"理解错了"，而是
    大写金额没换算、万元当元、日期没归一、表格里 0.05 当成 0.05%、
    把"以司法机关文书为准"编成具体法院、把"—"当成 0、把"合计"行当成一个股东。
    这些交给程序做，结果才能稳定、可复现；模型只负责按语境挑候选、填语义字段。
    """
    from . import announcement_extract as ax
    filename = str(args.get('filename') or '')
    doc = get_document(task_id, filename)
    pages = doc.get('pages') or []
    text = '\n'.join(str(p.get('text') or '') for p in pages)
    result = ax.scan(text)
    # 把 span 换成页码，便于直接引用原文
    offsets, cursor = [], 0
    for p in pages:
        seg = str(p.get('text') or '')
        offsets.append((cursor, cursor + len(seg), p.get('page')))
        cursor += len(seg) + 1

    def locate(pos: int):
        for start, end, page in offsets:
            if start <= pos < end:
                return page
        return None

    for bucket in ('amounts', 'shares', 'dates', 'percents', 'undisclosed'):
        for item in result.get(bucket, []):
            if 'start' in item:
                item['page'] = locate(item.pop('start'))
    return {
        'ok': True,
        'filename': filename,
        'page_count': len(pages),
        'ocr_used': bool(doc.get('ocr_pages') or doc.get('ocr_used')),
        'parsed': result,
        'field_alias_map': ax.FIELD_ALIASES,
        'usage': ('这是**机械解析的候选值**，不是最终字段值。'
                  '同一量可能有多处（如"本次"与"累计"），需按语境选择后写入 records；'
                  'undisclosed 里的短语表示对应字段原文未给出，必须填 null，禁止编造；'
                  'null_symbols 表示原文用该符号表示无数值，应填 null 而不是 0；'
                  'total_rows 是汇总行，不得当作数据条目。'
                  '填好后仍需用 validate_structured_output 校验证据支持。'),
    }


def build_valuation_model(task_id: str, args: dict) -> dict:
    from .valuation_service import build_model
    return build_model(task_id, args)


def read_valuation_model(task_id: str, args: dict) -> dict:
    from .valuation_service import read_model
    return read_model(task_id, args)


def get_valuation_market_data(task_id: str, args: dict) -> dict:
    from .valuation_service import market_data
    return market_data(task_id, args)


# ---------------------------------------------------------------- 新增：提取 / 纠错 / 质量 / 估值（增强补丁）
# 下面四个工具把补丁里新写的确定性模块接到 Agent 上。共同纪律：
#   1. 只做确定性计算，不调用模型；同一输入同一输出。
#   2. 结果全量落盘（run + 明细文件），返回里给 run_id 与下载地址；
#      返回值只放"下一步要用"的部分，避免把大表塞回上下文。
#   3. 缺数据就明说（返回 refusals / warnings / 待人工清单），不用默认值填平。

def _report_text_from_args(task_id: str, args: dict) -> tuple[str, str]:
    """研报正文的统一入口：附件名 / 路径 / 直接给正文，三者取其一。

    返回 (text, source_label)。source_label 用于在结果里标明"这段话是从哪来的"。
    """
    filename = str(args.get('filename') or '')
    report_path = str(args.get('report_path') or '')
    if filename:
        report_path = str(resolve_attachment(task_id, filename))
    text = str(args.get('text') or '')
    if not text and report_path:
        from . import report_audit
        text = report_audit.read_report_text(report_path)
    if not text.strip():
        raise ValueError('请提供报告附件名、报告路径或报告正文')
    if len(text) > 2_000_000:
        raise ValueError('报告正文超过 200 万字符，请拆分后核查')
    return text, (filename or report_path)


def extract_announcement_records(task_id: str, args: dict) -> dict:
    """公告**整条记录**的结构化抽取：类型识别 → 事件切分 → 字段映射 → 勾稽校验。

    与 `extract_announcement_fields` 的分工：
    - `extract_announcement_fields` 只做"机械解析"（大写金额、单位、日期、空值符号），
      输出的是**候选值**，由模型按语境挑选；
    - 本工具直接产出**结构化记录**（每条事件一条记录、每个字段带原文位置与规则 ID），
      并跑勾稽（持股数守恒、比例与股数互算、合计行核对），把对不上的地方放进 checks。
    抽完仍要用 `validate_structured_output` 过一遍必填与证据支持。
    """
    from . import announcement_doc as adoc
    from . import announcement_records as arec
    filename = str(args.get('filename') or '')
    doc = get_document(task_id, filename)
    ann = adoc.load_parsed_document(doc, filename=filename, kind='ocr')
    result = arec.extract_doc(ann)
    result['source_sha256'] = doc.get('sha256')
    saved = save_run(task_id, {**result, 'kind': 'announcement_records', 'filenames': [filename]})
    directory = run_dir(task_id, saved['run_id'])
    (directory / 'records.json').write_text(arec.to_json(result), encoding='utf-8')
    hard = [c for c in result.get('checks', []) if c.get('level') in ('error', 'warn')]
    return {
        'ok': True, 'run_id': saved['run_id'], 'filename': filename,
        'doc_type': result.get('doc_type'), 'doc_type_reason': result.get('doc_type_reason'),
        'records': result.get('records', []),
        'checks': result.get('checks', []),
        'review': result.get('review', []),
        'warnings': result.get('warnings', []),
        'n_records': len(result.get('records', [])),
        'n_checks_flagged': len(hard),
        'engine': result.get('engine'),
        'complete_result': f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/records.json',
        'usage': ('records 已按事件切分并带原文位置（page/loc/rule）；'
                  'checks 里 level=error/warn 的必须处理或写进"待人工核验"；'
                  '不要用模型改写 records 里的数值，需要修正请说明依据后重新解析。'),
    }


def check_report(task_id: str, args: dict) -> dict:
    """研报纠错核查：报告里的数与原始财报逐条对账，并做报告内部自洽复算。

    两类错误来源分开报，因为"错在哪、怎么发现的"完全不同：
    - `source`：报告写的数与原始财报不一致 → 必须给原始财报文件名与页码；
    - `internal`：不引入外部数据，仅凭报告自身即可复算出的矛盾（增速行与上一行不符、
      PE 与 EPS×股价不符、表头年份断裂、封面与结论评级不一致、CAGR 与表内数据不符等）。
    只报**能复算出正确值**的项；复算不出正确值的疑点进 suspects 待人工，不当成"报告错了"。
    """
    from .errorcheck_service import check_report as _check
    text, label = _report_text_from_args(task_id, args)
    names = args.get('source_filenames') or []
    if not isinstance(names, list) or any(not isinstance(n, str) for n in names):
        raise ValueError('source_filenames 必须是文件名列表')
    source_paths = [str(resolve_attachment(task_id, n)) for n in names]
    if not source_paths:
        raise ValueError('请用 source_filenames 提供原始财报附件（季报/半年报/年报）；'
                         '没有原始财报时只能做 internal 类检查，请改用 audit_report')
    result = _check(source_paths=source_paths, text=text)
    result['source_files'] = names
    saved = save_run(task_id, {**result, 'kind': 'errorcheck', 'filenames': [label],
                               'source_filenames': names})
    directory = run_dir(task_id, saved['run_id'])
    (directory / 'findings.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=1, default=str), encoding='utf-8')
    summary = result.get('summary', {})
    return {
        'ok': True, 'run_id': saved['run_id'], 'source_file': label,
        'findings': result.get('findings', []),
        'suspects': result.get('suspects', []),
        'summary': summary,
        'n_findings': summary.get('n_findings'),
        'n_suspects': summary.get('n_suspects'),
        'complete_result': f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/findings.json',
        'usage': ('findings 每条含 category/subtype/quote/claimed/expected/suggestion/reason，'
                  '可直接生成勘误清单；suspects 是"复算不出正确值"的疑点，必须人工判断，'
                  '不得直接写成"报告错误"。检测覆盖率取决于所提供的原始财报是否覆盖报告引用的期间。'),
    }


def assess_report_quality(task_id: str, args: dict) -> dict:
    """研报质量评估 v2：可确定性检测的子指标 → 维度分 → 总分 + 复核队列。

    评的是**可观察**的东西：来源标注率、数字密度、观点—证据对应、因果连接词与量化、
    估值方法与假设、风险具体性与量化、绝对化措辞、内部一致性、显式公式复算等。
    每个子指标给 0–5 分、命中原文行号与规则 ID；不适用的返回 None 并从加权均值中剔除，
    同时报出覆盖率。**这是项目内部量表，不是竞赛官方评分。**
    """
    from .report_quality_v2 import assess_quality
    from .quality_report_html import write_quality_html
    text, label = _report_text_from_args(task_id, args)
    result = assess_quality(text)
    result['source_file'] = label
    result['report_sha256'] = hashlib.sha256(text.encode()).hexdigest()
    saved = save_run(task_id, {**result, 'kind': 'quality', 'filenames': [label],
                               'quality_review': {'dimensions': [
                                   {'id': d, **info} for d, info in result['dimensions'].items()]}})
    directory = run_dir(task_id, saved['run_id'])
    exports = []
    try:
        html_path = directory / 'quality_report.html'
        write_quality_html(html_path, result, report_name=label)
        exports.append({'format': 'html', 'path': str(html_path),
                        'url': f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/quality_report.html',
                        'sha256': hashlib.sha256(html_path.read_bytes()).hexdigest(),
                        'size': html_path.stat().st_size})
    except Exception as error:                          # 导出失败不能吞掉评分结果
        result.setdefault('warnings', []).append(f'HTML 报告导出失败：{error}')
    _write_json(directory / 'run.json', {**saved, 'exports': exports})
    return {
        'ok': True, 'run_id': saved['run_id'], 'source_file': label,
        'total': result['total'], 'coverage': result['coverage'],
        'dimensions': result['dimensions'],
        'subs': result['subs'],
        'review_queue': result['review_queue'],
        'consistency': result['consistency'],
        'scale_note': result['scale_note'],
        'coverage_note': result['coverage_note'],
        'exports': exports,
        'complete_result': f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/run.json',
        'usage': ('引用总分时必须同时给出 coverage，否则"覆盖率低但分高"会被误读成报告很好；'
                  'review_queue 里 low_sub 是低于 2.5 分的子指标，not_applicable 是"该报告本来'
                  '没有这类内容"（不是低分）。多份报告排序请对每份各调一次再比 total。'
                  'exports 里的 HTML 是单文件可交互报告（离线可打开），可直接交付给用户查看。'),
    }


def build_pro_valuation_model(task_id: str, args: dict) -> dict:
    """A 股口径的自动估值建模（增强版）：逐年驱动预测 + 三表联动 + 活公式 Excel。

    与 `build_valuation_model` 的分工：本工具是补丁里重写的估值内核，强调
    (1) 驱动变量逐项标注来源（财报实算 / 市场输入 / 显式假设），不证实的不当实时报价；
    (2) 输出**活公式** Excel（引用单元格而非写死数值），可用 LibreOffice 重算校验；
    (3) 明确 refusals：缺关键数据就拒绝建模并给出补数清单，不用默认值填平。
    复用 `analyze_financial_report` 的 run_id 取事实，避免"用待核查报告自证"。
    """
    from .valuation_pro import build_pro_model
    from .valuation_xlsx_pro import export_with_cache
    run_id = str(args.get('run_id') or '')
    analysis = load_run(task_id, run_id) if run_id else _latest_run(task_id, ('analysis',))
    if not analysis:
        raise ValueError('请先调用 analyze_financial_report 完成财报分析，或指定其 run_id')
    if analysis.get('task_id') != _id(task_id):
        raise ValueError('运行记录不属于当前任务')
    if analysis.get('kind') != 'analysis':
        raise ValueError('run_id 必须指向财报分析结果')
    options = args.get('options') or {}
    if not isinstance(options, dict):
        raise ValueError('options 必须是对象')
    for key in ('industry', 'valuation_date', 'market_symbol', 'benchmark', 'currency', 'base_year', 'price', 'market_cap'):
        if args.get(key) not in (None, ''):
            options.setdefault(key, args[key])
    options.setdefault('company', analysis.get('company') or '')
    result = build_pro_model(analysis.get('facts') or [], options)
    saved = save_run(task_id, {**result, 'kind': 'valuation_pro', 'parameters': options,
                               'parent_run_id': analysis.get('run_id', ''),
                               'filenames': analysis.get('filenames') or []})
    directory = run_dir(task_id, saved['run_id'])
    exports = []
    if result.get('ok'):
        try:
            path = directory / 'valuation_pro.xlsx'
            export_with_cache(result, path)
            exports.append({'format': 'xlsx', 'path': str(path),
                            'url': f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/{quote(path.name)}',
                            'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                            'size': path.stat().st_size})
        except Exception as error:                      # 导出失败不能吞掉建模结果
            result.setdefault('warnings', []).append(f'活公式 Excel 导出失败：{error}')
    (directory / 'valuation_pro.json').write_text(
        json.dumps(result, ensure_ascii=False, indent=1, default=str), encoding='utf-8')
    _write_json(directory / 'run.json', {**saved, 'exports': exports})
    return {**result, 'kind': 'valuation_pro', 'run_id': saved['run_id'], 'exports': exports,
            'complete_result': f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/valuation_pro.json'}


def run_finance_workflow(task_id: str, args: dict, cancel_event=None) -> dict:
    from .finance_workflow import run_workflow
    return run_workflow(task_id,args,cancel_event)


TOOLS = []
def _tool(name, description, properties, required=()):
    TOOLS.append({'type': 'function', 'function': {'name': name, 'description': description,
        'parameters': {'type': 'object', 'properties': properties, 'required': list(required)}}})


_tool('inspect_document', '完整解析当前任务附件，自动OCR扫描页。先不传page看文档/表格目录，再按page/table_id/query和offset分段读取带evidence_id的原文，不能把预览当完整文档。',
      {'filename': {'type': 'string'}, 'page': {'type': 'integer'}, 'table_id': {'type': 'string'},
       'query': {'type': 'string'}, 'offset': {'type': 'integer'}, 'limit': {'type': 'integer'}}, ['filename'])
_tool('extract_financial_data', '从财报抽取并标准化财务事实，保留单位、期间、口径、原始值和出处；缺失/冲突进入复核。完整结果落盘，返回run_id。',
      {'filename': {'type': 'string'}, 'company': {'type': 'string'}, 'period': {'type': 'string'},
       'scope': {'type': 'string', 'enum': ['auto', 'consolidated', 'parent', 'unknown']}}, ['filename'])
_tool('analyze_financial_report', '程序核验上市公司财报并用Decimal计算财务指标、同比、环比、可推导单季、现金转化率、非经常性损益汇总占比、勾稽与异常信号；并补充单季抽取（年报「分季度主要财务指标」表，用四季度之和==年度值勾稽）、会计口径变化（同一科目同一期间多值）检测与"业绩变动原因"段落定位。可一次传多个报告文件；自动导出Excel/JSON/Markdown，返回公式和证据。不要用心算替代。环比指标 code 以 _qoq 结尾（derived=true），单季事实在 quarter_facts 段，都属本项目补充计算；restatement_changes 只收口径已知的组，口径未知的进 restatement_suspects（不得当口径变化引用）；口径变化成因与业绩归因不由程序判定，须人工或模型标注为推断。',
      {'filenames': {'type': 'array', 'items': {'type': 'string'}}, 'company': {'type': 'string'},
       'period': {'type': 'string'}, 'scope': {'type': 'string', 'enum': ['auto', 'consolidated', 'parent', 'unknown']}})
_tool('extract_announcement_fields', '对公告/表格做确定性机械解析：中文大写金额换算、万元/亿元/万股量级归一、日期格式归一、百分数小数形式识别、未披露短语定位、"—"空值符号、合计行识别、字段名笔误归一。返回带页码的候选值。公告类任务应先用本工具再填字段，避免单位与格式类错误；它不做语义判断，同一量有多处时需按语境选择。',
      {'filename': {'type': 'string'}}, ['filename'])
_tool('validate_structured_output', '验证公告等抽取候选的固定字段、必填完整性和证据原文支持。records=[{字段:{value, evidence_ids:[ID], raw_value:可选原文}}]。缺失设null；不支持的归一化需复核，不证明所有语义关系。',
      {'filename': {'type': 'string'}, 'records': {'type': 'array', 'items': {'type': 'object'}},
       'required_fields': {'type': 'array', 'items': {'type': 'string'}}}, ['filename', 'records'])
_tool('export_financial_report', '从已保存run_id直接导出带事实、指标、公式、来源和待复核清单的Excel/Word/JSON/Markdown，无需shell。返回下载链接。',
      {'run_id': {'type': 'string'}, 'format': {'type': 'string', 'enum': ['xlsx', 'docx', 'json', 'md']}}, ['run_id'])
_tool('read_financial_run', '分页读取已保存运行的全部事实、指标、勾稽、法证发现、估值情景或复核清单。用于结果预览被限制时继续读取；可按标准科目code和期间period筛选。section=metrics 里 code 以 _qoq 结尾的是**本项目补充计算的环比**（derived=true），与引擎原生指标区分；section=restatement_changes 读同一科目同一期间的多值情况（重述/口径变化候选，只给差异不判成因）；section=reason_candidates 读"业绩变动原因"候选段落（只定位，归纳须由模型完成并标为推断）。',
      {'run_id': {'type': 'string'}, 'section': {'type': 'string', 'enum': ['facts', 'metrics', 'checks', 'review_queue', 'records', 'field_results', 'signals', 'findings', 'tieout_diagnosis', 'scenarios', 'experiments', 'issues', 'competition_metrics', 'claims', 'quality_dimensions', 'restatement_changes', 'reason_candidates', 'disclosure_map', 'quarter_facts', 'restatement_suspects']},
       'offset': {'type': 'integer'}, 'limit': {'type': 'integer'}, 'code': {'type': 'string'}, 'period': {'type': 'string'}}, ['run_id'])
_tool('audit_report', '本地研报核查与质量评估：按科目、期间、口径绑定数字和原始证据，区分已验证/矛盾/无法验证；输出错误位置、修订建议和Word/Excel/JSON/Markdown审阅文件。质量是项目内部量表，不是竞赛官方评分，未评部分为null，须报告评估覆盖率。指定分析run_id，避免用报告自证或自动选错公司。报告交付前调用一次。',
      {'run_id': {'type': 'string', 'description': 'analyze_financial_report 的运行ID，提供事实与指标池'},
       'text': {'type': 'string', 'description': '报告正文（Markdown/纯文本）'},
       'report_path': {'type': 'string', 'description': '报告文件路径，支持 .md/.txt/.docx/.pdf'},
       'filename': {'type': 'string', 'description': '当前任务已上传的研报附件名，与text二选一'},
       'valuation_run_id': {'type': 'string', 'description': 'scenario_valuation 的运行ID，可选'},
       'forensic_run_id': {'type': 'string', 'description': 'forensic_scan 的运行ID，可选'}})
_tool('peer_benchmark', '把多家公司的财报分析运行摊成同行分位表（中位数/P25/P75/极值），并标出目标公司在同行中的分位位置与显著偏离项。只比可比指标（比率、周转天数），不比规模类绝对值。期间不一致会显式告警。返回的 peer_payload 可直接作为 forensic_scan 的 peer 参数，让 peer_extremeness 按真实四分位距加权。run_ids 传各公司 analyze_financial_report 的返回。',
      {'run_ids': {'type': 'array', 'items': {'type': 'string'},
                   'description': '各公司 analyze_financial_report 的 run_id，至少两个'},
       'target_run_id': {'type': 'string', 'description': '目标公司运行ID，用于标出它在同行中的位置'},
       'metrics': {'type': 'array', 'items': {'type': 'string'},
                   'description': '要比较的指标代码，默认一组可比比率与天数'},
       'period': {'type': 'string', 'description': '比较期间，默认取各公司都齐的最新期'}},
      ['run_ids'])
_tool('forensic_scan', '在已核验的财报事实上跑确定性财务法证规则，输出按风险分排序的待核验风险清单。每条含rule_id、severity、risk_score五因子分解、支持证据、反证检索方向、未解问题、估值传导路径；并自动反查勾稽失败是否为科目单位错位。必须在 analyze_financial_report 之后调用。',
      {'run_id': {'type': 'string'}, 'top_k': {'type': 'integer'},
       'peer': {'type': 'object', 'description': '{科目代码:{median,p25,p75,n}} 同行分位；缺失时 peer_extremeness 取中性 50 并标记'}})
_tool('scenario_valuation', '基准/乐观/压力三情景预测与DCF估值压力测试。DCF 的全部驱动变量（收入增速、毛利率、期间费用率、有效税率、折旧摊销率、资本开支率、周转天数、净债务、非经营性资产、总股本）默认由已核验事实与指标自动校准，并逐项记录来源；压力情景参数由 forensic_scan 的发现自动推导，输出权益价值区间、每股价值区间与WACC×永续增长率敏感性矩阵，不输出目标价。同时给出 DCF 隐含的 EV/EBITDA 与 P/E 自检。可在overrides里人工覆盖任何驱动变量，也可给 comparable_ev_ebitda / comparable_pe 让工具反推可比倍数法隐含价值，与 DCF 形成双方法交叉验证。',
      {'run_id': {'type': 'string'}, 'forensic_run_id': {'type': 'string'},
       'use_findings': {'type': 'boolean'},
       'overrides': {'type': 'object', 'description': '驱动变量覆盖，如 {wacc:0.11, revenue_growth:0.05, receivable_days:90, comparable_pe:20, comparable_ev_ebitda:15}'}})

_tool('run_ablation', '跑竞赛第三节要求的消融实验（五个对照），输出可复现的对照表：RAG 消融（证据检索）、Critic 消融（反证 Agent）、规则消融（确定性勾稽规则，含"注入万元/元量级错误"的检出能力测试）、同行消融（同行分位 vs 只看公司历史）、估值联动消融（风险参数化是否影响估值结论）。每个实验只改一个变量，输入与代码路径相同；数据不足的实验返回 insufficient_data 而不是用默认值填平。写完估值建模或方法学章节时应调用本工具取数，不要自己编对照数字。',
      {'run_id': {'type': 'string', 'description': 'analyze_financial_report 的运行ID，默认取最近一次'},
       'forensic_run_id': {'type': 'string', 'description': 'forensic_scan 的运行ID，默认取最近一次'},
       'peer_run_ids': {'type': 'array', 'items': {'type': 'string'},
                        'description': '同业公司的分析运行ID（同行消融用，≥1 家才有效）'},
       'synthetic_peers': {'type': 'integer',
                           'description': '拿不到真实同业时生成 N 家模拟同业；仅机制演示，报告必须写明'},
       'report_path': {'type': 'string', 'description': '报告文件路径，提供后把引用精确率纳入 RAG/Critic 消融'},
       'report_text': {'type': 'string', 'description': '报告正文，与 report_path 二选一'}})

_tool('build_valuation_model', '自动选择普通企业多阶段FCFF或金融企业剩余收益模型；逐年驱动预测、终值再投资、CAPM/WACC来源、风险三情景、敏感性、反向DCF、蒙特卡洛及有来源同行估值，自动导出可编辑Excel/Word/JSON。普通企业还会逐年生成**配平的资产负债表+利润表+现金流量表**（五项勾稽：资产=负债+权益、现金滚动、留存收益滚存、非流动资产滚存、现金非负），并用三表反推的 FCFF 与驱动法逐项对账。折现率会先用财报侧可实算项（债务成本、资本结构、行业β加杠杆）校准并给出参数敏感度归因。压力情景=证据驱动传导 + 结构性兜底（扫描无发现时压力情景仍成立），并自检三情景是否真的互不相同。复用财报分析run_id，缺失关键数据则返回补数清单；未证实市场输入不能当实时报价。',
      {**{k:{'type':'string'} for k in ('run_id','forensic_run_id','industry','valuation_date','market_symbol','benchmark')},
       **{k:{'type':'object'} for k in ('options','overrides','market','forecast','equity_model','monte_carlo','reverse_dcf')},
       'filenames':{'type':'array','items':{'type':'string'}}, 'peers':{'type':'array','items':{'type':'object'}},
       'export_formats':{'type':'array','items':{'type':'string','enum':['xlsx','docx','json','md']}},
       'allow_illustrative_model':{'type':'boolean'},'allow_estimated_shares':{'type':'boolean'}})
_tool('read_valuation_model','分页读取已保存自动估值模型。projection=逐年驱动预测；three_statements=该情景的**平衡三表**（含五项勾稽结果与 FCFF 交叉验证，报告可直接引用"资产=负债+权益残差 0 元"）；fcff_cross_check=三表反推 FCFF 与驱动法逐项对账；discount_rate_detail=折现率校准明细与参数归因；method_triangle=DCF/相对估值/资产基础法的加权区间与离散度；scenario_distinctness=三情景是否真的互不相同；stress_diagnostic=压力情景里哪些来自证据、哪些来自结构性兜底；financing_rates=三表计息用的借款成本与存款收益率。',
      {'run_id':{'type':'string'},'section':{'type':'string','enum':['summary','valuation_review','projection','scenarios','assumption_sources','checks','peers','monte_carlo','reverse_dcf','sensitivity','terminal_policy','inputs_used','options','market_snapshot','three_statements','three_statement_rows','fcff_cross_check','discount_rate_detail','method_triangle','scenario_distinctness','financing_rates','terminal_policy_checks','stress_diagnostic']},
       'scenario':{'type':'string','enum':['base','optimistic','stress']},'offset':{'type':'integer'},'limit':{'type':'integer'}},['run_id'])
_tool('get_valuation_market_data','获取Yahoo历史行情快照或当前任务CSV并估计beta；明确日期/币种/回归样本与来源。不能生成无风险利率、ERP、股数等缺失输入。失败可用本地复权CSV。',
      {k:{'type':'string'} for k in ('symbol','benchmark','period','as_of','price_file','benchmark_file','currency')})

_tool('extract_announcement_records', '公告整条记录的结构化抽取：先识别公告类型（质押/冻结/减持/担保/中标/业绩预告…），再按事件切分，逐字段映射并跑勾稽（持股数守恒、比例与股数互算、合计行核对）。每条记录带原文位置（page/loc/rule）与置信度，勾稽不符进 checks。与 extract_announcement_fields 的分工：后者只做机械解析给候选值，本工具直接给结构化记录。抽完仍需 validate_structured_output 校验。',
      {'filename': {'type': 'string', 'description': '当前任务已上传的公告附件名'}}, ['filename'])
_tool('check_report', '研报纠错核查：把报告里写的数与原始财报逐条对账（source 类），并做报告内部自洽复算（internal 类：增速行与上一行不符、PE与EPS×股价不符、表头年份断裂、封面与结论评级不一致、CAGR与表内数据不符、正文与预测表不一致）。每条含类别/原文/现值/应为值/修改建议/依据文件名与页码。只报能复算出正确值的项，复算不出的进 suspects 待人工。必须提供原始财报附件，否则只能做内部检查。',
      {'filename': {'type': 'string', 'description': '待核查研报的附件名，与 report_path/text 三选一'},
       'report_path': {'type': 'string', 'description': '研报文件路径（.md/.txt/.docx/.pdf）'},
       'text': {'type': 'string', 'description': '研报正文'},
       'source_filenames': {'type': 'array', 'items': {'type': 'string'},
                            'description': '原始财报附件名列表（季报/半年报/年报），报告引用了哪几期就传哪几份'}},
      ['source_filenames'])
_tool('assess_report_quality', '研报质量评估 v2：按可确定性检测的子指标打分（结构完整性、来源标注率、数字密度、图表来源、观点—证据对应、因果论证、事实/推论/观点区分、反证条件、估值方法与假设、区间与情景、风险具体性与量化、风险与盈利预测联动、绝对化措辞、免责声明、内部一致性、显式公式复算），每个子指标给 0–5 分并附命中原文行号与规则ID；不适用项从加权均值中剔除并报出覆盖率。同时给复核队列。这是项目内部量表，不是竞赛官方评分。交付前调用一次，多份报告排序就各调一次比 total。',
      {'filename': {'type': 'string', 'description': '研报附件名，与 report_path/text 三选一'},
       'report_path': {'type': 'string'}, 'text': {'type': 'string'}})
_tool('build_pro_valuation_model', 'A股口径自动估值建模（增强内核）：逐年驱动预测、三表联动、WACC与终值来源标注、三情景、敏感性、反向DCF、蒙特卡洛、同行相对估值，并导出**活公式**Excel（可LibreOffice重算校验）。驱动变量逐项标注来源（财报实算/市场输入/显式假设）；缺关键数据返回 refusals 与补数清单，不用默认值填平。复用 analyze_financial_report 的 run_id 取事实。',
      {'run_id': {'type': 'string', 'description': 'analyze_financial_report 的运行ID，缺省取最近一次'},
       'options': {'type': 'object', 'description': '估值参数，如 {industry, valuation_date, price, market_cap, wacc, base_year, currency, unit}'},
       'industry': {'type': 'string'}, 'valuation_date': {'type': 'string'},
       'market_symbol': {'type': 'string'}, 'benchmark': {'type': 'string'},
       'currency': {'type': 'string'}, 'base_year': {'type': 'integer'},
       'price': {'type': 'number'}, 'market_cap': {'type': 'number'}})

_tool('run_finance_workflow','按明确选择的附件联动五项能力。财报基础分析只执行一次，下游复用run_id；可选公告整条记录、专业估值、研报来源对账及质量评估。缺材料列needs_input，疑点列needs_review；不会自动选择最近公司。goals只传本轮需要的能力，默认五项。质量总分不作事实认证。',
      {'goals': {'type':'array','items':{'type':'string','enum':['extraction','financial_analysis','valuation','report_check','quality']}},
       'financial_filenames':{'type':'array','items':{'type':'string'}}, 'announcement_filenames':{'type':'array','items':{'type':'string'}},
       'analysis_run_id':{'type':'string'},'report_filename':{'type':'string'},'report_text':{'type':'string'},
       'valuation_options':{'type':'object'},'company':{'type':'string'},'scope':{'type':'string'},'period':{'type':'string'}})

TOOL_HANDLERS = {f.__name__: f for f in (run_finance_workflow, inspect_document, extract_financial_data, analyze_financial_report,
                                      extract_announcement_fields, validate_structured_output,
                                      export_financial_report, read_financial_run,
                                      forensic_scan, scenario_valuation, audit_report,
                                      peer_benchmark, run_ablation, build_valuation_model, read_valuation_model, get_valuation_market_data,
                                      extract_announcement_records, check_report, assess_report_quality,
                                      build_pro_valuation_model)}
