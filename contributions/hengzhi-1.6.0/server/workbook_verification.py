"""Independent, read-only Excel formula checking through the installed MCP."""
import asyncio,hashlib,json,re,tempfile,zipfile,xml.etree.ElementTree as ET
from decimal import Decimal,InvalidOperation
from pathlib import Path
from . import config
from .mcp_client import MCPClient

async def _read(server,path,sheet,cell):
    client=MCPClient(server['name'],server['command'],server.get('args') or [],env=server.get('env') or {},cwd=server.get('cwd'))
    try:
        await client.start()
        response=await client.call_tool('read_range',{'path':str(path),'sheet':sheet,'range':cell,'mode':'values'})
        if response.get('isError'):raise ValueError('Excel MCP未能读取该单元格')
        data=response.get('structuredContent') or json.loads(response['content'][0]['text'])
        if data.get('uncalculated'):return None
        return data['values'][0][0]
    finally:await client.stop()


def _without_caches(source,destination):
    """Preserve formulas and inputs; remove every formula's stored result in a copy."""
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    cleared=0
    with zipfile.ZipFile(source) as original,zipfile.ZipFile(destination,'w',zipfile.ZIP_DEFLATED) as copy:
        for entry in original.infolist():
            content=original.read(entry.filename)
            if entry.filename.startswith('xl/worksheets/') and entry.filename.endswith('.xml'):
                tree=ET.fromstring(content)
                for cell in tree.iter(ns+'c'):
                    if cell.find(ns+'f') is not None:
                        for tag in ('v','is'):
                            for value in list(cell.findall(ns+tag)):cell.remove(value)
                        cell.attrib.pop('t',None);cleared+=1
                content=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            copy.writestr(entry,content)
    return cleared


def verify_workbook(task_id,model):
    from . import finance_service as fs
    result={'status':'unavailable','scope':'仅验证Excel公式输出与模型引擎一致，不证明原始输入或预测正确'}
    server=next((s for s in config.load().get('mcp_servers') or [] if s.get('name')=='fin-excel-audit' and s.get('enabled')),None)
    if not server:return {**result,'reason':'未启用本地Excel MCP，未作独立复算'}
    if not model.get('run_id'):return {**result,'reason':'缺少模型运行ID'}
    folder=fs.run_dir(task_id,model['run_id']).resolve()
    files=[folder/'valuation_pro.xlsx',folder/'报告.xlsx']
    path=next((p for p in files if p.is_file() and not p.is_symlink()),None)
    if path is None:return {**result,'reason':'没有已保存的Excel模型'}
    before=hashlib.sha256(path.read_bytes()).hexdigest()
    try:
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False)
        try:
            named=book.defined_names.get('Equity_DCF')
            if named:
                destinations=list(named.destinations)
                if len(destinations)!=1:raise ValueError('权益价值名称未指向唯一单元格')
                sheet,cell=destinations[0];cell=cell.replace('$','')
                expected=(model.get('scenarios') or {}).get('base',{}).get('equity')
            elif 'DCF估值' in book.sheetnames:
                sheet,cell='DCF估值','B13';expected=(model.get('scenarios') or {}).get('base',{}).get('valuation',{}).get('equity_value')
            else:raise ValueError('未识别的模型布局，不猜测输出单元格')
            if sheet not in book.sheetnames or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d{0,6}',cell):raise ValueError('模型输出坐标无效')
            if not isinstance(book[sheet][cell].value,str) or not book[sheet][cell].value.startswith('='):raise ValueError('权益价值不是活公式，不能称独立公式复算')
        finally:book.close()
        target=Decimal(str(expected))
        if not target.is_finite():raise ValueError('模型基准值不是有限数值')
        with tempfile.NamedTemporaryFile(prefix='formula-audit-',suffix='.xlsx',dir=folder,delete=False) as temp:copy_path=Path(temp.name)
        try:
            cleared=_without_caches(path,copy_path)
            actual=asyncio.run(_read(server,copy_path,sheet,cell))
            result.update(formula_caches_cleared=cleared,calculation_basis='临时副本清除所有公式缓存后由MCP重算')
        finally:copy_path.unlink(missing_ok=True)
        if actual is None:raise ValueError('公式引擎未能计算该结果，不能按0或通过处理')
        actual=Decimal(str(actual))
        if not actual.is_finite():raise ValueError('Excel结果不是有限数值')
        tolerance=max(Decimal('.01'),abs(target)*Decimal('1e-10'))
        result.update(status='matched' if abs(actual-target)<=tolerance else 'mismatched',sheet=sheet,cell=cell,model_value=str(target),excel_value=str(actual),tolerance=str(tolerance),difference=str(actual-target))
    except Exception as error:result.update(status='needs_review',reason=str(error))
    after=hashlib.sha256(path.read_bytes()).hexdigest()
    if after!=before:result.update(status='needs_review',reason='复算期间工作簿发生变化，无法确认同一版本')
    result.update(workbook_sha256=after,model_run_id=model['run_id'])
    saved=fs.save_run(task_id,{**result,'kind':'workbook_verification','filenames':[],'parent_run_id':model['run_id']})
    return {**result,'run_id':saved['run_id'],'complete_result':f'/api/financial/artifacts/{task_id}/{saved["run_id"]}/run.json'}
