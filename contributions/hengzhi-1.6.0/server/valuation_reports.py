"""Chinese review artifacts from the canonical engine result, without new estimates."""
from __future__ import annotations
from decimal import Decimal, InvalidOperation
import json
import html
from pathlib import Path

from .valuation_workbook import export_valuation_workbook

LABELS = {
    'base':'基准','optimistic':'乐观','stress':'压力','multi_stage_fcff':'分阶段经营现金流估值',
    'residual_income':'金融企业剩余收益估值','explicit_override':'用户指定',
    'default_assumption':'默认建模假设','CAPM':'有来源市场参数推导',
    'user_override':'用户指定假设','fact':'财报事实','metric':'财报计算指标','derived':'财报推导',
    'user_forecast':'用户预测假设','user_forecast_path':'用户逐年预测路径',
    'revenue_base':'基期营业收入','revenue_growth':'收入增长率','gross_margin':'毛利率',
    'opex_ratio':'经营费用率','tax_rate':'所得税率','da_to_revenue':'折旧摊销比例',
    'capex_to_revenue':'资本开支比例','receivable_days':'应收周转天数','inventory_days':'存货周转天数',
    'payable_days':'应付周转天数','wacc':'折现率','terminal_growth':'永续增长率','terminal_roic':'稳态投入资本回报率',
    'net_debt':'净债务','minority_interest':'少数股东权益','non_operating_assets':'非经营性资产',
    'shares':'总股本','years':'预测年数','book_equity':'期初账面权益','net_income':'首年净利润',
    'cost_of_equity':'权益资本成本','roe':'初始净资产收益率','payout_ratio':'分红比例',
    'terminal_roe':'稳态净资产收益率','near_term_fcff_haircut':'近期现金流折减',
    'sustainable_profit_discount':'持续盈利折减','impairment_to_revenue':'首年减值比例',
    'CNY':'人民币元','USD':'美元','HKD':'港元',
}
RATIOS={'revenue_growth','gross_margin','opex_ratio','tax_rate','da_to_revenue','capex_to_revenue','wacc','terminal_growth','terminal_roic','roe','terminal_roe','payout_ratio','cost_of_equity','near_term_fcff_haircut','sustainable_profit_discount','impairment_to_revenue'}


def _text(value):
    if value is None or value == '': return '—'
    if isinstance(value,(dict,list)): return json.dumps(value,ensure_ascii=False,default=str,allow_nan=False)
    return str(value)


def _num(value, percent=False):
    if value is None or value=='': return '—'
    try:
        v=Decimal(str(value))
        if not v.is_finite(): return '—'
        return f'{v*100:,.2f}%' if percent else f'{v:,.2f}'
    except (InvalidOperation,ValueError,TypeError): return _text(value)


def _blocks(run):
    """One semantic report shared by Markdown and Word, so evidence stays consistent."""
    blocks=[]
    def title(text,level=1): blocks.append(('heading',level,text))
    def paragraph(text): blocks.append(('paragraph',text))
    def table(headers,rows):
        if rows: blocks.append(('table',headers,[[str(v) for v in row] for row in rows]))
        else: paragraph('暂无可用结果；见数据缺口与模型检查。')
    title('自动估值建模报告')
    plan=run.get('model_plan') or {}; currency_code=plan.get('currency') or run.get('currency'); currency=LABELS.get(currency_code,currency_code or '以输入财报为准')
    paragraph('状态：'+('模型可计算' if run.get('ok') else '建模受阻，保留诊断与已取得证据')+'。运行编号：'+_text(run.get('run_id'))+'；模型版本：'+_text(run.get('engine_version'))+'。')
    paragraph('金额均按财报币种基础单位记录，未缩放为万元或亿元。币种：'+str(currency)+'；股数单位为股。缺失值使用“—”，不替换为零。')
    title('模型选择与适用范围',2)
    paragraph('采用模型：'+LABELS.get(plan.get('method'),_text(plan.get('method')))+'；预测年数：'+_text(plan.get('years'))+'。')
    paragraph('经营驱动预测用于估算现金流，并未生成三张完整且平衡的预测财务报表。默认市场参数、增长与利润率等未来假设需要结合公司情况核验，不视为公司已观察事实。')
    paragraph('Excel 中的假设、逐年预测和估值保留可编辑公式。风险传导、可比公司、反向估值及蒙特卡洛为本次运行快照；修改 Excel 假设后须重新运行模型更新诊断。')
    paragraph('模型选择记录：'+_text(plan))
    title('三情景估值与权益价值区间',2)
    scenarios=run.get('scenarios') or {}; scenario_items=[(k,scenarios[k]) for k in ('base','optimistic','stress') if isinstance(scenarios.get(k),dict)]
    rows=[]
    for key,sc in scenario_items:
        v=sc.get('valuation') or {}
        rows.append([sc.get('label',LABELS[key]),_num(v.get('enterprise_value')),_num(v.get('equity_value')),_num(v.get('value_per_share')),_num(v.get('terminal_value_share'),True),'可计算' if v.get('ok') else _text(v.get('reason') or '数据不足')])
    table(['情景','企业价值','权益价值','每股价值','终值占比','状态'],rows)
    value_range=run.get('range') or {}
    paragraph('权益价值区间：'+_num(value_range.get('equity_value_low'))+' 至 '+_num(value_range.get('equity_value_high'))+'；情景中值：'+_num(value_range.get('equity_value_mid'))+'。股本缺失时每股价值为空。')
    title('折现率、终值与来源',2)
    rate=run.get('discount_rate') or {}
    paragraph('折现率依据：'+LABELS.get(rate.get('method'),_text(rate.get('method')))+'。')
    table(['字段','值 / 来源'],[[LABELS.get(k,k),_num(v,True) if k in RATIOS else _text(v)] for k,v in rate.items()])
    title('估值结论校验',2)
    review=run.get('valuation_review') or {}
    paragraph(_text(review.get('summary')))
    for reason in review.get('reasons',[]):paragraph(_text(reason))
    table(['同行方法','有效样本','同行权益价值中位','相对差异'],[[row.get('method'),row.get('sample_count'),_num(row.get('equity_value_median')),_num(row.get('relative_gap'),True)] for row in review.get('method_comparisons',[])])
    paragraph('账面净资产仅作会计参照，不是企业价值底线；多方法结果不自动平均。')
    paragraph('终值政策：'+_text(run.get('terminal_policy')))
    if plan.get('method')=='residual_income':
        paragraph('权益价值 = 期初账面权益 + 显性期剩余收益现值 + 终值剩余收益现值；剩余收益 = 净利润 − 期初账面权益 × 权益资本成本。终值增长由稳态 ROE 与分红比例决定。')
    else:
        paragraph('企业价值 = 显性期自由现金流现值 + 终值现值；权益价值 = 企业价值 − 净债务 − 少数股东权益 + 非经营性资产。终值现金流采用下一年经营税后利润 × (1 − 永续增长率 / 稳态投入资本回报率)，并传导持续盈利折减。')
    title('数据缺口、风险与模型检查',2)
    warnings=(run.get('warnings') or [])+(run.get('data_needs') or [])
    if warnings:
        for item in warnings: paragraph(_text(item))
    else: paragraph('本次运行没有新增数据缺口提示；仍请核验输入证据与默认假设。')
    checks=run.get('checks') or []
    if isinstance(checks,dict): checks=[{'name':k,'detail':v} for k,v in checks.items()]
    table(['检查','结果','说明'],[[_text(c.get('name',c.get('code'))),'通过' if c.get('ok') is True else '需核验' if c.get('ok') is False else _text(c.get('status')),_text(c.get('detail',c.get('message',c.get('reason'))))] if isinstance(c,dict) else ['检查','—',_text(c)] for c in checks])
    if run.get('stress_diagnostic'): paragraph('压力测试充分性：'+_text(run['stress_diagnostic']))
    title('逐年预测与情景假设',2)
    for key,sc in scenario_items:
        title(sc.get('label',LABELS[key]),3)
        params=sc.get('parameters') or {}; projection=sc.get('projection') or []
        table(['参数','本情景输入'],[[LABELS.get(k,k),_num(v,k in RATIOS)] for k,v in params.items()])
        if projection and 'opening_book_equity' in projection[0]:
            fields=[('year','预测年'),('opening_book_equity','期初账面权益'),('net_income','净利润'),('dividends','股利'),('residual_income','剩余收益'),('closing_book_equity','期末账面权益')]
            table([label for _,label in fields],[[_text(r.get(k)) if k=='year' else _num(r.get(k)) for k,_ in fields] for r in projection])
        else:
            fields=[('year','预测年'),('revenue','营业收入'),('ebit','经营利润'),('nopat','经营税后利润'),('fcff','自由现金流')]
            table([label for _,label in fields],[[_text(r.get(k)) if k=='year' else _num(r.get(k)) for k,_ in fields] for r in projection])
            fields=[('year','预测年'),('depreciation_amortization','折旧摊销'),('capex','资本开支'),('nwc','营运资本'),('delta_nwc','营运资本增加'),('pv_fcff','现金流现值')]
            table([label for _,label in fields],[[_text(r.get(k)) if k=='year' else _num(r.get(k)) for k,_ in fields] for r in projection])
        paragraph('风险传导记录：'+_text(sc.get('risk_adjustments') or []))
        if sc.get('scenario_status'): paragraph('情景适用状态：'+_text(sc['scenario_status']))
    title('可比公司交叉检验',2)
    paragraph('可比公司仅用于交叉核验；样本的期间、币种、口径与异常项以本次输入证据为准。')
    paragraph(_text(run.get('peers_summary') or run.get('peers')))
    title('敏感性、蒙特卡洛与反向估值',2)
    paragraph('以下结果为本次引擎快照。Excel 敏感性表固定已导出的逐年经营路径，仅改变折现率及终值增长，口径与引擎完整路径重算不同。')
    diagnostic_labels={'enabled':'是否启用','ok':'是否可计算','samples':'抽样次数','seed':'随机种子','valid_samples':'有效样本数','invalid_samples':'无效样本数','invalid_reasons':'无效原因','variables':'扰动变量与范围','distribution':'假设分布','correlations':'相关性设定','correlation_basis':'相关性口径','mean':'权益价值均值','note':'适用说明','reason':'诊断原因','status':'计算状态','parameter':'反推参数','implied_value':'隐含参数值','target_equity_value':'目标权益价值','calculated_equity_value':'匹配权益价值','iterations':'迭代次数','low_equity_value':'区间下端权益价值','high_equity_value':'区间上端权益价值'}
    for key,label in [('sensitivity','完整路径敏感性'),('monte_carlo','蒙特卡洛'),('reverse_dcf','反向估值')]:
        title(label,3); data=run.get(key) or {}
        if key=='sensitivity' and isinstance(data,dict):
            matrix=[]
            for row in data.get('matrix',[]):
                for c in row.get('cells',[]): matrix.append([_num(row.get('terminal_growth'),True),_num(c.get('wacc'),True),_num(c.get('equity_value')),'可计算' if c.get('ok') else _text(c.get('reason'))])
            table(['永续增长率','折现率','权益价值','状态'],matrix)
            table(['扰动参数','步长','向下扰动价值','向上扰动价值','绝对影响'],[[LABELS.get(r.get('parameter'),_text(r.get('parameter'))),_num(r.get('step'),True),_num(r.get('low_value')),_num(r.get('high_value')),_num(r.get('impact'))] for r in data.get('tornado',[])])
            paragraph(_text(data.get('note')))
        elif isinstance(data,dict):
            table(['项目','本次运行结果'],[[diagnostic_labels.get(k,k),('是' if v else '否') if isinstance(v,bool) else _num(v) if k.endswith('equity_value') or k=='mean' else _text(v)] for k,v in data.items() if k!='quantiles'])
            if data.get('quantiles'): table(['分位点','权益价值'],[[LABELS.get(k,k),_num(v)] for k,v in data['quantiles'].items()])
        else: paragraph(_text(data))
    title('输入证据与假设出处',2)
    for key,label in [('inputs_used','输入事实与推导'),('assumption_sources','建模假设来源'),('equity_assumption_sources','金融企业权益模型假设来源')]:
        title(label,3); data=run.get(key) or {}
        if isinstance(data,dict):
            table(['字段','完整来源与证据'],[[LABELS.get(k,k),_text(v)] for k,v in data.items()])
        else: paragraph(_text(data))
    paragraph('原文、页码、文档名、期间、币种、口径、证据编号与来源链接按输入记录保留。完整逐年驱动数据、数值精度和诊断结构保存在 JSON 与 Excel 中。')
    return blocks


def _markdown(blocks):
    def safe(value): return html.escape(str(value),quote=False).replace('|','\\|').replace('\r','').replace('\n','<br>')
    lines=[]
    for block in blocks:
        if block[0]=='heading': lines.extend(['#'*block[1]+' '+html.escape(block[2],quote=False),''])
        elif block[0]=='paragraph': lines.extend([html.escape(block[1],quote=False),''])
        else:
            headers,rows=block[1:]; lines+=['| '+' | '.join(map(safe,headers))+' |','| '+' | '.join(['---']*len(headers))+' |']
            lines+=['| '+' | '.join(map(safe,row))+' |' for row in rows]; lines.append('')
    return '\n'.join(lines)


def _word(blocks,path):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document(); section=doc.sections[0]
    section.page_width=Cm(21); section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(1.8); section.left_margin=section.right_margin=Cm(1.7)
    for style_name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[style_name]; style.font.name='Arial Unicode MS'; style._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    footer=section.footer.paragraphs[0]; footer.alignment=2
    footer.add_run('自动估值建模报告 · ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    for block in blocks:
        if block[0]=='heading': doc.add_heading(block[2],level=0 if block[1]==1 else block[1]-1)
        elif block[0]=='paragraph': doc.add_paragraph(block[1])
        else:
            headers,rows=block[1:]; table=doc.add_table(rows=1,cols=len(headers)); table.style='Light Shading Accent 1'; table.autofit=False
            widths=[17.6/len(headers)]*len(headers)
            if len(headers)==2: widths=[4.1,13.5]
            for i,width in enumerate(widths): table.columns[i].width=Cm(width)
            for i,label in enumerate(headers):
                table.rows[0].cells[i].text=label
                for run in table.rows[0].cells[i].paragraphs[0].runs: run.bold=True
            repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in rows:
                cells=table.add_row().cells
                for i,value in enumerate(values): cells[i].text=value
            for row in table.rows:
                for i,cell in enumerate(row.cells):
                    cell.width=Cm(widths[i])
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(3)
                        for run in para.runs:
                            run.font.size=Pt(8.5); run.font.name='Arial Unicode MS'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
            doc.add_paragraph()
    doc.core_properties.title='自动估值建模报告'; doc.core_properties.subject='财报证据、经营驱动预测与三情景估值'
    doc.save(path)


def export_model(run: dict, fmt: str, outputdir: Path) -> Path:
    """Export without mutating or recomputing any canonical result values."""
    fmt=str(fmt).lower().lstrip('.')
    if fmt not in {'xlsx','json','md','docx'}: raise ValueError('不支持的估值报告格式')
    outputdir=Path(outputdir); outputdir.mkdir(parents=True,exist_ok=True)
    path=outputdir/('报告.'+fmt)
    if fmt=='xlsx': return export_valuation_workbook(run,path)
    if fmt=='json': path.write_text(json.dumps(run,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
    else:
        blocks=_blocks(run)
        if fmt=='md': path.write_text(_markdown(blocks),encoding='utf-8')
        else: _word(blocks,path)
    return path
