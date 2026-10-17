"""Export the saved review snapshot without inventing scores or recomputing findings."""
from __future__ import annotations

from decimal import Decimal
import html
import json
import re
from pathlib import Path


_LABELS = {
    'matched': '核验一致', 'verified': '已核验', 'mismatch': '数值不一致',
    'ambiguous': '口径不明确', 'unsupported': '缺少支持证据',
    'unresolved': '待核验', 'missing': '缺失', 'not_assessed': '未评估',
    'human_review': '需人工评估', 'assessed': '已评估', 'passed': '通过',
    'failed': '不通过', 'blocked': '交付受阻', 'hard': '严重',
    'medium': '需关注', 'soft': '提示', 'contradicted': '证据与主张矛盾',
    'unverifiable': '暂无法核验', 'human_review_required': '需人工评估',
    'unassessable': '无法评估', 'partial': '部分评估', 'pass': '通过',
    'fail': '不通过', 'critical': '严重', 'incomplete': '评估未完成',
}
_SCOPE = (
    '机器核验只覆盖本次输入、已识别主张及已提供证据；未识别的主张、业务因果、'
    '预测合理性与决策价值仍需人工复核。未匹配不直接等于事实错误，匹配也不代表整份报告正确。'
)
_REPRO = (
    '可复现范围：本报告、Excel 与 JSON 为同一次审查结果快照，不新增计算或事实。'
    '复核应使用相同源文件、运行编号、输入事实、证据与规则版本；仅凭导出 JSON '
    '不能重建未随文件提供的原始财报。修改源报告或 Excel 后须重新运行审查。'
)


def _json_default(value):
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise ValueError('审查结果含非有限数值')
        return str(value)
    if isinstance(value, Path):
        return str(value)
    raise TypeError(f'不能序列化 {type(value).__name__}')


def _text(value):
    if value is None or value == '':
        return '—'
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=_json_default, allow_nan=False)
    return str(value)


def _label(value):
    return _LABELS.get(str(value), _text(value))


def _sections(result):
    """Shared semantic sections keep every human readable export consistent."""
    claim_review = result.get('claim_review') or {}
    quality = result.get('quality_review') or {}
    coverage = quality.get('assessment_coverage', quality.get('coverage'))
    summary = [
        ['运行编号', _text(result.get('run_id'))], ['源文件', _text(result.get('source_file'))],
        ['审查摘要', _text(result.get('summary'))],
        ['交付状态', '交付受阻' if quality.get('delivery_blocked') is True else
         '未发现已记录的阻断项，仍需完成待人工评估项目' if quality.get('delivery_blocked') is False else '未提供交付判定'],
        ['完整总分', _text(quality.get('overall_total')) if quality.get('overall_total') is not None else '未评估，不生成完整总分'],
        ['已评估权重', _text(quality.get('assessed_weight'))],
        ['评估覆盖率', _text(coverage)], ['暂定得分（仅已评估范围）', _text(quality.get('provisional_score'))],
        ['主张核验统计', _text(claim_review.get('stats'))],
        ['质量规则版本', _text(quality.get('rules_version'))],
        ['源报告内容指纹', _text(quality.get('report_hash'))],
        ['评分量表与适用范围', _text(quality.get('rubric'))],
    ]
    claims = []
    for claim in claim_review.get('claims') or []:
        claims.append([_text(claim.get(key)) if key != 'status' else _label(claim.get(key)) for key in
                       ('id', 'line', 'raw', 'code', 'period', 'scope', 'value', 'unit', 'status',
                        'expected_value', 'evidence_ids', 'reason', 'advice')])
    dimensions = [[_text(d.get('name')), _text(d.get('score')), _text(d.get('max_score')),
                   _label(d.get('status')), _text(d.get('basis')),
                   _text(d.get('assessed_weight')), _text(d.get('coverage')),
                   _text(d.get('human_review')), _text(d.get('checks'))]
                  for d in quality.get('dimensions') or []]
    issues = [[_label(i.get('severity')), _text(i.get('type')), _text(i.get('line')),
               _text(i.get('detail')), _text(i.get('context')), _text(i.get('advice'))]
              for i in result.get('issues') or []]
    limitations = quality.get('limitations') or []
    if isinstance(limitations, str):
        limitations = [limitations]
    notes = [['机器与人工边界', _SCOPE], ['可复现范围', _REPRO]]
    notes.extend(['本次评估限制', _text(item)] for item in limitations)
    if result.get('note'):
        notes.append(['规则说明', _text(result['note'])])
    formula_rows = [[_text(c.get('line')), _text(c.get('expression')), _label(c.get('status')), _text(c.get('calculated_value')), _text(c.get('stated_value')), _text(c.get('basis') or c.get('reason'))] for c in (result.get('formula_review') or {}).get('checks', [])]
    return [
        ('公式复算记录', ['源报告行号', '表达式', '结果', '复算值', '报告值', '说明'], formula_rows),
        ('审查概览', ['项目', '结果'], summary),
        ('主张核验明细', ['主张编号', '源报告行号', '原文', '科目', '期间', '口径', '报告数值', '单位',
                       '核验状态', '核验依据数值', '证据编号', '原因', '修订建议'], claims),
        ('独立质量评估', ['维度', '暂定得分', '满分', '评估状态', '评分依据',
                      '已评估权重', '维度覆盖率', '人工复核要求', '机器检查明细'], dimensions),
        ('质量阻断与风险提示', ['规则', '级别', '原因'],
         [[_text(f.get('code')), _label(f.get('severity')), _text(f.get('reason'))]
          for f in quality.get('flags') or []]),
        ('其他规则发现', ['级别', '问题类型', '源报告行号', '问题详情', '上下文', '修订建议'], issues),
        ('评估限制与复现说明', ['项目', '说明'], notes),
    ]


def _markdown(sections):
    def safe(value):
        return html.escape(str(value), quote=False).replace('|', '\\|').replace('\r', '').replace('\n', '<br>')
    lines = ['# 报告审查结果', '']
    for title, headers, rows in sections:
        lines.extend(['## ' + title, ''])
        if not rows:
            lines.extend(['本次未提供该项记录；不据此判定通过。', ''])
            continue
        lines.append('| ' + ' | '.join(map(safe, headers)) + ' |')
        lines.append('| ' + ' | '.join(['---'] * len(headers)) + ' |')
        lines.extend('| ' + ' | '.join(map(safe, row)) + ' |' for row in rows)
        lines.append('')
    return '\n'.join(lines)


def _office_text(value):
    # Office XML cannot represent these controls; JSON retains the source.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', str(value))


def _word(sections, path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    doc.styles['Normal'].font.size = Pt(10)
    doc.add_heading('报告审查结果', 0)
    for title, headers, rows in sections:
        rows = [[_office_text(value) for value in row] for row in rows]
        doc.add_heading(title, 1)
        if not rows:
            doc.add_paragraph('本次未提供该项记录；不据此判定通过。')
        elif title in ('主张核验明细', '独立质量评估'):
            # Vertical claim cards remain readable on an A4 page.
            for row in rows:
                heading = '主张 ' + row[0] + ' · 源报告第 ' + row[1] + ' 行' if title == '主张核验明细' else row[0]
                doc.add_heading(heading, 2)
                table = doc.add_table(rows=0, cols=2)
                table.style = 'Light Shading Accent 1'
                table.columns[0].width, table.columns[1].width = Cm(3.5), Cm(13.9)
                for label, value in zip(headers, row):
                    cells = table.add_row().cells
                    cells[0].text, cells[1].text = label, value
        else:
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Light Shading Accent 1'
            for cell, text in zip(table.rows[0].cells, headers):
                cell.text = text
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
    doc.core_properties.title = '报告审查结果'
    doc.core_properties.subject = '主张核验、独立质量评估与修订依据'
    doc.save(path)


def _excel(sections, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    for title, headers, rows in sections:
        sheet = book.create_sheet(title)
        for row_index, row in enumerate([headers] + rows, start=1):
            for column, value in enumerate(row, start=1):
                # Force all snapshot values to string cells, even =,+,-,@ or leading
                # whitespace/control characters. No source text can become a formula.
                cell = sheet.cell(row_index, column)
                cell.value = _office_text(value)
                cell.data_type = 's'
                cell.font = Font(name='Arial', size=10, bold=row_index == 1,
                                 color='FFFFFF' if row_index == 1 else '172B4D')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if row_index == 1:
                    cell.fill = PatternFill('solid', fgColor='24476B')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for i, header in enumerate(headers, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = 48 if header in (
                '原文', '原因', '修订建议', '评分依据', '说明', '结果', '问题详情', '上下文') else 20
    book.save(path)
    book.close()


def export_review(directory: Path, result: dict) -> dict:
    """Write four compatible artifacts and return their relative file names.

    Decimal values are serialized as exact strings. Scores and findings are kept
    as supplied; absent scores are never replaced with zero or fabricated totals.
    """
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    filenames = {'markdown': '报告.md', 'docx': '报告.docx', 'xlsx': '报告.xlsx', 'json': '报告.json'}
    payload = json.dumps(result, ensure_ascii=False, indent=2, default=_json_default, allow_nan=False)
    sections = _sections(result)
    (directory / filenames['json']).write_text(payload, encoding='utf-8')
    (directory / filenames['markdown']).write_text(_markdown(sections), encoding='utf-8')
    _word(sections, directory / filenames['docx'])
    _excel(sections, directory / filenames['xlsx'])
    return filenames
