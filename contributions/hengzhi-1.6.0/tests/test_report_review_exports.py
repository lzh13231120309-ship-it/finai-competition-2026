"""Review artifacts preserve findings, unknown scores and literal spreadsheet text."""
import copy
from decimal import Decimal
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from server.report_review_exports import export_review


def fixture():
    return {
        'run_id': 'review-123', 'source_file': '公司报告.md', 'summary': '营业收入不一致；质量评分尚待人工补齐。',
        'claim_review': {'claims': [
            {'id': 'C001', 'line': 12, 'raw': '2024年营业收入为9亿元', 'code': 'revenue',
             'period': '2024', 'scope': '合并', 'value': Decimal('900000000'), 'unit': '元',
             'status': 'mismatch', 'expected_value': Decimal('800000000.123456789'),
             'evidence_ids': ['abcdef12:p42:b1'], 'reason': '同科目同期间数值不一致',
             'advice': '将第12行收入修订为已核验财报值，并保留页码证据。'}],
            'stats': {'mismatch': 1}},
        'quality_review': {'dimensions': [
            {'name': '事实正确', 'score': Decimal('15.5'), 'max_score': 25,
             'status': 'assessed', 'basis': '按已识别主张核验'},
            {'name': '金融逻辑', 'score': None, 'max_score': 25,
             'status': 'human_review', 'basis': '需人工判断业务因果'}],
            'overall_total': None, 'assessed_weight': 25,
            'assessment_coverage': Decimal('0.25'), 'provisional_score': Decimal('15.5'),
            'delivery_blocked': True, 'limitations': ['待人工评估维度不得按满分或零分代入。']},
        'issues': [{'severity': 'hard', 'type': '计数不符', 'line': 28,
                    'detail': '报告写2项，工具为3项', 'context': '待复核2项', 'advice': '改为3项'}],
    }


class ReviewExportTests(unittest.TestCase):
    def test_paths_archives_precision_and_actionable_findings(self):
        result = fixture()
        original = copy.deepcopy(result)
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp) / 'review-output'
            files = export_review(directory, result)
            self.assertEqual(files, {'markdown': '报告.md', 'docx': '报告.docx', 'xlsx': '报告.xlsx', 'json': '报告.json'})
            self.assertEqual(result, original)
            for filename in files.values():
                path = directory / filename
                self.assertTrue(path.is_file())
                self.assertEqual(path.parent, directory)
                if path.suffix in ('.docx', '.xlsx'):
                    with ZipFile(path) as archive:
                        self.assertIsNone(archive.testzip())
            payload = json.loads((directory / files['json']).read_text())
            self.assertEqual(payload['claim_review']['claims'][0]['expected_value'], '800000000.123456789')
            self.assertIsNone(payload['quality_review']['overall_total'])
            md = (directory / files['markdown']).read_text()
            document = Document(directory / files['docx'])
            word = '\n'.join(p.text for p in document.paragraphs) + '\n' + '\n'.join(
                c.text for table in document.tables for row in table.rows for c in row.cells)
            for text in (md, word):
                for required in ('review-123', '公司报告.md', '2024年营业收入为9亿元', '800000000.123456789',
                                 'abcdef12:p42:b1', '第12行', '独立质量评估', '评估覆盖率', '0.25',
                                 '需人工评估', '未评估，不生成完整总分', '可复现范围', '不能重建', '交付受阻', '计数不符'):
                    self.assertIn(required, text)
            book = load_workbook(directory / files['xlsx'], data_only=True)
            self.assertEqual(book['主张核验明细']['B2'].value, '12')
            self.assertEqual(book['主张核验明细']['J2'].value, '800000000.123456789')
            self.assertEqual(book['独立质量评估']['B3'].value, '—')
            book.close()

    def test_formula_injection_in_all_untrusted_fields_is_literal(self):
        for attack in ('=HYPERLINK("https://evil.example","open")', '+1+1', '-1+1', '@SUM(A1)',
                       '\t=1+1', '\n=1+1', '\x00=1+1'):
            result = fixture()
            result['run_id'] = result['source_file'] = result['summary'] = attack
            claim = result['claim_review']['claims'][0]
            for key in claim:
                claim[key] = attack
            for key in result['quality_review']['dimensions'][0]:
                result['quality_review']['dimensions'][0][key] = attack
            result['quality_review']['limitations'] = [attack]
            for key in result['issues'][0]:
                result['issues'][0][key] = attack
            with tempfile.TemporaryDirectory() as tmp:
                files = export_review(Path(tmp), result)
                book = load_workbook(Path(tmp) / files['xlsx'], data_only=False)
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            self.assertNotEqual(cell.data_type, 'f', cell.coordinate)
                            self.assertIsNone(cell.hyperlink)
                self.assertEqual(book['主张核验明细']['C2'].value, attack.replace('\x00', ''))
                book.close()
                with ZipFile(Path(tmp) / files['xlsx']) as archive:
                    self.assertFalse(any('externalLink' in name for name in archive.namelist()))

    def test_empty_review_has_no_fabricated_total_or_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            files = export_review(Path(tmp), {'run_id': 'empty'})
            md = (Path(tmp) / files['markdown']).read_text()
            self.assertIn('未提供交付判定', md)
            self.assertIn('未评估，不生成完整总分', md)
            self.assertIn('不据此判定通过', md)
            self.assertNotIn('100', md)
            self.assertNotIn('完整总分 | 0', md)

    def test_markdown_escapes_source_html_and_table_separators(self):
        result = fixture()
        result['claim_review']['claims'][0]['raw'] = '<script>bad()</script>|a\nb'
        with tempfile.TemporaryDirectory() as tmp:
            files = export_review(Path(tmp), result)
            md = (Path(tmp) / files['markdown']).read_text()
            self.assertNotIn('<script>', md)
            self.assertIn('&lt;script&gt;bad()&lt;/script&gt;\\|a<br>b', md)


if __name__ == '__main__':
    unittest.main()
