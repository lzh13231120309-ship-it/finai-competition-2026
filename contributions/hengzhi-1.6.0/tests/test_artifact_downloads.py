"""Download bytes, filenames and UTF-8 text remain usable outside the app."""
import asyncio,io,json
import httpx
from openpyxl import load_workbook
from server import finance_service as fs
from test_agent_longrun import env


def test_exports_download_as_valid_files_with_explicit_utf8(env):
    from server.app import app
    run=fs.save_run(env,{'kind':'analysis','company':'演示公司','facts':[{'fact_id':'demo','code':'revenue','raw_label':'营业收入','company':'演示公司','period':'2025-12-31','currency':'CNY','scope':'consolidated','value':'1200','unit':'元','status':'verified','evidence_ids':[]}],'metrics':[],'checks':[],'review_queue':[],'warnings':[]})
    async def check():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
            for fmt in ('xlsx','docx','md','json'):
                response=await client.get(f'/api/financial/artifacts/{env}/{run["run_id"]}/报告.{fmt}')
                assert response.status_code==200
                assert 'attachment' in response.headers['content-disposition']
                if fmt=='xlsx':
                    book=load_workbook(io.BytesIO(response.content));assert book.sheetnames;book.close()
                elif fmt=='docx':assert response.content[:2]==b'PK'
                else:
                    assert 'charset=utf-8' in response.headers['content-type']
                    assert '营业收入' in response.content.decode('utf-8')
                    if fmt=='json':assert json.loads(response.content)['company']=='演示公司'
            assert (await client.get(f'/api/financial/artifacts/{env}/{run["run_id"]}/private.txt')).status_code==404
    asyncio.run(check())
