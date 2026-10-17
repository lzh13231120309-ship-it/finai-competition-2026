"""Regression tests for complete, source-addressable financial documents."""
import hashlib
import importlib
import json
import platform
import tempfile
import unicodedata
import unittest
from pathlib import Path


class FinancialDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        try:
            self.api = importlib.import_module('server.financial_documents')
        except ModuleNotFoundError:
            self.fail('Structured document extraction is not implemented')

    def test_text_is_complete_and_paged_reader_reports_remaining_content(self):
        text = '营业收入 -1,234.50 元\n' * 6000
        source = self.root / '财务.txt'
        source.write_text(text, encoding='utf-8')
        doc = self.api.parse_document(source, self.root / 'cache')
        self.assertEqual(doc['text'], text)
        self.assertEqual(doc['sha256'], hashlib.sha256(source.read_bytes()).hexdigest())
        chunk = self.api.read_document(doc, offset=len(text)-80, limit=40)
        self.assertEqual(chunk['text'], text[-80:-40])
        self.assertTrue(chunk['has_more'])
        self.assertEqual(chunk['total_characters'], len(text))
        self.assertTrue(chunk['evidence'][0]['evidence_id'])
        json.dumps(doc, ensure_ascii=False, allow_nan=False)

    def test_cache_does_not_mix_identical_contents_under_different_names(self):
        a = self.root / 'a.txt'; b = self.root / 'b.txt'
        a.write_text('相同证据'); b.write_text('相同证据')
        first = self.api.parse_document(a, self.root / 'cache')
        cached = self.api.parse_document(b, self.root / 'cache')
        self.assertEqual(first['document_id'], cached['document_id'])
        self.assertEqual(cached['filename'], 'b.txt')
        self.assertEqual(cached['text'], '相同证据')
        a.write_text('新证据')
        self.assertNotEqual(self.api.parse_document(a, self.root / 'cache')['document_id'], first['document_id'])

    def test_csv_keeps_all_rows_columns_signs_and_coordinates(self):
        source = self.root / 'statement.csv'
        lines = ['项目,' + ','.join('列'+str(i) for i in range(50))]
        lines += ['收入'+str(i)+','+','.join(['(-1,234.50)']*50).replace('(-1,234.50)', '"(-1,234.50)"') for i in range(350)]
        source.write_text('\n'.join(lines), encoding='utf-8-sig')
        table = self.api.parse_document(source)['tables'][0]
        self.assertEqual(len(table['rows']), 351)
        self.assertEqual(len(table['rows'][-1]), 51)
        self.assertEqual(table['rows'][-1][-1], '(-1,234.50)')
        self.assertEqual(table['cell_coordinates'][-1][-1], 'AY351')
        doc = self.api.parse_document(source)
        read = self.api.read_document(doc, table_id=table['table_id'], offset=0, limit=100)
        self.assertEqual(read['tables'][0]['rows'][0][0], '项目')
        self.assertLess(len(read['tables'][0]['rows']), 351)
        self.assertTrue(read['has_more'])

    def test_xlsx_preserves_formulas_without_inventing_uncached_results(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = '利润表'
        ws.append(['项目', '2025', '2024']); ws.append(['收入', -1234.5, 900])
        ws['B3'] = '=B2*2'; ws['C3'] = 0; ws['AY351'] = '最后单元格'
        source = self.root / 'statement.xlsx'; wb.save(source)
        doc = self.api.parse_document(source)
        table = doc['tables'][0]
        self.assertEqual(table['rows'][350][50], '最后单元格')
        self.assertEqual(table['formulas']['B3'], '=B2*2')
        self.assertEqual(table['rows'][2][1], '=B2*2')
        self.assertTrue(any('B3' in w and '缓存' in w for w in doc['warnings']))
        self.assertEqual(table['cell_coordinates'][1][1], 'B2')
        block_by_id = {b['evidence_id']: b for p in doc['pages'] for b in p['blocks']}
        cell = block_by_id[table['cell_evidence_ids'][1][1]]
        self.assertEqual(cell['text'], '-1234.5')
        self.assertEqual(cell['coordinate'], 'B2')
        self.assertEqual(cell['sheet'], '利润表')

    def test_excel_percentage_format_and_formula_errors_are_explicit(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(['指标', '值']); ws.append(['增长率', 0.125])
        ws['B2'].number_format = '0.0%'; ws['B3'] = '#DIV/0!'
        source = self.root / 'formats.xlsx'; wb.save(source)
        doc = self.api.parse_document(source)
        self.assertEqual(doc['tables'][0]['rows'][1][1], '12.5%')
        self.assertEqual(doc['tables'][0]['number_formats'][1][1], '0.0%')
        self.assertTrue(any('B3' in w and '#DIV/0!' in w for w in doc['warnings']))

    def test_borderless_pdf_right_aligned_financial_columns(self):
        import pymupdf as fitz
        pdf = fitz.open(); page = pdf.new_page(width=600,height=600)
        for r,row in enumerate([['Item','2025','2024'],['Revenue','1,234,567.00','8.00'],['Profit','(100.00)','50.00'],['Cashflow','250.00','2.00']]):
            y = 50+r*30
            page.insert_text((30,y+12),row[0],fontsize=11)
            for c in (1,2): page.insert_textbox(fitz.Rect(150*c,y,150*c+140,y+25),row[c],fontsize=11,align=2)
        source = self.root / 'right-aligned.pdf'; pdf.save(source); pdf.close()
        doc = self.api.parse_document(source)
        self.assertTrue(any(t['rows'][1] == ['Revenue','1,234,567.00','8.00'] for t in doc['tables']))
        self.assertTrue(any(t.get('review_required') for t in doc['tables']))

    def test_docx_order_and_tables_remain_auditable(self):
        from docx import Document
        d = Document(); d.add_paragraph('开始')
        t = d.add_table(rows=2, cols=2)
        t.cell(0,0).text = '收入'; t.cell(0,1).text = '2025'
        t.cell(1,0).text = '净利润'; t.cell(1,1).text = '(100)'
        d.add_paragraph('结束')
        source = self.root / 'statement.docx'; d.save(source)
        doc = self.api.parse_document(source)
        self.assertLess(doc['text'].index('开始'), doc['text'].index('净利润'))
        self.assertLess(doc['text'].index('净利润'), doc['text'].index('结束'))
        self.assertEqual(doc['tables'][0]['rows'][1][1], '(100)')
        self.assertTrue(any('页码' in w for w in doc['warnings']))

    def test_query_reader_returns_source_evidence(self):
        source = self.root / 'notes.txt'; source.write_text('收入增加\n净利润下降 -50%\n现金流稳定', encoding='utf-8')
        doc = self.api.parse_document(source)
        found = self.api.read_document(doc, query='净利润', limit=20)
        self.assertIn('净利润', found['text'])
        self.assertTrue(found['evidence'])
        self.assertEqual(found['document_id'], doc['document_id'])
        with self.assertRaises(ValueError):
            self.api.read_document(doc, table_id='missing')
        with self.assertRaises(ValueError):
            self.api.read_document(doc, page=9)

    def test_pdf_native_text_keeps_all_pages_and_numeric_tables(self):
        import pymupdf as fitz
        source = self.root / 'native.pdf'; pdf = fitz.open()
        page = pdf.new_page(width=500, height=700)
        for r, row in enumerate([['Item','2025','2024'],['Revenue','-1,234.50','900.00'],['Profit','(100.00)','50.00']]):
            for c, cell in enumerate(row):
                page.insert_text((40+c*150, 70+r*30), cell)
        for y in [50,80,110,140]: page.draw_line((30,y),(480,y))
        for x in [30,180,330,480]: page.draw_line((x,50),(x,140))
        second = pdf.new_page(); second.insert_text((50,50), 'last-page-evidence')
        pdf.save(source); pdf.close()
        doc = self.api.parse_document(source)
        self.assertEqual(doc['page_count'], 2)
        self.assertEqual(doc['pages'][0]['method'], 'text')
        self.assertIn('last-page-evidence', doc['text'])
        self.assertTrue(any('-1,234.50' in str(t['rows']) and '(100.00)' in str(t['rows']) for t in doc['tables']))
        self.assertTrue(all(block['bbox'] for block in doc['pages'][0]['blocks']))

    def test_pdf_table_context_excludes_distant_statement_heading(self):
        import pymupdf as fitz
        pdf = fitz.open(); page = pdf.new_page(width=500,height=700)
        page.insert_text((40,50),'Parent company statement')
        page.insert_text((40,280),'Consolidated statement Unit: CNY')
        for r,row in enumerate([['Item','2025','2024'],['Revenue','100','90'],['Profit','10','9']]):
            for c,cell in enumerate(row): page.insert_text((40+c*150,320+r*30),cell)
        for y in [300,330,360,390]: page.draw_line((30,y),(480,y))
        for x in [30,180,330,480]: page.draw_line((x,300),(x,390))
        source = self.root / 'context.pdf'; pdf.save(source); pdf.close()
        table = self.api.parse_document(source)['tables'][0]
        self.assertIn('Consolidated statement',table['context'])
        self.assertNotIn('Parent company',table['context'])

    @unittest.skipUnless(platform.system() == 'Darwin', 'Native Apple Vision OCR requires macOS')
    def test_chinese_image_and_mixed_pdf_use_offline_ocr(self):
        import pymupdf as fitz
        from PIL import Image, ImageDraw, ImageFont
        font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Songti.ttc', 46)
        img = Image.new('RGB',(1600,800),'white'); draw = ImageDraw.Draw(img)
        draw.text((70,50),'财务报表 2025 年',font=font,fill='black')
        rows = [('项目','2025','2024'),('营业收入','-1,234.50','900.00'),('净利润','(100.00)','50.00'),('现金流','250.00','200.00')]
        for r,row in enumerate(rows):
            for c,cell in enumerate(row): draw.text((70+c*510,180+r*110),cell,font=font,fill='black')
        source = self.root / 'scan.png'; img.save(source)
        doc = self.api.parse_document(source)
        self.assertEqual(doc['pages'][0]['method'], 'ocr')
        self.assertIn('营业收入', doc['text'])
        self.assertIn('1,234.50', doc['text'])
        self.assertTrue(any('-1,234.50' in str(t['rows']) for t in doc['tables']))
        self.assertTrue(any('(100.00)' in unicodedata.normalize('NFKC', str(t['rows'])) for t in doc['tables']))
        self.assertTrue(all(b['bbox'] and b['confidence'] > 0 for b in doc['pages'][0]['blocks']))
        pdf = fitz.open(); p = pdf.new_page(width=800,height=400)
        p.insert_image(p.rect,filename=str(source)); p.insert_text((10,390),'stamp')
        p = pdf.new_page(); p.insert_text((50,50),'digital page')
        mixed = self.root / 'mixed.pdf'; pdf.save(mixed); pdf.close()
        result = self.api.parse_document(mixed)
        self.assertEqual([p['method'] for p in result['pages']], ['ocr','text'])
        self.assertIn('营业收入', result['text'])
        self.assertIn('digital page', result['text'])

    def test_unsupported_and_resource_limits_fail_explicitly(self):
        source = self.root / 'unknown.binary'; source.write_bytes(b'\x00\xff\x00')
        with self.assertRaises(ValueError): self.api.parse_document(source)
        source = self.root / 'huge.txt'; source.write_text('too large')
        from unittest.mock import patch
        with patch.object(self.api, 'MAX_FILE_BYTES', 2):
            with self.assertRaisesRegex(ValueError, '大小'): self.api.parse_document(source)

if __name__ == '__main__':
    unittest.main()
