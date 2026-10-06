"""Behaviour checks for disclosed estimation and readable report delivery."""
import json,sys
from pathlib import Path
from decimal import Decimal as D
import pytest
ROOT=Path(__file__).resolve().parents[1];sys.path[:0]=[str(ROOT),str(ROOT.parent/'第一阶段_金融提取')]
from server import estimation, research_engine as e,presentation,report_documents,parsers
from test_research_engine import record,assumptions,facts

def inputs():return facts()+[record('revenue','2023',960),record('revenue_cost','2024',720)]

def test_no_json_produces_history_and_partial_enterprise():
    a,audit=estimation.prepare(inputs());r=e.valuation(inputs(),a,allow_partial=True)
    assert D(a['growth_rates'][0])==D('.2') # observed 25%, deliberately capped with explicit basis
    item=next(x for x in audit['parameters'] if x['parameter']=='growth_rates')
    assert item['origin']=='historical_estimate' and '限幅' in item['basis']
    assert r['dcf']['enterprise_value'] is not None and r['dcf']['equity_value'] is None and r['dcf']['price_per_share'] is None
    assert all(x['price_per_share'] is None for x in r['dcf']['sensitivity'])
    assert {x['field'] for x in audit['missing']}=={'net_debt','shares_outstanding','comparables'}

def test_known_inputs_are_never_overwritten():
    a=assumptions();result,audit=estimation.prepare(inputs(),a)
    for k in ('growth_rates','ebit_margin','wacc','terminal_growth','shares_outstanding','net_debt'):assert result[k]==a[k]
    assert all(x['origin']=='user_supplied' for x in audit['parameters'] if x['parameter']=='wacc')

@pytest.mark.parametrize('bad',[{'wacc':.01,'terminal_growth':.02},{'shares_outstanding':-1},{'growth_rates':[None,.1]},{'tax_rate':25},{'ebit_margin':2}])
def test_invalid_explicit_input_is_not_estimated_away(bad):
    a=assumptions();a.update(bad)
    with pytest.raises(ValueError):estimation.prepare(inputs(),a)

def test_missing_growth_is_disclosed_scenario_not_factual_prediction():
    a,audit=estimation.prepare(facts());x=next(x for x in audit['parameters'] if x['parameter']=='growth_rates')
    assert x['origin']=='scenario_assumption' and '不是市场预测' in x['basis'] and x['sources']==[]

def test_observed_shares_and_debt_cash_bridge_preserve_units():
    rows=inputs()+[record('shares_outstanding','2024-year-end',100,base_unit='股'),record('interest_bearing_debt_total','2024-year-end',30),record('cash','2024-year-end',20)]
    a,audit=estimation.prepare(rows);r=e.valuation(rows,a,allow_partial=True)
    assert a['net_debt']=='10' and a['shares_outstanding']=='100'
    assert D(r['dcf']['price_per_share'])==pytest.approx((D(r['dcf']['enterprise_value'])-10)/100,abs=D('.000001'))
    assert next(x for x in audit['parameters'] if x['parameter']=='net_debt')['origin']=='historical_estimate'

def test_peer_business_date_currency_filter():
    a=assumptions();a.update(industry='electronics',business_tags=['connectors'])
    a['comparables']=[{'name':'Good','industry':'electronics','business_tags':['connectors'],'currency':'CNY','pe':10,'pb':1,'as_of':a['valuation_date'],'basis':'source'},
      {'name':'Bank','industry':'bank','pe':10,'pb':1,'as_of':a['valuation_date'],'basis':'source'},
      {'name':'Old','industry':'electronics','pe':10,'pb':1,'as_of':'2024-01-01','basis':'source'},
      {'name':'WrongBusiness','industry':'electronics','business_tags':['food'],'pe':10,'pb':1,'as_of':a['valuation_date'],'basis':'source'}]
    prepared,audit=estimation.prepare(inputs(),a)
    assert [x['name'] for x in prepared['comparables']]==['Good']
    assert len([x for x in audit['peer_screening'] if not x['included']])==3

def test_peer_estimated_growth_and_computed_multiples():
    a=assumptions();a.update(industry='electronics',growth_rates=None)
    a['comparables']=[{'name':name,'industry':'electronics','as_of':a['valuation_date'],'basis':'教学同业报价','market_cap':cap,'unit':'元','currency':'CNY'} for name,cap in [('B',1200),('C',3600)]]
    rows=facts()
    for company,income in [('B',1100),('C',1300)]:
        rows += [record('revenue','2023',1000,issuer=company),record('revenue','2024',income,issuer=company),record('net_profit_parent','2024',100,issuer=company),record('equity_parent','2024-year-end',1000,issuer=company)]
    prepared,audit=estimation.prepare(rows,a)
    assert D(prepared['growth_rates'][0])==D('.2')
    assert next(x for x in audit['parameters'] if x['parameter']=='growth_rates')['origin']=='peer_estimate'
    assert [D(x['pe']) for x in prepared['comparables']]==[D(12),D(36)]

def test_no_cross_basis_parameter_estimation():
    rows=inputs()+[record('revenue','2022',100,basis='调整前')]
    a,audit=estimation.prepare(rows)
    assert D(a['growth_rates'][0])==D('.2')

def test_short_chat_combines_figures_interpretation_and_word_link():
    r=e.analysis(inputs());text=presentation.chat(r,{'report.docx':'/api/report.docx','result.json':'/api/result.json'})
    assert '25.00%' in text and '40.00%' in text and '盈利空间' in text and 'Word' in text
    assert '工具流程' not in text and 'revenue' not in text and len(text)<1800

def test_word_contains_hypotheses_sources_and_formulas(tmp_path):
    from docx import Document
    a,audit=estimation.prepare(inputs());r=e.valuation(inputs(),a,allow_partial=True);r['estimation_audit']=audit
    p=tmp_path/'report.docx';report_documents.write_research(r,p);doc=Document(p)
    text='\n'.join(x.text for x in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert 'FCFF' in text and '探索情景' in text and '每股' in text and '总股数' in text

@pytest.mark.parametrize('encoding',['utf-8-sig','utf-16','gb18030'])
def test_text_encodings_do_not_generate_mojibake(tmp_path,encoding):
    p=tmp_path/'input.txt';p.write_text('营业收入与现金流，测试文字。',encoding=encoding)
    assert parsers._read_text_file(p)=='营业收入与现金流，测试文字。'
