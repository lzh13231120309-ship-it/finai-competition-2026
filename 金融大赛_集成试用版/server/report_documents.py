"""Offline Word exports from financial artifacts; no desktop application needed by users."""
from decimal import Decimal as D
import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .presentation import LABELS,ORIGINS,num,human,readable_excerpt,period_label,analysis_paragraphs,valuation_paragraphs

def document(title):
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.8)
    for name,size in [('Normal',10.5),('Title',23),('Heading 1',15),('Heading 2',12)]:
        s=doc.styles[name];s.font.name='Microsoft YaHei';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
        s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.2
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.add_paragraph(title,'Title');doc.add_paragraph('衡知金融研究助手  ·  '+datetime.datetime.now().strftime('%Y年%m月%d日'))
    doc.core_properties.title=title;doc.core_properties.author='衡知';doc.core_properties.subject='事实、分析与预测假设分开记录'
    foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('衡知  ·  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    return doc

def table(doc,headers,rows,widths=None):
    if not rows:doc.add_paragraph('现有材料不足，未形成该项结果。');return
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    total=6.65
    for i,(c,text) in enumerate(zip(t.rows[0].cells,headers)):
        c.text=text;c.width=Inches(widths[i] if widths else total/len(headers))
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EFF1');c._tc.get_or_add_tcPr().append(sh)
        for run in c.paragraphs[0].runs:run.bold=True
        c.paragraphs[0].paragraph_format.keep_with_next=True
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,(c,text) in enumerate(zip(cells,row)):
            c.text=str(text if text is not None else '未计算');c.width=Inches(widths[i] if widths else total/len(headers))
            if ri%2:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F6F8F8');c._tc.get_or_add_tcPr().append(sh)
            for p in c.paragraphs:p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
    props=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
    props.append(borders)
    for row in t.rows:
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:r.font.size=Pt(9)
    doc.add_paragraph('')

def percent(v):return num(D(str(v))*100)+'%' if v is not None else '未提供'

def write_extraction(result,rows,validation,path):
    doc=document('公告解读与关键事项');doc.add_heading('材料概览',1)
    doc.add_paragraph('文件：'+result['source']['filename']+f"。共整理{len(rows)}项字段，{len(validation['issues'])}项带有复核提示。原文未出现的项目保持缺失，不作为零值。")
    events=result.get('events',[])
    doc.add_paragraph(f"识别到{len(events)}个可关联事项。中标金额不直接等于已确认收入；质押、计划与权益变动应结合主体、日期、数量及原公告核验。")
    doc.add_heading('关键字段与来源',1)
    from finance_extract.extract import METRICS
    labels={v:k for k,v in METRICS.items()};labels.update({'pledged_shares':'本次质押股数','shareholder':'股东','pledgee':'质权人','award_amount':'中标金额','project_name':'项目名称','capital_ratio':'总股本占比','holder_ratio':'持股占比','pledge_start':'质押起始日','pledge_end':'质押到期日'})
    table(doc,['字段','主体/期间','原值','单位','来源'],[[labels.get(f['name'],f['name']),f.get('subject') or f.get('period') or '',f.get('raw') or f.get('value'),f.get('unit') or '',('第'+str(f['evidence'].get('row'))+'行') if f.get('source_location_kind')=='csv_row' else '第'+str(f['evidence'].get('page'))+'页'] for f in rows[:500]])
    if len(rows)>500:doc.add_paragraph('报告展示前500项；完整字段保存在CSV及JSON中。')
    doc.add_heading('需要复核的项目',1)
    for issue in validation['issues'][:100]:doc.add_paragraph(labels.get(issue['name'],issue['name'])+'；原文第'+str(issue['page'])+'页；'+human('、'.join(issue['reasons'])))
    doc.add_heading('材料与依据',1);doc.add_paragraph('SHA256：'+result['source']['sha256']);doc.add_paragraph('字段与引用来自本任务材料，完整原文定位和复核标记保存在JSON。自动格式核验不等于所有语义已人工确认。');doc.save(path)

def write_research(data,path):
    analysis=data['schema']=='finai.analysis.v1';title='财务报告分析' if analysis else '估值分析与假设说明'
    doc=document(title);doc.add_heading('核心判断',1)
    for overview in data.get('financial_overview',[]):doc.add_paragraph(overview)
    for paragraph in analysis_paragraphs(data) if analysis else valuation_paragraphs(data):doc.add_paragraph(paragraph)
    if analysis:
        doc.add_heading('经营指标与可比变化',1)
        rows=[]
        for c in data.get('changes',[]):
            rate=num(c['percent_change'])+'%' if c['percent_change'] is not None else '非正基数，不计算常规增长率'
            rows.append([c['issuer'],LABELS.get(c['metric'],c['metric']),period_label(c['period']),period_label(c['comparison_period']),c.get('basis','未明确'),('同比' if c['mode']=='yoy' else '环比')+' '+rate])
        table(doc,['公司','指标','期间','比较期','会计口径','变化'],rows)
        doc.add_paragraph('增长率 =（本期－比较期）/比较期。累计差额只用于流量指标的单季还原，期末资产负债余额不相减生成单季指标。完整口径、原值、绝对变化及来源字段随JSON交付。')
        doc.add_heading('盈利质量与财务结构',1)
        table(doc,['公司','期间','会计口径','指标','数值'],[[r['issuer'],period_label(r['period']),r.get('basis','未明确'),LABELS.get(r['metric'],r['metric']),num(r['value'])+('%' if r['metric'] in ('gross_margin','net_margin_parent','debt_asset_ratio') else '倍')] for r in data.get('ratios',[])])
        for ratio in data.get('ratios',[]):
            if ratio['metric']=='cash_profit_ratio':doc.add_paragraph(ratio['issuer']+' '+ratio['period']+'（'+ratio.get('basis','口径未明确')+'）现金利润比的分母为'+LABELS.get(ratio.get('denominator_metric','net_profit_parent'),'利润')+'；只有归母利润时属于统计范围有差异的代理比较。')
        for s in data.get('signals',[]):doc.add_paragraph(s['issuer']+' '+s['period']+'：'+s['text'])
        if data.get('profit_bridges'):
            doc.add_heading('毛利变化的定量桥接',1)
            for b in data['profit_bridges']:
                doc.add_paragraph(b['issuer']+' '+b['period']+'相对'+b['comparison_period']+'；金额单位为'+b['currency']+'。'+b['formula'])
                table(doc,['项目','金额'],[['毛利变化',num(b['gross_profit_change'])],['收入变化贡献',num(b['revenue_effect'])],['毛利率变化贡献',num(b['margin_effect'])]])
                doc.add_paragraph(b['interpretation']+'；加总勾稽：'+('通过' if b['reconciled'] else '需要复核'))
        if data.get('turnover'):
            doc.add_heading('应收与存货周转核查',1)
            table(doc,['公司','期间','指标','天数或缺失原因'],[[r['issuer'],period_label(r['period']),LABELS[r['metric']],num(r['value'])+'天' if r['status']=='ok' else r['reason']] for r in data['turnover']])
            doc.add_paragraph('年度平均余额采用（年初＋年末）/2，按本年实际天数计算。应收以营业收入代理赊销收入，存货以营业成本代理匹配成本；只属同口径财务近似，季节性及行业适用性需核对。缺年初余额不使用年末余额替代。')
        for question in data.get('investigation_questions',[]):doc.add_paragraph(question['issuer']+' '+period_label(question['period'])+'：'+question['observation']+question['question']+' '+question['threshold_basis'])
        if data.get('reconciliations') or data.get('causal_evidence'):doc.add_heading('勾稽与会计口径',1)
        else:doc.add_paragraph('本材料缺少可勾稽的资产、负债和总权益组合或相关会计说明；未将缺少核验误写为勾稽通过。')
        for c in data.get('reconciliations',[]):doc.add_paragraph(c['issuer']+' '+c['period']+f"：资产－负债－总权益差额为{num(c['difference'])}，显示精度容差为{num(c['tolerance'])}；"+('在容差内。' if c['passed'] else '超出容差，需要复核。'))
        for n in data.get('causal_evidence',[]):doc.add_paragraph(f"原文第{n['page']}页披露：{readable_excerpt(n['quote'])}\n此摘录仍需核对是否解释同期间指标，不自动视为确认原因。")
        doc.add_heading('输入问题与适用边界',1)
        if data.get('input_issues'):table(doc,['问题类别','受影响指标/口径'],[[human(i['code']),human(i.get('metric') or i.get('notice') or '见对应字段及JSON来源')] for i in data['input_issues']])
        else:doc.add_paragraph('本次未发现影响计算的输入冲突；仍应结合原文核对业务含义。')
        for limit in data.get('limitations',[]):doc.add_paragraph(limit)
    elif data['status']=='needs_input':
        doc.add_heading('优先补充的资料',1);doc.add_paragraph(human(data['reason']))
        for item in data.get('estimation_audit',{}).get('missing',[]):doc.add_paragraph(item['label']+'：'+item['impact'])
    else:
        a=data['assumptions'];d=data.get('dcf',{});audit=data.get('estimation_audit',{})
        doc.add_heading('经营假设及其依据',1)
        doc.add_paragraph(f"目标公司：{a['issuer']}；基期：{a['base_period']}；会计口径：{a.get('financial_basis','以已核验原表为准')}；币种：{a['currency']}；估值参考日：{a['valuation_date']}。")
        rows=[]
        for item in audit.get('parameters',[]):
            k=item['parameter'];v=item['value']
            if k not in ('growth_rates','ebit_margin','tax_rate','da_ratio','capex_ratio','nwc_ratio','wacc','terminal_growth','net_debt','shares_outstanding'):continue
            value=' / '.join(percent(g) for g in v) if k=='growth_rates' else num(v) if k in ('net_debt','shares_outstanding') else percent(v)
            rows.append([LABELS.get(k,k),value,ORIGINS.get(item['origin'],item['origin'])+'；'+item['basis']])
        table(doc,['参数','采用值','类型与依据'],rows,[1.15,1.4,4.1])
        doc.add_paragraph('历史估计与可比估计也是预测假设，探索情景缺少企业事实支持。所有参数、置信标记、范围和证据字段保存在JSON及假设工作表中，可以修改后重算。')
        doc.add_heading('现金流折现模型',1)
        if d.get('status')=='ok':
            doc.add_paragraph('FCFF = EBIT－max(EBIT,0)×税率＋折旧摊销－资本开支－经营营运资金增加。稳定期各项按永续增长率重新计算，再按WACC折现；权益价值还需企业价值与净债务、少数权益等项目桥接。')
            table(doc,['结果','数值'],[['企业价值（'+data['currency']+'）',num(d['enterprise_value'])],['权益价值（'+data['currency']+'）',num(d['equity_value'])],['每股价值（'+data['currency']+'/股）',num(d['price_per_share'])],['终值占企业价值',percent(d['terminal_weight'])]])
            doc.add_paragraph('下表金额统一为万'+data['currency']+'，预测年份为基期后的第1至第N年。')
            table(doc,['年','收入','EBIT','税后经营利润','折旧摊销','资本开支','营运资金增加','FCFF'],[[r['year']]+[num(D(r[k])/10000) for k in ('revenue','ebit','nopat','da','capex','delta_nwc','fcff')] for r in d['forecast']])
            doc.add_heading('折现率与增长率敏感性',1)
            unit='每股价值（'+data['currency']+'/股）' if d.get('price_per_share') is not None else '企业价值（万'+data['currency']+'）'
            doc.add_paragraph('以下为'+unit+'。网格最小值至最大值是条件区间，不是统计置信区间。')
            gs=a['terminal_growth_grid'];ws=a['wacc_grid'];lookup={(r['wacc'],r['terminal_growth']):r for r in d['sensitivity']}
            matrix=[]
            for w in ws:
                cells=[]
                for g in gs:
                    point=lookup[(str(w),str(g))]
                    cells.append('不适用' if point.get('status') not in (None,'ok') else num(point['price_per_share']) if d.get('price_per_share') is not None else num(D(point['enterprise_value'])/10000))
                matrix.append([percent(w)]+cells)
            table(doc,['WACC / 增长']+[percent(g) for g in gs],matrix)
            for point in d['sensitivity']:
                if point.get('reason'):doc.add_paragraph('WACC '+percent(point['wacc'])+'、增长 '+percent(point['terminal_growth'])+'：'+point['reason'])
        else:
            doc.add_paragraph(human(d.get('reason','DCF尚未形成结果')))
            if d.get('forecast'):
                doc.add_paragraph('保留预测期现金流，金额单位为万'+data['currency']+'；不能将预测期表格误认为有效的永续估值。')
                table(doc,['年','收入','EBIT','税后经营利润','折旧摊销','资本开支','营运资金增加','FCFF'],[[r['year']]+[num(D(r[k])/10000) for k in ('revenue','ebit','nopat','da','capex','delta_nwc','fcff')] for r in d['forecast']])
        doc.add_heading('可比公司与相对估值',1)
        for p in audit.get('peer_screening',[]):doc.add_paragraph(str(p['name'])+'：'+('入选' if p['included'] else '排除')+'；'+p['reason']+('；可比匹配分'+str(p['score']) if 'score' in p else ''))
        for peer in a.get('comparables',[]):doc.add_paragraph(str(peer.get('name'))+'；报价日：'+str(peer.get('as_of'))+'；依据：'+str(peer.get('basis')))
        rows=[]
        for r in data.get('relative',[]):
            if r.get('status')=='ok':rows.append([r['method'],num(r['median_multiple'])+'倍',num(r['equity_low']),num(r['equity_mid']),num(r['equity_high'])])
            else:doc.add_paragraph(r['method']+'暂不形成区间：利润/净资产基数或有效可比数量不足。')
        table(doc,['方法','中位倍数','权益下限','权益中值','权益上限'],rows)
        doc.add_paragraph('相对估值金额单位为'+data['currency']+'；利润口径与时点采用本任务提供的依据。未联网自动筛选全市场公司，未证明所提供公司一定可比。')
        doc.add_heading('缺失资料与结论边界',1)
        doc.add_heading('经营情景交叉检验',2)
        doc.add_paragraph('在基准增长路径上分别变动±5个百分点，同时经营利润率变动±3个百分点，其余参数相同。这是预设探索压力测试，参数依据和范围应继续用业务材料校准。')
        table(doc,['情景','企业价值','权益价值','每股价值'],[[s['label'],num(s.get('enterprise_value')),num(s.get('equity_value')),num(s.get('price_per_share'))] for s in data.get('operating_scenarios',[])])
        for scenario in data.get('operating_scenarios',[]):
            if scenario.get('reason'):doc.add_paragraph(scenario['label']+'情景没有有效DCF结果：'+scenario['reason'])
        for item in audit.get('missing',[]):doc.add_paragraph(item['label']+'：'+item['impact'])
        for warning in data.get('warnings',[]):doc.add_paragraph(warning)
    doc.add_heading('资料来源与复现',1)
    for source in data.get('source_snapshots',[]):
        doc.add_paragraph(source['source']['filename']+'；处理范围：'+('全部' if source['page_range']=='all' else str(source['page_range']))+'；SHA256：'+source['source']['sha256'])
        if source.get('user_confirmation'):doc.add_paragraph('币种采用用户明确确认的'+source['user_confirmation']['currency']+'，来源为对话补充，未改写原始材料。该确认不会解决其他期间、单位或数值冲突。')
    doc.add_paragraph('同目录JSON保存源字段、公式、工具输入与模型标识。原文事实、计算解释和未来假设分开存储，缺失值不填零。本报告随任务生成，后续修改输入应重新生成。')
    doc.save(path)
