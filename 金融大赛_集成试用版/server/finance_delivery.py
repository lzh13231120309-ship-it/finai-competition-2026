"""One verified Word per response. Technical artifacts remain private to the workflow."""
from copy import deepcopy
from pathlib import Path
import hashlib,json,re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from . import financial_agent as fa,research_agent as ra,report_documents as rd,presentation as p

def latest(task_id,exports,research):
    valid_exports={x['run_id']:x for x in exports if fa.export_is_current(task_id,x)}
    valid_research={x['mode']:x for x in research if ra.current(task_id,x)}
    return list(valid_exports.values()),list(valid_research.values())

def build(task_id,exports,research,intent='analysis',interruption=''):
    exports,research=latest(task_id,exports,research)
    signature={'exports':exports,'research':research,'intent':intent,'interruption':interruption,'format':'single-word-v4'}
    rid=hashlib.sha256(json.dumps(signature,ensure_ascii=False,sort_keys=True).encode()).hexdigest()[:32]
    folder=fa.task_folder(task_id)/'deliveries'/rid
    folder.mkdir(parents=True,exist_ok=True)
    payloads={x['mode']:json.loads((ra.scoped_folder(task_id,x['research_id'])/'result.json').read_text(encoding='utf-8')) for x in research}
    analysis=payloads.get('analysis');valuation=payloads.get('valuation')
    paragraphs=[]
    if analysis:paragraphs+=p.analysis_paragraphs(analysis)
    if valuation:paragraphs+=p.valuation_paragraphs(valuation)
    if not paragraphs and exports:
        text=fa.delivery_text(task_id,exports)
        # Retain the readable explanation; downloads are handled exactly once below.
        paragraphs=[line for line in text.splitlines() if line.strip() and not line.startswith('[下载')]
    if not paragraphs:
        paragraphs=['这次尚未形成足够可靠的财务结论。我不会把未核实的数值或模型草稿作为分析结果。报告说明了当前进度、缺少的依据和可以继续处理的事项。']
    issues=[]
    confirmations={s['source']['filename']:s['user_confirmation'] for payload in payloads.values() for s in payload.get('source_snapshots',[]) if s.get('user_confirmation')}
    for out in exports:
        _,_,source,rows=fa.completed(task_id,out['run_id'])
        validation=json.loads((fa.completed(task_id,out['run_id'])[0]/'agent_validation.json').read_text(encoding='utf-8'))
        fields={f['field_id']:f for f in rows}
        for item in validation.get('issues',[]):
            resolved_currency=source['source']['filename'] in confirmations and fields.get(item['field_id'],{}).get('currency') in (None,'','未知','未明确')
            reasons=[r for r in item['reasons'] if not (r=='currency_unknown' and resolved_currency)]
            if reasons:issues.append({**item,'reasons':reasons,'filename':source['source']['filename']})
    if issues:
        names=list(dict.fromkeys(p.LABELS.get(i['name'],'相关项目') for i in issues))
        paragraphs.append('仍需核对的项目包括'+'、'.join(names[:4])+'。这些问题只限制相关指标，已核验的其他数据仍可用于分析；具体原文位置和影响都列入同一份Word。')
    if interruption:paragraphs.append('本轮处理尚有未完成事项：'+p.human(interruption)+'。已保留可靠结果；详细报告区分已完成内容与待处理内容。')
    if intent=='valuation' and not valuation:
        paragraphs.append('估值尚未完成，'+('当前可交付的是已核验材料及财务分析' if analysis else '当前可交付的是已核验的材料说明' if exports else '当前只能说明处理进度与尚缺依据')+'，不能据此认定已经得到可靠股价。')
    title={'extraction':'金融材料解读报告','analysis':'财务分析报告','valuation':'财务研究与估值报告'}.get(intent,'金融研究报告')
    if not (folder/'report.docx').is_file():
        doc=rd.document(title)
        doc.add_paragraph('本报告集中呈现本次分析、判断依据、假设和仍待核对的事项。所有详细内容在这一份Word中，后台中间文件无需逐个打开。')
        if interruption:doc.add_paragraph('处理状态：部分完成。'+p.human(interruption))
        if not research and not exports:
            doc.add_heading('当前结论与处理情况',1)
            for paragraph in paragraphs:doc.add_paragraph(paragraph)
        parts=[]
        for mode in ('analysis','valuation'):
            if mode in payloads:
                target=folder/('_'+mode+'.docx');rd.write_research(payloads[mode],target)
                parts.append((('财务表现与解释' if mode=='analysis' else '估值方法与假设'),target))
        if not research:
            parts += [('材料解读',fa.completed(task_id,out['run_id'])[0]/'report.docx') for out in exports]
        for heading,path in parts:
            if not path.is_file():continue
            doc.add_heading(heading,1)
            part=Document(path)
            # Same style family and no external relationships in generated financial docs.
            for child in list(part.element.body)[2:]:
                if research and child.tag==qn('w:p') and ''.join(n.text or '' for n in child.iter(qn('w:t')))=='资料来源与复现':break
                if child.tag!=qn('w:sectPr'):doc.element.body.insert(-1,deepcopy(child))
        if issues:
            doc.add_heading('原文复核位置及问题',1)
            grouped={}
            for item in issues:
                key=(item['filename'],item['name'],tuple(item['reasons']))
                grouped.setdefault(key,set()).add(str(item.get('page','未知')))
            rd.table(doc,['材料','项目','原文位置','需要核对的内容'],[[filename,p.LABELS.get(name,'相关项目'),'第'+ '、'.join(sorted(pages))+'页',p.human('、'.join(reasons))] for (filename,name,reasons),pages in list(grouped.items())[:100]],[1.6,1.15,.65,3.25])
            if len(issues)>100:doc.add_paragraph('复核提示超过100项，本文列出前100项；这意味着材料质量尚需进一步核验，不能作为全部字段已经确认。')
        if exports and not research:
            doc.add_paragraph('以上已完成材料解读；本轮没有生成财务分析或估值，不把解读结果当作已完成估值。')
        if research:
            doc.add_heading('资料来源与复现',1)
            note=doc.add_paragraph('原文事实、计算解释和预测假设分别保存，缺失值不填零；来源指纹用于追溯。材料修改后应重新生成报告。')
            note.paragraph_format.space_after=Pt(3)
            note.paragraph_format.line_spacing=1.1
            for run in note.runs:run.font.size=Pt(9)
            sources={}
            for payload in payloads.values():
                for source in payload.get('source_snapshots',[]):
                    sources[(source['source']['filename'],source['source']['sha256'],str(source['page_range']))]=source
            for source in sources.values():
                doc.add_paragraph(source['source']['filename']+'；处理范围：'+('全部' if source['page_range']=='all' else str(source['page_range']))+'；SHA256：'+source['source']['sha256'])
                if source.get('user_confirmation'):doc.add_paragraph('币种采用用户明确确认的'+source['user_confirmation']['currency']+'，来源为对话补充，未改写原始材料。该确认不会解决其他期间、单位或数值冲突。')
        # Remove references inviting readers to open technical formats, not their audit data.
        for para in doc.paragraphs:
            if any(word in para.text for word in ('随JSON交付','保存在JSON','保存于JSON','CSV及JSON','同目录JSON','见对应字段及JSON')):
                text=para.text.replace('完整口径、原值、绝对变化及来源字段随JSON交付。','原文位置、口径与核查提示在本报告中说明，完整计算记录由后台保存。').replace('完整字段保存在CSV及JSON中。','完整字段由后台保存。').replace('保存在JSON及假设工作表中','由后台保存').replace('保存在JSON中','由后台保存').replace('见对应字段及JSON来源','见本报告原文复核说明')
                if text.startswith('同目录JSON'):
                    text='原文事实、计算解释和预测假设分别保存，缺失值不填零；来源指纹用于追溯。材料修改后应重新生成报告。'
                    para.paragraph_format.space_after=Pt(3)
                    para.paragraph_format.line_spacing=1.1
                para.text=text
                if text.startswith('原文事实、计算解释和预测假设'):
                    for run in para.runs:run.font.size=Pt(9)
        doc.save(folder/'report.docx')
        for path in folder.glob('_*.docx'):path.unlink()
    url=f'/api/finance-delivery/{task_id}/{rid}/report.docx'
    content='### '+('材料解读与分析' if intent=='extraction' else '分析与判断')+'\n\n'+'\n\n'.join(paragraphs)+'\n\n'+f'[下载完整分析 Word]({url})'+'\n'
    meta={**signature,'content':content,'title':title,'research_ids':[x['research_id'] for x in research]}
    fa.dump(folder/'delivery.json',meta)
    return {'content':content,'id':rid,'url':url,'title':title,'partial':bool(interruption or (valuation and valuation['status']=='needs_input'))}

def folder(task_id,rid):
    if not re.fullmatch('[a-f0-9]{32}',rid or ''):raise ValueError('报告标识无效')
    path=fa.task_folder(task_id)/'deliveries'/rid
    meta=json.loads((path/'delivery.json').read_text(encoding='utf-8'))
    exports,research=latest(task_id,meta['exports'],meta['research'])
    if len(exports)!=len(meta['exports']) or len(research)!=len(meta['research']):raise ValueError('材料已改变，请重新生成报告')
    return path,meta

def history(task_id,messages):
    """Present legacy replies through verified tools without changing stored conversations."""
    exports=[];research=[];intent=None;out=[]
    from .agent import financial_intent,financial_followup
    for msg in messages:
        item=dict(msg)
        if item['role']=='user':intent=financial_intent(item.get('content')) or (intent if financial_followup(item.get('content')) else None)
        if item['role']=='tool':
            try:
                value=json.loads(item['content']);name=(item.get('tool_calls') or {}).get('name')
                if value.get('ok') and name=='export_financial_result':exports.append(value)
                if value.get('ok') and name in ra.NAMES:research.append(value)
            except (ValueError,TypeError):pass
        if item['role']=='assistant' and not item.get('tool_calls') and intent and '/api/finance-delivery/' not in (item.get('content') or ''):
            valid_exports,valid_research=latest(task_id,exports,research)
            if valid_exports or valid_research:
                reason='此前这轮处理没有完成全部步骤，当前展示已核验的可用结果' if '⚠️' in (item.get('content') or '') else ''
                item['content']=build(task_id,valid_exports,valid_research,intent,reason)['content']
        out.append(item)
    return out
